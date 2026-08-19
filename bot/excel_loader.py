from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from itertools import chain, islice
from pathlib import Path
import threading
from typing import Any

from openpyxl import load_workbook


ALIASES = {
    "account": {"hisob raqam", "xisob raqam", "account", "abonent", "ls", "лицевой счет", "лиц.счет"},
    "fio": {"fio", "f.i.o", "fish", "f.i.sh", "iste'molchi", "istemolchi", "фио", "ф.и.о. владельца"},
    "phone": {"telefon", "phone", "tel", "телефон", "сотовый тел"},
    "new_phone": {"yangi telefon", "new phone", "новый тел"},
    "address": {"manzil", "address", "адрес"},
    "area": {"hudud", "uchastka", "mahalla", "area", "territory", "участок", "территория"},
    "mahalla": {"mahalla", "махалля"},
    "street": {"ko'cha", "kucha", "street", "улица"},
    "house": {"uy", "uy raqami", "house", "№ дома"},
    "flat": {"xonadon", "kvartira", "flat", "кв-ра"},
    "has_meter": {"hisoblagich", "с водомерами", "водомер"},
    "inspector_fio": {"inspektor fio", "inspector fio", "inspektor", "инспектор"},
    "inspector_phone": {"inspektor telefon", "inspector phone", "инспектор телефон"},
    "last_meter_value": {"oxirgi korsatgich", "oxirgi ko'rsatgich", "last meter", "показания"},
    "last_meter_date": {"korsatgich sana", "ko'rsatgich sana", "last meter date", "дата послед.показ"},
    "debit": {"debitor", "qarzdorlik", "общий дебит"},
    "credit": {"kreditor", "oldindan to'lov", "общий кредит"},
    "balance": {"balans", "balance", "сальдо", "баланс"},
    "updated_at": {"yuklangan sana", "updated at", "дата загрузки"},
}

REAL_ABONENT_COLUMN_MAP = {
    "mahalla": 0,
    "street": 1,
    "house": 2,
    "flat": 3,
    "fio": 4,
    "account": 5,
    "phone": 6,
    "new_phone": 7,
    "has_meter": 8,
    "last_meter_date": 9,
    "last_meter_value": 10,
    "debit": 11,
    "credit": 12,
    "area": 13,
    "inspector_fio": 14,
}


@dataclass(frozen=True)
class Abonent:
    account: str
    fio: str
    phone: str
    address: str
    area: str
    inspector_fio: str
    inspector_phone: str
    last_meter_value: str
    last_meter_date: str
    balance: str
    updated_at: str
    mahalla: str = ""
    street: str = ""
    house: str = ""
    flat: str = ""
    original_phone: str = ""
    edited_phone: str = ""
    has_meter: str = ""
    debit: str = ""
    credit: str = ""


class AbonentBook:
    def __init__(self, path: str, inspectors_path: str | None = None) -> None:
        self.path = Path(path)
        self.inspectors_path = Path(inspectors_path) if inspectors_path else None
        self.loaded_at = ""
        self._items: dict[str, Abonent] = {}
        self._write_lock = threading.RLock()
        self.reload()

    def count(self) -> int:
        return len(self._items)

    def all_items(self) -> list[Abonent]:
        return list(self._items.values())

    def reload(self) -> None:
        if not self.path.exists():
            raise FileNotFoundError(f"Excel fayl topilmadi: {self.path}")

        inspectors = self._load_inspectors()
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        items: dict[str, Abonent] = {}
        try:
            sheet = workbook.active
            row_iterator = sheet.iter_rows(values_only=True)
            preview = list(islice(row_iterator, 8))
            if not preview:
                raise RuntimeError("Excel fayl bo'sh")

            mapping = self._detect_mapping(preview)
            if "account" not in mapping:
                raise RuntimeError(
                    "Excel faylda hisob raqam ustuni topilmadi"
                )

            start_index = self._data_start_index(preview, mapping)
            rows = chain(preview, row_iterator)
            for row_index, row in enumerate(rows):
                if row_index < start_index:
                    continue
                account = self.clean_account(
                    self._cell(row, mapping, "account")
                )
                if not account:
                    continue
                fio = self._cell(row, mapping, "fio")
                if fio and self._normalize(fio) in {
                    "ф.и.о. владельца",
                    "fio",
                }:
                    continue

                area = self._cell(row, mapping, "area")
                inspector = inspectors.get(self._normalize(area), {})
                original_phone = self._cell(row, mapping, "phone")
                edited_phone = self._cell(row, mapping, "new_phone")
                debit = self._cell(row, mapping, "debit")
                credit = self._cell(row, mapping, "credit")
                address_parts = {
                    "mahalla": self._cell(row, mapping, "mahalla"),
                    "street": self._cell(row, mapping, "street"),
                    "house": self._cell(row, mapping, "house"),
                    "flat": self._cell(row, mapping, "flat"),
                }
                address = self._address(address_parts)

                items[account] = Abonent(
                    account=account,
                    fio=fio or "-",
                    phone=edited_phone or original_phone,
                    address=address,
                    area=area,
                    inspector_fio=(
                        inspector.get("fio")
                        or self._cell(row, mapping, "inspector_fio")
                    ),
                    inspector_phone=(
                        inspector.get("phone")
                        or self._cell(row, mapping, "inspector_phone")
                    ),
                    last_meter_value=self._cell(
                        row,
                        mapping,
                        "last_meter_value",
                    ),
                    last_meter_date=self._cell(
                        row,
                        mapping,
                        "last_meter_date",
                    ),
                    balance=self._balance(debit, credit),
                    updated_at=self._cell(row, mapping, "updated_at"),
                    mahalla=address_parts["mahalla"],
                    street=address_parts["street"],
                    house=address_parts["house"],
                    flat=address_parts["flat"],
                    original_phone=original_phone,
                    edited_phone=edited_phone,
                    has_meter=self._cell(row, mapping, "has_meter"),
                    debit=debit,
                    credit=credit,
                )
        finally:
            workbook.close()

        self._items = items
        timestamp = datetime.fromtimestamp(self.path.stat().st_mtime)
        self.loaded_at = timestamp.strftime("%d.%m.%Y")

    def find(self, account: str) -> Abonent | None:
        return self._items.get(self.clean_account(account))

    def update_edited_phone(self, account: str, phone: str) -> None:
        clean_account = self.clean_account(account)
        if not clean_account:
            raise ValueError("Hisob raqami noto'g'ri.")

        with self._write_lock:
            self._update_edited_phone_locked(clean_account, phone)

    def _update_edited_phone_locked(self, clean_account: str, phone: str) -> None:
        temporary = self.path.with_name(
            f".{self.path.stem}.phone.tmp{self.path.suffix}"
        )
        temporary.unlink(missing_ok=True)
        workbook = load_workbook(self.path)
        try:
            sheet = workbook.active
            preview = [
                tuple(cell.value for cell in row)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row, 8),
                )
            ]
            mapping = self._detect_mapping(preview)
            account_column = mapping.get("account")
            new_phone_column = mapping.get("new_phone")
            if account_column is None:
                raise RuntimeError(
                    "Excel faylda hisob raqami ustuni topilmadi."
                )
            if new_phone_column is None:
                raise RuntimeError(
                    "Excel faylda yangi telefon ustuni topilmadi."
                )

            matched_row = None
            for row_number in range(1, sheet.max_row + 1):
                value = sheet.cell(
                    row=row_number,
                    column=account_column + 1,
                ).value
                if self.clean_account(self._format(value)) == clean_account:
                    matched_row = row_number
                    break
            if matched_row is None:
                raise LookupError(
                    "Hisob raqami Excel fayldan topilmadi."
                )

            sheet.cell(
                row=matched_row,
                column=new_phone_column + 1,
                value=phone,
            )
            workbook.save(temporary)
        finally:
            workbook.close()

        try:
            temporary.replace(self.path)
        finally:
            temporary.unlink(missing_ok=True)
        self.reload()

    @staticmethod
    def clean_account(value: str) -> str:
        return "".join(ch for ch in value.strip() if ch.isdigit())

    @staticmethod
    def _normalize(value: Any) -> str:
        return str(value or "").strip().lower().replace("’", "'")

    @staticmethod
    def _format(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%d.%m.%Y")
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def _detect_mapping(self, rows: list[tuple[Any, ...]]) -> dict[str, int]:
        real_mapping = self._real_abonent_mapping(rows)
        if real_mapping:
            return real_mapping

        header_row = rows[0]
        headers = [self._normalize(value) for value in header_row]
        return self._build_mapping(headers)

    def _real_abonent_mapping(self, rows: list[tuple[Any, ...]]) -> dict[str, int] | None:
        for row in rows[:8]:
            normalized = [self._normalize(value) for value in row]
            if "махалля" in normalized and "лиц.счет" in normalized and "инспектор" in normalized:
                return REAL_ABONENT_COLUMN_MAP.copy()
        return None

    def _data_start_index(self, rows: list[tuple[Any, ...]], mapping: dict[str, int]) -> int:
        if mapping == REAL_ABONENT_COLUMN_MAP:
            return 0
        return 1

    def _build_mapping(self, headers: list[str]) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for index, header in enumerate(headers):
            for field, aliases in ALIASES.items():
                if header in aliases and field not in mapping:
                    mapping[field] = index
        return mapping

    def _cell(self, row: tuple[Any, ...], mapping: dict[str, int], field: str) -> str:
        index = mapping.get(field)
        if index is None or index >= len(row):
            return ""
        return self._format(row[index])

    def _address(self, parts: dict[str, str]) -> str:
        values = []
        if parts["mahalla"]:
            values.append(f"Mahalla: {parts['mahalla']}")
        if parts["street"]:
            values.append(parts["street"])
        if parts["house"]:
            values.append(f"{parts['house']}-uy")
        if parts["flat"]:
            values.append(f"{parts['flat']}-xonadon")
        return ", ".join(values)

    def _balance(self, debit: str, credit: str) -> str:
        debit_value = self._number(debit)
        credit_value = self._number(credit)
        if debit_value:
            return f"Qarzdorlik: {debit}"
        if credit_value:
            return f"Oldindan to'lov: {credit}"
        return "0"

    def _number(self, value: str) -> float:
        cleaned = value.replace(" ", "").replace(",", ".")
        try:
            return float(cleaned)
        except ValueError:
            return 0.0

    def _load_inspectors(self) -> dict[str, dict[str, str]]:
        if not self.inspectors_path or not self.inspectors_path.exists():
            return {}

        workbook = load_workbook(self.inspectors_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            return {}

        headers = [self._normalize(value) for value in rows[0]]
        mapping = self._build_mapping(headers)
        if "area" not in mapping:
            return {}

        inspectors: dict[str, dict[str, str]] = {}
        for row in rows[1:]:
            area = self._cell(row, mapping, "area")
            if not area:
                continue
            inspectors[self._normalize(area)] = {
                "fio": self._cell(row, mapping, "inspector_fio"),
                "phone": self._cell(row, mapping, "inspector_phone"),
            }
        return inspectors
