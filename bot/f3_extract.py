"""
To'liq F3 (104 ustunli) fayldan umumiy_shablon.xlsx (20 ustunli) yaratadi.

Agar umumiy_shablon.xlsx mavjud bo'lsa → faqat data qatorlari XML injection bilan
almashtiriladi (sarlavha formatlanishi va footer SUM qatori saqlanadi).
Agar mavjud bo'lmasa → openpyxl bilan yangi fayl yaratiladi (footer SUM qatori qo'shiladi).
"""
from __future__ import annotations

import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from bot.passport_report import (
    _XLSX_NS,
    _ci,
    _col_letter_to_idx,
    _excel_serial_to_date,
    _load_shared_strings,
    _extract_ref_date,
)


# ---------------------------------------------------------------------------
# Ustun mapping: umumiy_shablon pozitsiyasi → to'liq F3 ustun harfi
# ---------------------------------------------------------------------------

_UMUMIY_COLS: list[str] = [
    "B",    # 0  A: Mahalla
    "C",    # 1  B: Ko'cha
    "D",    # 2  C: Uy
    "E",    # 3  D: Kv
    "I",    # 4  E: FIO
    "K",    # 5  F: Hisob raqami (Лицевой счет)
    "M",    # 6  G: Telefon
    "U",    # 7  H: C vodomerom (+/-)
    "W",    # 8  I: Kol-vo lits Fakt (propiska)
    "X",    # 9  J: Kol-vo lits Tekush (prozhivayushchiy)
    "Z",    # 10 K: Kol.schet (hisoblagich soni)
    "AA",   # 11 L: Data posled.pokaz (oxirgi ko'rsatgich sanasi) — sana
    "AB",   # 12 M: Narastayushchiy (jami sarflangan)
    "CH",   # 13 N: Vsego nachisleno s NDS (hisob summasi)
    "CO",   # 14 O: Balans Debet (qarz summasi, >0)
    "CR",   # 15 P: Balans Kredit (kredit summasi, <0)
    "CV",   # 16 Q: Data posled. oplaty (oxirgi to'lov sanasi) — sana
    "CW",   # 17 R: Status (holat)
    "CX",   # 18 S: Adres (manzil)
    "CY",   # 19 T: Inspektor
]

_DATE_F3_COLS = {"AA", "CV"}
_ZERO_AS_NONE = {"W", "X", "Z", "AB", "CO", "CR"}

_F3_NEEDED_IDX: set[int] = {_ci(col) for col in _UMUMIY_COLS}
_F3_DATE_IDX: set[int] = {_ci(col) for col in _DATE_F3_COLS}
_F3_ZERO_IDX: set[int] = {_ci(col) for col in _ZERO_AS_NONE}

_HDR3 = [
    "Mahalla", "Ko'cha", "Uy", "Kv", "F.I.Sh. abonent",
    "Lits.schet", "Tel.nomer", "C vodomerom",
    "Kol-vo lits", "", "Po vodomerom", "", "",
    "Vsego nachisleno", "Balans", "", "Data posled.", "", "Adres", "Inspektor",
]
_HDR4 = [
    "", "", "", "", "", "", "", "",
    "Fakt", "Tekush", "Kol.schet", "Data pokaz.", "Narastayushchiy",
    "", "Debet", "Kredit", "", "", "", "",
]

_DATA_ROW_START = 6

# Raqam ustunlari pozitsiyalari (0-based) — footer SUM formulasi uchun
_SUM_COL_POSITIONS = {8, 9, 10, 12, 13, 14, 15}  # I,J,K,M,N,O,P

_EXCEL_EPOCH = date(1899, 12, 30)
_ROW_PATTERN = re.compile(rb"<row\b[^>]*\br=\"(\d+)\"[^>]*>.*?</row>", re.DOTALL)


# ---------------------------------------------------------------------------
# F3 XML Streaming → 20-ustunli qatorlar ro'yxati
# ---------------------------------------------------------------------------

def _read_f3_rows_for_shablon(path: Path) -> tuple[date | None, list[list]]:
    col_order = [_ci(c) for c in _UMUMIY_COLS]
    ref_date: date | None = None
    rows: list[list] = []

    with zipfile.ZipFile(str(path), "r") as z:
        shared = _load_shared_strings(z)
        with z.open("xl/worksheets/sheet1.xml") as f:
            row_buf: dict[int, object] = {}
            row_num = 0

            for event, elem in ET.iterparse(f, events=("start", "end")):
                tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag

                if event == "start" and tag == "row":
                    row_buf = {}
                    row_num = int(elem.get("r", 0))

                elif event == "end" and tag == "c":
                    ref = elem.get("r", "")
                    col_idx = _col_letter_to_idx(ref)
                    if col_idx not in _F3_NEEDED_IDX and (row_num != 1 or col_idx != 0):
                        elem.clear()
                        continue

                    typ = elem.get("t", "")
                    if typ == "inlineStr":
                        if col_idx not in _F3_DATE_IDX:
                            t_el = elem.find(f".//{{{_XLSX_NS}}}t")
                            raw = (t_el.text or "").strip() if t_el is not None else ""
                            if raw:
                                row_buf[col_idx] = raw
                    else:
                        v_el = elem.find(f"{{{_XLSX_NS}}}v")
                        raw = v_el.text if v_el is not None else None
                        if raw is not None:
                            if typ == "s":
                                idx_s = int(raw)
                                row_buf[col_idx] = shared[idx_s] if idx_s < len(shared) else ""
                            elif col_idx in _F3_DATE_IDX:
                                d = _excel_serial_to_date(raw)
                                if d is not None:
                                    row_buf[col_idx] = d
                            else:
                                try:
                                    row_buf[col_idx] = float(raw)
                                except ValueError:
                                    row_buf[col_idx] = raw
                    elem.clear()

                elif event == "end" and tag == "row":
                    if row_num == 1:
                        ref_date = _extract_ref_date(str(row_buf.get(0, "")))
                    elif row_num >= _DATA_ROW_START:
                        out = []
                        for ci in col_order:
                            val = row_buf.get(ci)
                            if val is None:
                                out.append(None)
                            elif isinstance(val, float):
                                if ci in _F3_ZERO_IDX and val == 0.0:
                                    out.append(None)
                                elif val == int(val):
                                    out.append(int(val))
                                else:
                                    out.append(val)
                            else:
                                out.append(val)
                        rows.append(out)
                    row_buf = {}
                    elem.clear()

    return ref_date, rows


# ---------------------------------------------------------------------------
# XML yordamchi funksiyalar
# ---------------------------------------------------------------------------

def _col_label(n: int) -> str:
    """1-indexed ustun raqamidan harf (A, B, ..., T, ...)."""
    result = ""
    while n:
        n, r = divmod(n - 1, 26)
        result = chr(65 + r) + result
    return result


def _values_to_row_xml(
    row_idx: int,
    values: list,
    date_style: int,
) -> bytes:
    """Python qiymatlar ro'yxatini Excel XML qatoriga aylantiradi."""
    cells: list[str] = []
    for col_pos, val in enumerate(values):
        if val is None:
            continue
        col = _col_label(col_pos + 1)
        ref = f"{col}{row_idx}"
        if isinstance(val, date):
            serial = (val - _EXCEL_EPOCH).days
            cells.append(f'<c r="{ref}" s="{date_style}" t="n"><v>{serial}</v></c>')
        elif isinstance(val, str):
            escaped = (
                val.replace("&", "&amp;")
                   .replace("<", "&lt;")
                   .replace(">", "&gt;")
                   .replace('"', "&quot;")
            )
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escaped}</t></is></c>')
        else:
            cells.append(f'<c r="{ref}" t="n"><v>{val}</v></c>')

    if not cells:
        return b""
    return (
        f'<row r="{row_idx}" spans="1:20">{"".join(cells)}</row>'
    ).encode("utf-8")


def _make_sum_footer_xml(footer_row: int, last_data_row: int) -> bytes:
    """SUM formulali footer qatori (N ustunlar uchun)."""
    cells: list[str] = []
    for pos in sorted(_SUM_COL_POSITIONS):
        col = _col_label(pos + 1)
        formula = f"SUM({col}{_DATA_ROW_START}:{col}{last_data_row})"
        cells.append(f'<c r="{col}{footer_row}"><f>{formula}</f></c>')
    if not cells:
        return f'<row r="{footer_row}"/>'.encode()
    return (
        f'<row r="{footer_row}" spans="1:20">{"".join(cells)}</row>'
    ).encode("utf-8")


def _find_date_style(template_bytes: bytes) -> int:
    """Template XML dan sana format style indeksini topadi (s="N")."""
    # L ustun date cell → s atribut qiymatini o'qish
    for row_num in range(_DATA_ROW_START, _DATA_ROW_START + 200):
        cell = re.search(
            rb'<c\b[^>]*r="L' + str(row_num).encode() + rb'"[^>]*s="(\d+)"',
            template_bytes,
        )
        if cell:
            return int(cell.group(1))
    return 3  # default openpyxl date style


def _update_title_in_prefix(prefix: bytes, ref_date: date | None) -> bytes:
    """Sarlavha qatori ichidagi sanani yangilaydi."""
    if not ref_date:
        return prefix
    new_title = f"F3 HISOBOT  {ref_date.strftime('%d.%m.%Y')}".encode("utf-8")
    return re.sub(
        rb'(<c\s+r="A1"[^>]*t="inlineStr"><is><t>)[^<]*(</t>)',
        rb"\g<1>" + new_title + rb"\g<2>",
        prefix,
        count=1,
    )


# ---------------------------------------------------------------------------
# Asosiy: XML injection orqali mavjud templateni yangilash
# ---------------------------------------------------------------------------

def _inject_data_xml(
    template_path: Path,
    data_rows: list[list],
    ref_date: date | None,
    output_path: Path,
) -> None:
    """
    Mavjud umumiy_shablon.xlsx templatega F3 data qatorlarini kiritadi.
    Sarlavha formatlanishi (rows 1-5) va footer (SUM) qatori saqlanadi.
    """
    with zipfile.ZipFile(str(template_path), "r") as tmpl:
        sheet_xml = tmpl.read("xl/worksheets/sheet1.xml")

        dim = re.search(br'<dimension ref="A1:T(\d+)"', sheet_xml)
        if not dim:
            raise RuntimeError("Template dimension A1:T{N} topilmadi")
        old_footer_row = int(dim.group(1))

        sdo = sheet_xml.find(b"<sheetData>")
        sdc = sheet_xml.find(b"</sheetData>")
        if sdo < 0 or sdc < 0:
            raise RuntimeError("sheetData topilmadi")

        content_start = sdo + len(b"<sheetData>")
        prefix = sheet_xml[:content_start]
        content = sheet_xml[content_start:sdc]
        suffix = sheet_xml[sdc:]

        date_style = _find_date_style(sheet_xml)

        # Header qatorlarini (1-5) saqlaymiz
        header_rows_bytes = b""
        old_footer_bytes = b""
        for m in _ROW_PATTERN.finditer(content):
            rnum = int(m.group(1))
            if rnum <= 5:
                header_rows_bytes += m.group(0)
            elif rnum == old_footer_row:
                old_footer_bytes = m.group(0)

        # Yangi data qatorlari
        new_data_bytes = b""
        with tempfile.TemporaryFile() as tf:
            for row_idx, values in enumerate(data_rows, _DATA_ROW_START):
                row_xml = _values_to_row_xml(row_idx, values, date_style)
                if row_xml:
                    tf.write(row_xml)
            tf.seek(0)
            new_data_bytes = tf.read()

        new_last_data_row = _DATA_ROW_START + len(data_rows) - 1
        new_footer_row = new_last_data_row + 1

        # Footer qatorini yangilash (row number va SUM oraliqlarini o'zgartirish)
        # Agar footer <f> (formula) ni o'z ichiga olmasa → u data qator, SUM yangi yaratiladi
        has_formula = b"<f>" in old_footer_bytes if old_footer_bytes else False

        if old_footer_bytes and has_formula:
            new_footer = old_footer_bytes
            # Row number yangilash
            new_footer = re.sub(
                rb'(<row\b[^>]*\br=")' + str(old_footer_row).encode() + rb'"',
                rb'\g<1>' + str(new_footer_row).encode() + rb'"',
                new_footer, count=1,
            )
            # Cell reference raqamlarini yangilash
            new_footer = re.sub(
                rb'r="([A-Z]+)' + str(old_footer_row).encode() + rb'"',
                lambda m: b'r="' + m.group(1) + str(new_footer_row).encode() + b'"',
                new_footer,
            )
            # SUM oraliqlarini yangilash (N6:N{old_last} → N6:N{new_last})
            old_last = str(old_footer_row - 1).encode()
            new_last = str(new_last_data_row).encode()
            new_footer = new_footer.replace(old_last, new_last)
        else:
            # Template da SUM footer yo'q — yangi yaratamiz
            new_footer = _make_sum_footer_xml(new_footer_row, new_last_data_row)

        # Sarlavhadagi sanani yangilash
        prefix = _update_title_in_prefix(prefix, ref_date)

        # Dimension yangilash
        prefix = re.sub(
            br'<dimension ref="A1:T\d+"',
            f'<dimension ref="A1:T{new_footer_row}"'.encode(),
            prefix, count=1,
        )

        new_sheet_xml = prefix + header_rows_bytes + new_data_bytes + new_footer + suffix

        # Temp faylga yozib, keyin asl faylni almashtiramiz
        import shutil as _shutil
        tmp = output_path.parent / (output_path.stem + "_tmp_inject.xlsx")
        try:
            with zipfile.ZipFile(
                str(tmp), "w",
                compression=zipfile.ZIP_DEFLATED,
                compresslevel=6,
            ) as out_zip:
                for item in tmpl.infolist():
                    if item.filename == "xl/calcChain.xml":
                        continue
                    data = tmpl.read(item.filename)
                    if item.filename == "xl/worksheets/sheet1.xml":
                        data = new_sheet_xml
                    out_zip.writestr(item, data)
        except Exception:
            tmp.unlink(missing_ok=True)
            raise
    _shutil.move(str(tmp), str(output_path))


# ---------------------------------------------------------------------------
# Yangi yaratish (birinchi marta, template yo'q bo'lganda)
# ---------------------------------------------------------------------------

def _make_header_rows(ws, ref_date: date | None) -> None:
    """Rows 1-5: sarlavha va ustun nomlari."""
    thin = Side(style="thin")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    bold = Font(bold=True, size=9)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title = "F3 HISOBOT" + (f"  {ref_date.strftime('%d.%m.%Y')}" if ref_date else "")
    ws.merge_cells("A1:T1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=11)
    ws["A1"].alignment = center

    for col_idx, (h3, h4) in enumerate(zip(_HDR3, _HDR4), 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        c3 = ws[f"{col_letter}3"]
        c3.value = h3
        c3.font = bold
        c3.alignment = center
        c3.fill = hdr_fill
        c3.border = brd
        c4 = ws[f"{col_letter}4"]
        c4.value = h4 if h4 else None
        c4.font = bold
        c4.alignment = center
        c4.fill = hdr_fill
        c4.border = brd


def _create_from_scratch(
    rows: list[list],
    ref_date: date | None,
    output_path: Path,
) -> None:
    """Yangi umumiy_shablon.xlsx yaratadi (birinchi yuklash yoki template yo'q)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F3"

    _make_header_rows(ws, ref_date)

    # Row 5 — bo'sh ajratgich
    ws.append([None] * 20)

    # Ustun kengliklari
    widths = [20, 20, 8, 5, 30, 15, 12, 5, 5, 5, 5, 12, 10, 14, 14, 14, 12, 15, 20, 25]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Data rows 6+
    for row_data in rows:
        ws.append(row_data)

    # Footer SUM qatori
    last_data_row = _DATA_ROW_START + len(rows) - 1
    footer_row_num = last_data_row + 1
    for pos in _SUM_COL_POSITIONS:
        col_letter = openpyxl.utils.get_column_letter(pos + 1)
        ws.cell(footer_row_num, pos + 1).value = (
            f"=SUM({col_letter}{_DATA_ROW_START}:{col_letter}{last_data_row})"
        )

    wb.save(str(output_path))


# ---------------------------------------------------------------------------
# Ochiq API
# ---------------------------------------------------------------------------

def extract_umumiy_shablon(
    f3_path: str | Path,
    output_path: str | Path,
) -> tuple[int, date | None]:
    """
    To'liq F3 dan umumiy_shablon.xlsx yaratadi yoki yangilaydi.

    - Agar output_path mavjud va A1:T{N} dimension bo'lsa →
        faqat data qatorlari (6-dan footer-1 gacha) almashtiriladi,
        sarlavha + footer formatlanishi TO'LIQ SAQLANADI.
    - Agar mavjud bo'lmasa → yangi fayl yaratiladi (SUM footer bilan).

    Returns: (data_rows_count, reference_date)
    """
    source = Path(f3_path)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    ref_date, rows = _read_f3_rows_for_shablon(source)

    # Mavjud template tekshirish
    _can_inject = False
    if output.exists():
        try:
            with zipfile.ZipFile(str(output), "r") as z:
                xml = z.read("xl/worksheets/sheet1.xml")
                if re.search(br'<dimension ref="A1:T\d+"', xml):
                    _can_inject = True
        except Exception:
            pass

    if _can_inject:
        _inject_data_xml(output, rows, ref_date, output)
    else:
        _create_from_scratch(rows, ref_date, output)

    return len(rows), ref_date
