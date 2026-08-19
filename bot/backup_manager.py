from __future__ import annotations

import hashlib
import json
import shutil
import sqlite3
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from bot.encryption import decrypt_file, encrypt_file
from bot.security_hardening import validate_safe_xlsx


MANIFEST_NAME = "manifest.json"
MANIFEST_VERSION = 1


@dataclass(frozen=True)
class BackupFile:
    label: str
    source: Path
    backup_name: str
    required: bool = True


@dataclass(frozen=True)
class BackupSet:
    directory: Path
    created_at: datetime
    reason: str
    valid: bool
    files: tuple[Path, ...]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_backup_child(folder: Path, name: object) -> Path:
    text = str(name or "")
    if not text or Path(text).name != text:
        raise RuntimeError("Backup manifestidagi fayl nomi xavfsiz emas.")
    path = folder / text
    if path.parent.resolve() != folder.resolve():
        raise RuntimeError("Backup fayli papkadan tashqariga chiqmoqda.")
    return path


def validate_sqlite(path: Path) -> None:
    with closing(
        sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    ) as database:
        result = database.execute("pragma integrity_check").fetchone()
    if not result or str(result[0]).lower() != "ok":
        raise RuntimeError(f"SQLite integrity_check muvaffaqiyatsiz: {result}")


def validate_excel(path: Path) -> None:
    validate_safe_xlsx(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if not workbook.sheetnames:
            raise RuntimeError(f"Excel sheet topilmadi: {path.name}")
        sheet = workbook[workbook.sheetnames[0]]
        if sheet.max_row < 1 or sheet.max_column < 1:
            raise RuntimeError(f"Excel bo'sh: {path.name}")
    finally:
        workbook.close()


def validate_backup_file(label: str, path: Path) -> None:
    if label == "database":
        validate_sqlite(path)
    elif label in {"abonents", "inspectors"} or ".xlsx" in path.suffixes:
        validate_excel(path)


def create_backup_set(
    backup_root: str | Path,
    files: Iterable[BackupFile],
    reason: str,
    now: datetime | None = None,
    encryption_password: str | None = None,
) -> BackupSet:
    created_at = now or datetime.now()
    root = Path(backup_root)
    root.mkdir(parents=True, exist_ok=True)
    directory = root / created_at.strftime("%Y%m%d_%H%M%S_%f")
    directory.mkdir()

    manifest_files = []
    created_files: list[Path] = []
    try:
        for item in files:
            if not item.source.exists():
                if item.required:
                    raise FileNotFoundError(f"Backup manbasi topilmadi: {item.source}")
                continue
            plaintext_target = directory / item.backup_name
            temporary = plaintext_target.with_name(
                f".{plaintext_target.stem}.backup.tmp{plaintext_target.suffix}"
            )
            if item.label == "database":
                with closing(sqlite3.connect(item.source)) as source_database:
                    with closing(sqlite3.connect(temporary)) as target_database:
                        source_database.backup(target_database)
            else:
                shutil.copy2(item.source, temporary)
            validate_backup_file(item.label, temporary)
            encrypted = bool(encryption_password)
            if encrypted:
                target = directory / f"{item.backup_name}.uwenc"
                encrypt_file(
                    temporary,
                    target,
                    encryption_password or "",
                    label=item.label,
                )
                temporary.unlink(missing_ok=True)
            else:
                target = plaintext_target
                temporary.replace(target)
            created_files.append(target)
            manifest_files.append(
                {
                    "label": item.label,
                    "source": str(item.source),
                    "name": target.name,
                    "size": target.stat().st_size,
                    "sha256": sha256_file(target),
                    "encrypted": encrypted,
                    "original_name": item.backup_name,
                }
            )

        manifest = {
            "version": MANIFEST_VERSION,
            "created_at": created_at.isoformat(timespec="seconds"),
            "reason": reason,
            "valid": True,
            "files": manifest_files,
        }
        (directory / MANIFEST_NAME).write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return BackupSet(
            directory=directory,
            created_at=created_at,
            reason=reason,
            valid=True,
            files=tuple(created_files),
        )
    except Exception:
        shutil.rmtree(directory, ignore_errors=True)
        raise


def load_backup_set(
    directory: str | Path,
    verify_hashes: bool = True,
    encryption_password: str | None = None,
) -> BackupSet:
    folder = Path(directory)
    manifest_path = folder / MANIFEST_NAME
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if int(manifest.get("version") or 0) != MANIFEST_VERSION:
        raise RuntimeError("Backup manifest versiyasi noto'g'ri.")
    if manifest.get("valid") is not True:
        raise RuntimeError("Backup manifest sog'lom deb belgilanmagan.")
    manifest_files = manifest.get("files")
    if not isinstance(manifest_files, list) or not manifest_files:
        raise RuntimeError("Backup manifestida fayllar yo'q.")
    files: list[Path] = []
    labels: set[str] = set()
    for item in manifest_files:
        if not isinstance(item, dict):
            raise RuntimeError("Backup manifest fayl yozuvi noto'g'ri.")
        label = str(item.get("label") or "")
        if not label or label in labels:
            raise RuntimeError("Backup manifestida takroriy fayl turi bor.")
        labels.add(label)
        path = _safe_backup_child(folder, item.get("name"))
        if not path.exists():
            raise FileNotFoundError(f"Backup fayli topilmadi: {path}")
        if path.stat().st_size != int(item.get("size") or -1):
            raise RuntimeError(f"Backup hajmi mos emas: {path.name}")
        if verify_hashes and sha256_file(path) != str(item.get("sha256") or ""):
            raise RuntimeError(f"Backup hash mos emas: {path.name}")
        if item.get("encrypted"):
            if not encryption_password:
                raise RuntimeError("Shifrlangan backup uchun parol berilmagan.")
            original_name = str(item.get("original_name") or "")
            if not original_name or Path(original_name).name != original_name:
                raise RuntimeError(
                    "Backupdagi asl fayl nomi xavfsiz emas."
                )
            temporary = folder / f".verify_{original_name}"
            try:
                decrypt_file(
                    path,
                    temporary,
                    encryption_password,
                    expected_label=label,
                )
                validate_backup_file(label, temporary)
            finally:
                temporary.unlink(missing_ok=True)
        else:
            validate_backup_file(label, path)
        files.append(path)
    return BackupSet(
        directory=folder,
        created_at=datetime.fromisoformat(manifest["created_at"]),
        reason=str(manifest.get("reason") or ""),
        valid=bool(manifest.get("valid")),
        files=tuple(files),
    )


def list_backup_sets(
    backup_root: str | Path,
    encryption_password: str | None = None,
) -> list[BackupSet]:
    root = Path(backup_root)
    if not root.exists():
        return []
    sets: list[BackupSet] = []
    for directory in root.iterdir():
        if not directory.is_dir() or not (directory / MANIFEST_NAME).exists():
            continue
        try:
            sets.append(
                load_backup_set(
                    directory,
                    verify_hashes=False,
                    encryption_password=encryption_password,
                )
            )
        except Exception:
            continue
    return sorted(sets, key=lambda item: item.created_at, reverse=True)


def latest_valid_backup(
    backup_root: str | Path,
    encryption_password: str | None = None,
) -> BackupSet | None:
    for backup in list_backup_sets(backup_root, encryption_password):
        try:
            return load_backup_set(
                backup.directory,
                verify_hashes=True,
                encryption_password=encryption_password,
            )
        except Exception:
            continue
    return None


def apply_retention(
    backup_root: str | Path,
    keep_sets: int,
    retention_days: int,
    now: datetime | None = None,
    encryption_password: str | None = None,
) -> list[Path]:
    current = now or datetime.now()
    backups = list_backup_sets(backup_root, encryption_password)
    removed: list[Path] = []
    for index, backup in enumerate(backups):
        too_many = index >= max(keep_sets, 1)
        too_old = current - backup.created_at > timedelta(
            days=max(retention_days, 1)
        )
        if too_many or too_old:
            shutil.rmtree(backup.directory)
            removed.append(backup.directory)
    return removed


def restore_backup_set(
    backup: BackupSet,
    destinations: dict[str, Path],
    encryption_password: str | None = None,
) -> list[Path]:
    verified = load_backup_set(
        backup.directory,
        verify_hashes=True,
        encryption_password=encryption_password,
    )
    manifest = json.loads(
        (verified.directory / MANIFEST_NAME).read_text(encoding="utf-8")
    )
    staged: list[tuple[str, Path, Path]] = []
    rollback_files: list[tuple[Path, Path | None]] = []
    try:
        for item in manifest["files"]:
            label = str(item["label"])
            destination = destinations.get(label)
            if destination is None:
                continue
            source = _safe_backup_child(
                verified.directory,
                item.get("name"),
            )
            destination.parent.mkdir(parents=True, exist_ok=True)
            temporary = destination.with_name(
                f".{destination.stem}.restore.tmp{destination.suffix}"
            )
            temporary.unlink(missing_ok=True)
            if item.get("encrypted"):
                if not encryption_password:
                    raise RuntimeError(
                        "Shifrlangan backupni tiklash uchun parol berilmagan."
                    )
                decrypt_file(
                    source,
                    temporary,
                    encryption_password,
                    expected_label=label,
                )
            else:
                shutil.copy2(source, temporary)
            validate_backup_file(label, temporary)
            staged.append((label, temporary, destination))

        restored: list[Path] = []
        try:
            for _, temporary, destination in staged:
                rollback = destination.with_name(
                    f".{destination.name}.restore.rollback"
                )
                rollback.unlink(missing_ok=True)
                if destination.exists():
                    shutil.copy2(destination, rollback)
                    rollback_files.append((destination, rollback))
                else:
                    rollback_files.append((destination, None))
                temporary.replace(destination)
                restored.append(destination)
            return restored
        except Exception:
            for destination, rollback in reversed(rollback_files):
                if rollback is None:
                    destination.unlink(missing_ok=True)
                elif rollback.exists():
                    rollback.replace(destination)
            raise
    finally:
        for _, temporary, _ in staged:
            temporary.unlink(missing_ok=True)
        for _, rollback in rollback_files:
            if rollback is not None:
                rollback.unlink(missing_ok=True)
