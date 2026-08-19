from __future__ import annotations

import csv
from collections import defaultdict
from functools import wraps
import weakref
import html
import io
import json
import logging
from datetime import datetime, timedelta
from pathlib import Path
import shutil
import sys
import threading
import time
import traceback
from typing import Any
from uuid import uuid4
import zipfile

import telebot
from openpyxl import load_workbook
from telebot import types
from telebot.apihelper import ApiTelegramException

from bot.backup_manager import (
    BackupFile,
    apply_retention,
    create_backup_set,
    latest_valid_backup,
    list_backup_sets,
    load_backup_set,
    restore_backup_set,
)
from bot.buttons import (
    ADMIN_STATS,
    ADMIN_RELOAD,
    ADMIN_SEARCH,
    ADMIN_BROADCAST,
    ADMIN_RECENT_SUBMISSIONS,
    ADMIN_PENDING_USERS,
    ADMIN_UPLOAD_ABONENTS,
    ADMIN_UPLOAD_INSPECTORS,
    ADMIN_ADD_OPERATOR,
    ADMIN_LIST_OPERATORS,
    ADMIN_REMOVE_OPERATOR,
    ADMIN_CLEAR_TEST_DATA,
    ADMIN_DEBT_ANALYSIS,
    ADMIN_DEBT_SEND,
    ADMIN_AUDIT_LOGS,
    ADMIN_AUDIT_EXPORT,
    ADMIN_BACKUP,
    ADMIN_BACKUP_LIST,
    ADMIN_BACKUP_RESTORE,
    ADMIN_SYSTEM_STATUS,
    ADMIN_USER_STATS,
    ADMIN_SETTINGS,
    ADMIN_DEBT_HISTORY,
    ADMIN_EXPORT_USERS,
    ADMIN_EXPORT_DEBTORS,
    ADMIN_EXPORT_SUBMISSIONS,
    ADMIN_EXPORT_DEBT_HISTORY,
    ADMIN_INSPECTORS,
    ADMIN_ADD_INSPECTOR,
    ADMIN_LIST_INSPECTORS,
    ADMIN_INSPECTOR_LOCATION_COMMAND,
    ADMIN_INSPECTOR_LOCATION_STATUS,
    INSPECTOR_LOCATION_STATUS,
    INSPECTOR_LOCATION_HELP,
    INSPECTOR_EXIT,
    ADMIN_SECTION_EXCEL,
    ADMIN_SECTION_ABONENTS,
    ADMIN_SECTION_REQUESTS,
    ADMIN_SECTION_OPERATORS,
    ADMIN_SECTION_BROADCAST,
    ADMIN_SECTION_SERVICE,
    ADMIN_SECTION_REPORTS,
    ADMIN_SECTION_TABLE_REPORTS,
    ADMIN_TABLE_TEMPLATES,
    ADMIN_TABLE_UPLOAD_F3,
    ADMIN_TABLE_DISTRICT_PASSPORT,
    ADMIN_TEMPLATE_THREE_MONTHS,
    ADMIN_TEMPLATE_DEBT_LIST,
    ADMIN_TEMPLATE_CREDITOR_LIST,
    ADMIN_TEMPLATE_NO_ACCOUNTING,
    ADMIN_TABLE_BACK,
    ADMIN_BACK,
    ADMIN_EXIT,
    CONFIRM_SEND,
    CANCEL,
    CLEAR_TEST_CONFIRM,
    DEBT_SEND_CONFIRM,
    EXCEL_REPLACE_CONFIRM,
    REMOVE_OPERATOR_CONFIRM,
    BACKUP_RESTORE_CONFIRM,
    LANGUAGE_BY_BUTTON,
    LANGUAGE_BUTTONS,
)
from bot.config import Config, load_config
from bot.code_integrity import verify_code_integrity
from bot.debt_texts import debt_warning_text
from bot.encryption import decrypt_bytes, encrypt_bytes, encrypt_file
from bot.excel_loader import Abonent, AbonentBook
from bot.formatters import balance_text, cabinet_text, meter_intro_text, report_caption
from bot.health_checks import (
    PreflightReport,
    preflight_summary,
    run_preflight_checks,
    write_preflight_report,
)
from bot.masking import (
    mask_account,
    mask_telegram_id,
    masked_row,
    sanitize_audit_details,
)
from bot.image_matching import InvalidMeterImage, compare_photo_sets, fingerprint_image
from bot.meter_vault import load_meter_photo, remove_meter_photo, store_meter_photo
from bot.security import (
    ADMIN_BLOCK_MINUTES,
    admin_access_blocked,
    register_denied_attempt,
    session_deadline,
)
from bot.security_hardening import (
    redact_sensitive_text,
    safe_exception_summary,
    validate_safe_xlsx,
)
from bot.runtime import (
    AlreadyRunningError,
    ProcessLock,
    health_worker,
    install_shutdown_handlers,
    write_health_status,
)
from bot.rate_limit import (
    RequestRateLimiter,
    file_size_allowed,
    item_count_allowed,
)
from bot.table_reports import (
    generate_creditor_report,
    generate_debt_report,
    generate_no_accounting_report,
    generate_three_month_report,
    latest_creditor_report,
    latest_debt_report,
    latest_no_accounting_report,
    latest_three_month_report,
)
from bot.passport_report import (
    generate_district_passport,
    latest_district_passport,
    validate_f3_full,
)
from bot.f3_extract import extract_umumiy_shablon
from bot import states
from bot.storage import BotStorage
from bot.texts import GUIDE, TEXTS, WELCOME

config: Config
book: AbonentBook
storage: BotStorage
bot: telebot.TeleBot
sessions: dict[int, dict[str, Any]] = {}
session_lock = threading.RLock()
_user_message_locks_store: weakref.WeakValueDictionary[int, threading.RLock] = weakref.WeakValueDictionary()
_user_message_locks_guard = threading.Lock()
maintenance_lock = threading.RLock()
error_log_lock = threading.RLock()
debt_job_lock = threading.Lock()
shutdown_event = threading.Event()
process_started_at = ""
last_preflight_status = ""
request_limiter: RequestRateLimiter
_location_last_update: dict[int, float] = {}
_LOCATION_THROTTLE_SECONDS = 30
ERROR_LOG_MAX_BYTES = 2 * 1024 * 1024
ERROR_LOG_BACKUP_COUNT = 5
INSPECTOR_LOCATION_RESPONSE_MINUTES = 10
INSPECTOR_LOCATION_FINAL_GRACE_MINUTES = 10
INSPECTOR_MINIMUM_LIVE_PERIOD_SECONDS = 60 * 60


def keyboard(rows: list[list[types.KeyboardButton | str]], one_time: bool = False) -> types.ReplyKeyboardMarkup:
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=one_time)
    for row in rows:
        markup.add(*row)
    return markup


def start_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([["▶️ Start"], [TEXTS["uz"]["service"]]], one_time=True)


def language_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[button] for button in LANGUAGE_BUTTONS], one_time=True)


def main_menu(language: str) -> types.ReplyKeyboardMarkup:
    text = TEXTS[language]
    return keyboard(
        [
            [text["cabinet"]],
            [text["send_meter"]],
            [text["balance"], text["my_accounts"]],
            [text["service"]],
        ]
    )


def cabinet_keyboard(language: str) -> types.ReplyKeyboardMarkup:
    text = TEXTS[language]
    return keyboard([[text["edit_phone"]], [text["back"]]])


def service_keyboard(language: str) -> types.ReplyKeyboardMarkup:
    text = TEXTS[language]
    return keyboard([[text["contact_center"]], [text["work_hours"]], [text["faq"]], [text["change_language"]], [text["guide"]], [text["back"]]])


def public_service_keyboard(language: str = "uz") -> types.ReplyKeyboardMarkup:
    text = TEXTS.get(language, TEXTS["uz"])
    return keyboard([[text["contact_center"]], [text["work_hours"]], [text["faq"]], [text["back"]]])


def phone_keyboard(language: str) -> types.ReplyKeyboardMarkup:
    text = TEXTS[language]
    contact_button = types.KeyboardButton(text["share_phone"], request_contact=True)
    return keyboard([[contact_button], [text["back"]]])


def meter_keyboard(language: str) -> types.ReplyKeyboardMarkup:
    text = TEXTS[language]
    return keyboard([[text["finish"]], [text["back"]]])


def meter_review_keyboard(submission_id: int) -> types.InlineKeyboardMarkup:
    markup = types.InlineKeyboardMarkup()
    markup.add(
        types.InlineKeyboardButton(
            "✅ To'g'ri rasm",
            callback_data=f"meterapprove:{submission_id}",
        ),
        types.InlineKeyboardButton(
            "❌ Noto'g'ri rasm",
            callback_data=f"meterreject:{submission_id}",
        ),
    )
    return markup


def approval_keyboard(telegram_id: int, account: str) -> types.InlineKeyboardMarkup:
    markup = types.InlineKeyboardMarkup()
    markup.add(
        types.InlineKeyboardButton(
            "✅ Tasdiqlash",
            callback_data=f"approve:{telegram_id}:{account}",
        ),
        types.InlineKeyboardButton(
            "❌ Rad etish",
            callback_data=f"reject:{telegram_id}:{account}",
        ),
    )
    return markup


def account_list_keyboard(
    language: str,
    accounts: list[dict[str, str]],
) -> tuple[types.ReplyKeyboardMarkup, dict[str, str]]:
    text = TEXTS[language]
    mapping: dict[str, str] = {}
    rows: list[list[str]] = []
    for index, item in enumerate(accounts, 1):
        label = f"🏠 {index}. {mask_account(item['account'])}"
        mapping[label] = item["account"]
        rows.append([label])
    rows.extend([[text["add_account"]], [text["back"]]])
    return keyboard(rows), mapping


def admin_keyboard(user_id: int | None = None) -> types.ReplyKeyboardMarkup:
    rows = [
        [ADMIN_SECTION_EXCEL, ADMIN_SECTION_ABONENTS],
        [ADMIN_SECTION_REQUESTS],
    ]
    if user_id is None or user_id in config.admin_ids:
        rows.extend(
            [
                [ADMIN_SECTION_OPERATORS],
                [ADMIN_SECTION_BROADCAST],
                [ADMIN_SECTION_REPORTS, ADMIN_SECTION_TABLE_REPORTS],
                [ADMIN_INSPECTORS],
                [ADMIN_SECTION_SERVICE],
            ]
        )
    rows.append([ADMIN_EXIT])
    return keyboard(rows)


def excel_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    return keyboard([[ADMIN_RELOAD], [ADMIN_UPLOAD_ABONENTS], [ADMIN_UPLOAD_INSPECTORS], [ADMIN_BACK]])


def abonents_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    return keyboard([[ADMIN_SEARCH], [ADMIN_BACK]])


def requests_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    return keyboard([[ADMIN_PENDING_USERS], [ADMIN_RECENT_SUBMISSIONS], [ADMIN_BACK]])


def operators_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard([[ADMIN_ADD_OPERATOR], [ADMIN_LIST_OPERATORS], [ADMIN_REMOVE_OPERATOR], [ADMIN_BACK]])


def broadcast_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard([[ADMIN_BROADCAST], [ADMIN_BACK]])


def admin_input_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[CANCEL], [ADMIN_BACK]])


def service_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard([
        [ADMIN_STATS, ADMIN_SYSTEM_STATUS],
        [ADMIN_USER_STATS, ADMIN_SETTINGS],
        [ADMIN_DEBT_ANALYSIS],
        [ADMIN_DEBT_SEND],
        [ADMIN_BACKUP, ADMIN_BACKUP_LIST],
        [ADMIN_BACKUP_RESTORE],
        [ADMIN_CLEAR_TEST_DATA],
        [ADMIN_BACK],
    ])


def reports_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard([
        [ADMIN_AUDIT_LOGS, ADMIN_AUDIT_EXPORT],
        [ADMIN_DEBT_HISTORY, ADMIN_EXPORT_DEBT_HISTORY],
        [ADMIN_EXPORT_USERS],
        [ADMIN_EXPORT_DEBTORS],
        [ADMIN_EXPORT_SUBMISSIONS],
        [ADMIN_BACK],
    ])


def table_reports_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard(
        [
            [ADMIN_TABLE_TEMPLATES],
            [ADMIN_TABLE_UPLOAD_F3],
            [ADMIN_TABLE_DISTRICT_PASSPORT],
            [ADMIN_BACK],
        ]
    )


def table_templates_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard(
        [
            [ADMIN_TEMPLATE_THREE_MONTHS],
            [ADMIN_TEMPLATE_DEBT_LIST],
            [ADMIN_TEMPLATE_CREDITOR_LIST],
            [ADMIN_TEMPLATE_NO_ACCOUNTING],
            [ADMIN_TABLE_BACK],
        ]
    )


def inspectors_admin_keyboard(user_id: int) -> types.ReplyKeyboardMarkup:
    if not is_super_admin_id(user_id):
        return admin_keyboard(user_id)
    return keyboard(
        [
            [ADMIN_INSPECTOR_LOCATION_COMMAND],
            [ADMIN_INSPECTOR_LOCATION_STATUS],
            [ADMIN_ADD_INSPECTOR],
            [ADMIN_LIST_INSPECTORS],
            [ADMIN_BACK],
        ]
    )


def inspector_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard(
        [
            [INSPECTOR_LOCATION_STATUS],
            [INSPECTOR_LOCATION_HELP],
            [INSPECTOR_EXIT],
        ]
    )

def broadcast_confirm_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[CONFIRM_SEND], [CANCEL]])


def clear_test_data_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[CLEAR_TEST_CONFIRM], [CANCEL]])


def debt_send_confirm_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[DEBT_SEND_CONFIRM], [CANCEL]])


def excel_confirm_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[EXCEL_REPLACE_CONFIRM], [CANCEL]])


def remove_operator_confirm_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[REMOVE_OPERATOR_CONFIRM], [CANCEL]])


def backup_restore_confirm_keyboard() -> types.ReplyKeyboardMarkup:
    return keyboard([[BACKUP_RESTORE_CONFIRM], [CANCEL]])


def session(message: types.Message) -> dict[str, Any]:
    user_id = int(message.from_user.id)
    now = time.time()
    with session_lock:
        expired = [
            telegram_id
            for telegram_id, data in sessions.items()
            if telegram_id != user_id
            and now - float(data.get("_touched_at") or now) > 3600
        ]
        for telegram_id in expired:
            expired_session = sessions.pop(telegram_id, None)
            cleanup_meter_session(expired_session)
        current = sessions.setdefault(user_id, {})
        current["_touched_at"] = now
        return current


def clear_session(message: types.Message) -> None:
    uid = message.from_user.id
    with session_lock:
        expired_session = sessions.pop(uid, None)
    cleanup_meter_session(expired_session)


def cleanup_meter_session(current: dict[str, Any] | None) -> None:
    if not current:
        return
    for photo in current.get("meter_photos") or []:
        path = photo.get("encrypted_path") if isinstance(photo, dict) else None
        if not path:
            continue
        try:
            remove_meter_photo(str(path))
        except Exception:
            continue


def _get_user_message_lock(uid: int) -> threading.RLock:
    with _user_message_locks_guard:
        lock = _user_message_locks_store.get(uid)
        if lock is None:
            lock = threading.RLock()
            _user_message_locks_store[uid] = lock
        return lock


def serialized_message_handler(callback):
    @wraps(callback)
    def wrapper(message: types.Message):
        user_id = int(message.from_user.id)
        lock = _get_user_message_lock(user_id)
        with lock:
            try:
                return callback(message)
            finally:
                with session_lock:
                    current = sessions.get(user_id)
                    if current and set(current) <= {"_touched_at"}:
                        sessions.pop(user_id, None)

    return wrapper


def ensure_runtime_dirs() -> None:
    for dirname in [
        "data",
        "data/report_templates",
        "storage",
        "outputs",
        "outputs/table_reports",
        "work",
        "work/uploads",
        "work/f3",
        "logs",
        "encrypted_vault/meter_photos",
        config.backup_dir,
    ]:
        Path(dirname).mkdir(parents=True, exist_ok=True)


def user_context(message: types.Message) -> tuple[dict[str, str], Abonent, str] | None:
    user = storage.get_user(message.from_user.id)
    if not user:
        return None
    if user.get("status") != "approved":
        return None
    abonent = book.find(user["account"])
    if not abonent:
        return None
    language = user.get("language") or "uz"
    return user, abonent, language


def notify_channel(text: str) -> bool:
    if not config.report_channel_id:
        return True
    try:
        bot.send_message(config.report_channel_id, text)
        return True
    except Exception as error:
        notify_error("Hisobot kanaliga xabar yuborilmadi", error)
        return False


def notify_channel_photo(file_id: str, caption: str) -> bool:
    if not config.report_channel_id:
        return True
    try:
        bot.send_photo(
            config.report_channel_id,
            file_id,
            caption=caption,
        )
        return True
    except Exception as error:
        notify_error("Hisobot kanaliga rasm yuborilmadi", error)
        return False


def notify_meter_submission(
    submission_id: int,
    abonent: Abonent,
    phone: str,
    language: str,
    *,
    decision: str,
) -> bool:
    if not config.report_channel_id:
        return False
    if not config.encryption_password:
        notify_error("Hisoblagich rasmi ochilmadi", "Shifrlash kaliti mavjud emas.")
        return False
    photos = storage.meter_submission_photos(submission_id)
    try:
        for index, photo in enumerate(photos, 1):
            payload = load_meter_photo(
                str(photo["encrypted_path"]),
                config.encryption_password,
                str(photo["sha256"]),
            )
            stream = io.BytesIO(payload)
            stream.name = f"meter_{submission_id}_{index}.jpg"
            bot.send_photo(
                config.report_channel_id,
                stream,
                caption=(
                    f"📸 Ariza #{submission_id}, rasm {index}/{len(photos)}\n\n"
                    + report_caption(
                        abonent,
                        phone,
                        language,
                        masked=False,
                    )
                ),
            )
        status_text = {
            "review": "⏳ Operator tekshiruvi talab etiladi.",
            "approve": "✅ Etalon bilan avtomatik mos keldi.",
            "reject": "❌ Etalon bilan aniq mos kelmadi; qayta rasm so'raldi.",
        }.get(decision, "⏳ Tekshiruv holati noma'lum.")
        channel_msg = bot.send_message(
            config.report_channel_id,
            f"🧾 Hisoblagich arizasi #{submission_id}\n"
            f"📸 Rasmlar: {len(photos)}\n"
            f"{status_text}",
            reply_markup=(
                meter_review_keyboard(submission_id)
                if decision == "review"
                else None
            ),
        )
        if decision == "review":
            storage.save_review_message(
                submission_id,
                config.report_channel_id,
                channel_msg.message_id,
                source="channel",
            )
            staff_ids = set(config.admin_ids)
            staff_ids.update(
                int(operator["telegram_id"])
                for operator in storage.operators()
            )
            for staff_id in staff_ids:
                try:
                    for index, photo in enumerate(photos, 1):
                        bot.send_photo(
                            staff_id,
                            str(photo["telegram_file_id"]),
                            caption=(
                                f"📸 Tekshiruv #{submission_id}, "
                                f"rasm {index}/{len(photos)}\n\n"
                                + report_caption(
                                    abonent,
                                    phone,
                                    language,
                                    masked=False,
                                )
                            ),
                        )
                    staff_msg = bot.send_message(
                        staff_id,
                        f"🧾 Hisoblagich arizasi #{submission_id}\n"
                        "Rasmlarni tekshirib, qarorni tanlang.",
                        reply_markup=meter_review_keyboard(submission_id),
                    )
                    storage.save_review_message(
                        submission_id,
                        staff_id,
                        staff_msg.message_id,
                        source="staff",
                    )
                except Exception as error:
                    write_error_log(
                        f"Operatorga hisoblagich tekshiruvi yuborilmadi: "
                        f"staff_id={staff_id}; error={error}"
                    )
        return True
    except Exception as error:
        notify_error("Hisoblagich arizasi kanalga yuborilmadi", error)
        return False


def notify_error(title: str, error: Exception | str) -> None:
    safe_title = redact_sensitive_text(title, 200)
    safe_error = (
        safe_exception_summary(error)
        if isinstance(error, BaseException)
        else redact_sensitive_text(error, 800)
    )
    safe_error = sanitize_audit_details(safe_error)
    text = f"🚨 {safe_title}\n\n{safe_error}"
    write_error_log(text)
    target_ids: list[int] = []
    if config.error_channel_id:
        target_ids.append(config.error_channel_id)
    target_ids.extend(config.admin_ids)
    for target_id in dict.fromkeys(target_ids):
        try:
            bot.send_message(target_id, text[:3900])
        except Exception:
            continue


def write_error_log(text: str) -> None:
    try:
        text = sanitize_audit_details(redact_sensitive_text(text, 4000))
        with error_log_lock:
            log_dir = Path("logs")
            log_dir.mkdir(parents=True, exist_ok=True)
            line = (
                f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] "
                f"{text}\n\n"
            )
            if config.encryption_enabled and config.encryption_password:
                encrypted_path = log_dir / "errors.log.uwenc"
                existing = b""
                if encrypted_path.exists():
                    try:
                        existing, _ = decrypt_bytes(
                            encrypted_path.read_bytes(),
                            config.encryption_password,
                        )
                    except Exception:
                        existing = b""
                if len(existing) + len(line.encode("utf-8")) > ERROR_LOG_MAX_BYTES:
                    for index in range(ERROR_LOG_BACKUP_COUNT, 1, -1):
                        older = log_dir / f"errors.{index - 1}.log.uwenc"
                        newer = log_dir / f"errors.{index}.log.uwenc"
                        if older.exists():
                            older.replace(newer)
                    encrypted_path.replace(log_dir / "errors.1.log.uwenc")
                    existing = b""
                payload = encrypt_bytes(
                    existing + line.encode("utf-8"),
                    config.encryption_password,
                    label="error_log",
                    original_name="errors.log",
                )
                temporary = encrypted_path.with_name(
                    f".{encrypted_path.name}.tmp"
                )
                try:
                    temporary.write_bytes(payload)
                    temporary.replace(encrypted_path)
                finally:
                    temporary.unlink(missing_ok=True)
                (log_dir / "errors.log").unlink(missing_ok=True)
                return
            plaintext_path = log_dir / "errors.log"
            if (
                plaintext_path.exists()
                and plaintext_path.stat().st_size
                + len(line.encode("utf-8"))
                > ERROR_LOG_MAX_BYTES
            ):
                for index in range(ERROR_LOG_BACKUP_COUNT, 1, -1):
                    older = log_dir / f"errors.{index - 1}.log"
                    newer = log_dir / f"errors.{index}.log"
                    if older.exists():
                        older.replace(newer)
                plaintext_path.replace(log_dir / "errors.1.log")
            with plaintext_path.open("a", encoding="utf-8") as file:
                file.write(line)
    except Exception as _log_err:
        sys.stderr.write(f"[write_error_log failed] {_log_err}: {line}\n")


def audit(actor_id: int | None, action: str, details: str = "") -> None:
    try:
        storage.add_audit_log(actor_id, action, sanitize_audit_details(details)[:2000])
    except Exception as error:
        notify_error("Audit log yozilmadi", error)


def send_message_slow(chat_id: int, text: str) -> bool:
    try:
        bot.send_message(chat_id, text)
        time.sleep(max(config.bulk_send_delay, 0))
        return True
    except ApiTelegramException as error:
        retry_after = 0
        result = getattr(error, "result_json", {}) or {}
        parameters = result.get("parameters", {}) if isinstance(result, dict) else {}
        retry_after = int(parameters.get("retry_after") or 0)
        if retry_after:
            time.sleep(retry_after + 1)
            try:
                bot.send_message(chat_id, text)
                time.sleep(max(config.bulk_send_delay, 0))
                return True
            except Exception:
                return False
        return False
    except Exception:
        return False


def backup_sources() -> list[BackupFile]:
    sources = [
        BackupFile("database", Path(config.database_file), "bot.sqlite3"),
        BackupFile("abonents", Path(config.excel_file), "abonents.xlsx"),
        BackupFile("inspectors", Path(config.inspectors_file), "inspectors.xlsx"),
    ]
    reference_archive = Path("work/meter_reference_photos.zip")
    if reference_archive.exists():
        sources.append(
            BackupFile(
                "meter_references",
                reference_archive,
                "meter_reference_photos.zip",
                required=False,
            )
        )
    return sources


def create_meter_reference_archive() -> Path | None:
    paths = [Path(path).resolve() for path in storage.meter_reference_paths()]
    if not paths:
        return None
    vault_root = Path("encrypted_vault/meter_photos").resolve()
    archive = Path("work/meter_reference_photos.zip")
    temporary = archive.with_name(f".{archive.name}.tmp")
    temporary.unlink(missing_ok=True)
    archive.unlink(missing_ok=True)
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in paths:
            if path.parent != vault_root or not path.exists():
                raise RuntimeError(f"Etalon rasm topilmadi yoki yo'li xavfsiz emas: {path}")
            bundle.write(path, arcname=path.name)
    temporary.replace(archive)
    return archive


def restore_meter_reference_archive(archive: Path) -> None:
    vault_root = Path("encrypted_vault/meter_photos").resolve()
    vault_root.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(archive, "r") as bundle:
        entries = bundle.infolist()
        if len(entries) > config.max_meter_photos * 100_000:
            raise RuntimeError("Etalon rasmlar backupida haddan tashqari ko'p fayl bor.")
        total_size = 0
        for info in entries:
            name = Path(info.filename)
            if (
                info.is_dir()
                or name.name != info.filename
                or name.suffix != ".uwenc"
            ):
                raise RuntimeError("Etalon rasmlar backupida xavfsiz bo'lmagan fayl bor.")
            if info.file_size > int(config.max_meter_photo_mb * 1024 * 1024 * 2):
                raise RuntimeError("Etalon backupidagi fayl xavfsizlik limitidan katta.")
            total_size += info.file_size
            if total_size > 20 * 1024 * 1024 * 1024:
                raise RuntimeError("Etalon rasmlar backupi ochilganda juda katta.")
            if (
                info.file_size > 1024 * 1024
                and info.file_size / max(info.compress_size, 1) > 100
            ):
                raise RuntimeError("Etalon rasmlar backupida ZIP bomb belgisi bor.")
            target = vault_root / name.name
            temporary = target.with_name(f".{target.name}.restore.tmp")
            temporary.write_bytes(bundle.read(info))
            temporary.replace(target)


def cleanup_expired_meter_photos(retention_days: int = 90) -> int:
    removed_paths = storage.purge_expired_rejected_meter_submissions(
        retention_days,
    )
    removed = 0
    for path in removed_paths:
        try:
            remove_meter_photo(path)
            removed += 1
        except Exception as error:
            write_error_log(
                f"Eski rad etilgan hisoblagich rasmi o'chirilmadi: {error}"
            )
    return removed


def backup_encryption_password() -> str | None:
    if not config.encryption_enabled:
        return None
    return config.encryption_password


def run_system_preflight() -> PreflightReport:
    report = run_preflight_checks(
        database_file=config.database_file,
        excel_file=config.excel_file,
        inspectors_file=config.inspectors_file,
        backup_dir=config.backup_dir,
        encryption_password=backup_encryption_password(),
        backup_maximum_age_hours=config.backup_maximum_age_hours,
        minimum_free_gb=config.minimum_disk_free_gb,
        minimum_free_percent=config.minimum_disk_free_percent,
    )
    write_preflight_report(report, "logs/preflight.json")
    return report


def integrity_monitor_worker() -> None:
    global last_preflight_status
    while not shutdown_event.is_set():
        try:
            report = run_system_preflight()
            signature = "|".join(
                f"{item.name}:{item.level}:{item.message}"
                for item in report.results
            )
            if signature != last_preflight_status:
                if report.status != "ok":
                    notify_error(
                        "Tizim yaxlitligi ogohlantirishi",
                        preflight_summary(report),
                    )
                elif last_preflight_status:
                    for admin_id in config.admin_ids:
                        try:
                            bot.send_message(
                                admin_id,
                                "✅ UZWATERMU tizim yaxlitligi qayta tiklandi.",
                            )
                        except Exception:
                            continue
                last_preflight_status = signature
        except Exception as error:
            notify_error("Integrity monitoring xatosi", error)
        shutdown_event.wait(
            max(config.health_check_interval_hours, 0.25) * 3600
        )


def start_integrity_monitor_worker() -> None:
    threading.Thread(
        target=integrity_monitor_worker,
        name="integrity-monitor",
        daemon=True,
    ).start()


def create_backup(actor_id: int | None = None, reason: str = "manual") -> list[Path]:
    with maintenance_lock:
        password = backup_encryption_password()
        reference_archive = create_meter_reference_archive()
        try:
            backup = create_backup_set(
                config.backup_dir,
                backup_sources(),
                reason,
                encryption_password=password,
            )
        finally:
            if reference_archive:
                reference_archive.unlink(missing_ok=True)
        removed = apply_retention(
            config.backup_dir,
            config.backup_keep_sets,
            config.backup_retention_days,
            encryption_password=password,
        )
    storage.set_setting("last_backup_at", backup.created_at.isoformat(timespec="seconds"))
    audit(
        actor_id,
        "backup_created",
        f"reason={reason}; files={len(backup.files)}; removed={len(removed)}",
    )
    return list(backup.files)


def backup_is_due(now: datetime | None = None) -> bool:
    current = now or datetime.now()
    backups = list_backup_sets(
        config.backup_dir,
        backup_encryption_password(),
    )
    if not backups:
        return True
    return current - backups[0].created_at >= timedelta(
        hours=max(config.backup_interval_hours, 0.25)
    )


def automatic_backup_worker() -> None:
    while not shutdown_event.is_set():
        try:
            if backup_is_due():
                create_backup(None, "automatic")
        except Exception as error:
            notify_error("Avtomatik backup yaratilmadi", error)
        check_seconds = min(max(config.backup_interval_hours * 900, 300), 3600)
        shutdown_event.wait(check_seconds)


def start_automatic_backup_worker() -> None:
    worker = threading.Thread(
        target=automatic_backup_worker,
        name="automatic-backup",
        daemon=True,
    )
    worker.start()


def notify_registration(abonent: Abonent, phone: str = "") -> None:
    notify_channel(
        "Ro'yxatdan o'tdi:\n"
        + report_caption(abonent, phone, masked=False)
    )


def approval_request_text(abonent: Abonent, telegram_id: int) -> str:
    return (
        "🆕 Yangi ulanish so'rovi\n\n"
        f"{report_caption(abonent, masked=False)}\n"
        f"🏠 Manzil: {abonent.address or '-'}\n"
        f"🆔 Telegram ID: {telegram_id}\n\n"
        "Tasdiqlaysizmi?"
    )


def notify_approval_request(abonent: Abonent, telegram_id: int) -> None:
    operator_ids = {int(operator["telegram_id"]) for operator in storage.operators()}
    staff_ids = set(config.admin_ids) | operator_ids
    for admin_id in staff_ids:
        try:
            bot.send_message(
                admin_id,
                approval_request_text(abonent, telegram_id),
                reply_markup=approval_keyboard(telegram_id, abonent.account),
            )
        except Exception:
            continue


def access_waiting_text(status: str, language: str = "uz") -> str:
    text = TEXTS.get(language, TEXTS["uz"])
    if status == "pending":
        return text["access_pending"]
    if status == "rejected":
        return text["access_rejected"]
    return WELCOME


def handle_existing_consumer_start(message: types.Message) -> bool:
    user = storage.get_user(message.from_user.id)
    if not user:
        return False

    language = user.get("language") or "uz"
    if user.get("status") == "approved":
        bot.send_message(
            message.chat.id,
            TEXTS[language]["main_menu"],
            reply_markup=main_menu(language),
        )
    else:
        bot.send_message(
            message.chat.id,
            access_waiting_text(user.get("status", ""), language),
            reply_markup=start_keyboard(),
        )
    return True


def is_super_admin_id(telegram_id: int) -> bool:
    return telegram_id in config.admin_ids


def is_operator_id(telegram_id: int) -> bool:
    return storage.get_operator(telegram_id) is not None


def is_staff_id(telegram_id: int) -> bool:
    return is_super_admin_id(telegram_id) or is_operator_id(telegram_id)


def is_admin(message: types.Message) -> bool:
    return message.from_user.id in config.admin_ids


def is_staff(message: types.Message) -> bool:
    return is_staff_id(message.from_user.id)


def flood_warning_text(user_id: int, retry_after: int) -> str:
    user = storage.get_user(user_id)
    language = (user or {}).get("language") or "uz"
    if language == "ru":
        return (
            "⚠️ Слишком много запросов. "
            f"Повторите через {retry_after} секунд."
        )
    if language == "uz_cyrl":
        return (
            "⚠️ Жуда кўп сўров юборилди. "
            f"{retry_after} сониядан кейин қайта урининг."
        )
    return (
        "⚠️ Juda ko'p so'rov yuborildi. "
        f"{retry_after} soniyadan keyin qayta urinib ko'ring."
    )


def request_allowed(user_id: int, chat_id: int | None = None) -> bool:
    staff = is_staff_id(user_id)
    limit = (
        config.staff_request_limit
        if staff
        else config.user_request_limit
    )
    decision = request_limiter.check(user_id, limit)
    if decision.allowed:
        return True
    if decision.newly_blocked:
        audit(
            user_id,
            "request_flood_blocked",
            f"staff={staff}; retry_after={decision.retry_after}",
        )
        if chat_id is not None:
            try:
                bot.send_message(
                    chat_id,
                    flood_warning_text(user_id, decision.retry_after),
                )
            except Exception:
                logging.debug("Flood ogohlantirish yuborishda xato (chat_id=%s)", chat_id, exc_info=True)
    return False


def is_private_chat(message: types.Message) -> bool:
    return getattr(message.chat, "type", "") == "private"


def private_chat_required_text(user_id: int) -> str:
    user = storage.get_user(user_id)
    language = (user or {}).get("language") or "uz"
    return TEXTS.get(language, TEXTS["uz"])["private_chat_required"]


def activate_admin_session(message: types.Message) -> None:
    current = session(message)
    current["admin_session_until"] = session_deadline()
    storage.clear_admin_access_attempts(message.from_user.id)


def clear_admin_flow(message: types.Message) -> None:
    deadline = session(message).get("admin_session_until")
    with session_lock:
        sessions[message.from_user.id] = {"_touched_at": time.time()}
        if deadline and float(deadline) > time.time():
            sessions[message.from_user.id]["admin_session_until"] = deadline


def admin_session_active(message: types.Message) -> bool:
    deadline = float(session(message).get("admin_session_until") or 0)
    return deadline > time.time()


def require_admin_session(message: types.Message) -> bool:
    if not is_private_chat(message):
        bot.send_message(
            message.chat.id,
            "⛔ Admin panel faqat bot bilan shaxsiy chatda ishlaydi.",
        )
        return False
    if admin_session_active(message):
        return True
    clear_session(message)
    bot.send_message(
        message.chat.id,
        "🔒 Admin sessiyasi tugagan. Panelga qayta kirish uchun /admin yuboring.",
    )
    return False


def register_unauthorized_admin_attempt(telegram_id: int) -> dict[str, Any]:
    previous = storage.admin_access_attempt(telegram_id)
    if admin_access_blocked(previous):
        return {**previous, "blocked": True, "just_blocked": False}

    result = register_denied_attempt(previous)
    storage.save_admin_access_attempt(
        telegram_id,
        int(result["failed_count"]),
        str(result["window_started"]),
        str(result["blocked_until"]),
        str(result["last_attempt"]),
    )
    audit(
        telegram_id,
        "admin_access_denied",
        f"telegram_id={telegram_id}; failures={result['failed_count']}",
    )
    if result["just_blocked"]:
        warning = (
            "🚨 Admin panelga takroriy ruxsatsiz kirish urinishi\n\n"
            f"Telegram ID: {telegram_id}\n"
            f"Vaqtinchalik blok: {ADMIN_BLOCK_MINUTES} daqiqa"
        )
        for admin_id in config.admin_ids:
            try:
                bot.send_message(admin_id, warning)
            except Exception:
                continue
    return {**result, "blocked": bool(result["blocked_until"])}


def admin_denied_text(message: types.Message) -> str:
    return (
        "⛔ Admin panelga kirish uchun ruxsat yo'q.\n\n"
        "Agar siz tashkilot xodimi bo'lsangiz, super adminga murojaat qiling."
    )


def admin_stats_text() -> str:
    stats = storage.stats()
    return (
        "📊 Statistika\n\n"
        f"👥 Ro'yxatdan o'tganlar: {stats['users']}\n"
        f"⏳ Tasdiqlash kutayotganlar: {stats['pending']}\n"
        f"👨‍💼 Operatorlar: {stats['operators']}\n"
        f"🧾 Ko'rsatgich yuborishlar: {stats['submissions']}\n"
        f"📸 Qabul qilingan rasmlar: {stats['photos']}\n"
        f"📚 Excel abonentlar soni: {book.count()}"
    )


def parse_money(value: str) -> float:
    cleaned = str(value or "").replace(" ", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def format_money(value: str) -> str:
    amount = parse_money(value)
    formatted = f"{amount:,.2f}"
    return (
        formatted.replace(",", "\u00a0")
        .replace(".", ",")
        .rstrip("0")
        .rstrip(",")
    )


def parse_db_datetime(value: str) -> datetime | None:
    if not value:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def excel_statement_date() -> str:
    return book.loaded_at or datetime.now().strftime("%d.%m.%Y")


def debt_calendar_due(now: datetime) -> bool:
    return now.day in {10, 25}


def next_debt_calendar_date(now: datetime | None = None) -> datetime:
    current = now or datetime.now()
    for day in (10, 25):
        if day > current.day:
            return current.replace(
                day=day,
                hour=9,
                minute=0,
                second=0,
                microsecond=0,
            )
    year = current.year + (1 if current.month == 12 else 0)
    month = 1 if current.month == 12 else current.month + 1
    return current.replace(
        year=year,
        month=month,
        day=10,
        hour=9,
        minute=0,
        second=0,
        microsecond=0,
    )


def debt_warning_sent_today(
    user: dict[str, str],
    account: str,
    now: datetime,
) -> bool:
    last_warning = storage.last_debt_warning(
        int(user["telegram_id"]),
        account,
    )
    if not last_warning:
        return False
    sent_at = parse_db_datetime(last_warning.get("sent_at", ""))
    return sent_at is not None and sent_at.date() == now.date()


def _run_debt_analysis_core(
    send_messages: bool = False,
    force: bool = False,
    calendar_only: bool = False,
    run_id: str = "",
) -> dict[str, Any]:
    now = datetime.now()
    statement_date = excel_statement_date()
    result = {
        "checked": 0,
        "debtors": 0,
        "sent": 0,
        "skipped": 0,
        "failed": 0,
        "calendar_wait": 0,
        "busy": 0,
        "statement_date": statement_date,
    }
    if send_messages and calendar_only and not debt_calendar_due(now):
        result["calendar_wait"] = 1
        return result

    if send_messages:
        if not run_id:
            run_id = datetime.now().strftime("%Y%m%dT%H%M%S%f")
        storage.set_setting("debt_job_status", "running")
        storage.set_setting("debt_job_run_id", run_id)
        storage.set_setting(
            "debt_job_started_at",
            now.isoformat(timespec="seconds"),
        )
        storage.set_setting("debt_job_progress", "0")

    for user in storage.all_approved_accounts():
        result["checked"] += 1
        if send_messages and result["checked"] % 100 == 0:
            storage.set_setting(
                "debt_job_progress",
                str(result["checked"]),
            )
        abonent = book.find(user["account"])
        if not abonent:
            result["skipped"] += 1
            continue
        if parse_money(abonent.debit) <= 0:
            continue

        result["debtors"] += 1
        language = user.get("language") or "uz"
        if (
            send_messages
            and storage.debt_warning_succeeded_in_run(
                int(user["telegram_id"]),
                abonent.account,
                run_id,
            )
        ):
            result["skipped"] += 1
            continue
        if (
            send_messages
            and not force
            and debt_warning_sent_today(user, abonent.account, now)
        ):
            result["skipped"] += 1
            storage.save_debt_warning_log(
                int(user["telegram_id"]),
                abonent.account,
                abonent.debit,
                language,
                statement_date,
                "skipped_today",
                run_id=run_id,
            )
            continue
        if not send_messages:
            continue

        try:
            delivered = send_message_slow(
                int(user["telegram_id"]),
                debt_warning_text(
                    language,
                    format_money(abonent.debit),
                    statement_date,
                    mask_account(abonent.account),
                ),
            )
            if not delivered:
                result["failed"] += 1
                storage.save_debt_warning_log(
                    int(user["telegram_id"]),
                    abonent.account,
                    abonent.debit,
                    language,
                    statement_date,
                    "failed",
                    "send_message_slow returned False",
                    run_id,
                )
                continue
            storage.save_debt_warning(int(user["telegram_id"]), abonent.account, abonent.debit)
            storage.save_debt_warning_log(
                int(user["telegram_id"]),
                abonent.account,
                abonent.debit,
                language,
                statement_date,
                "sent",
                run_id=run_id,
            )
            result["sent"] += 1
        except Exception as error:
            storage.save_debt_warning_log(
                int(user["telegram_id"]),
                abonent.account,
                abonent.debit,
                language,
                statement_date,
                "failed",
                str(error),
                run_id,
            )
            result["failed"] += 1

    if send_messages:
        storage.set_setting("debt_job_status", "completed")
        storage.set_setting(
            "debt_job_progress",
            str(result["checked"]),
        )
        storage.set_setting(
            "debt_job_completed_at",
            datetime.now().isoformat(timespec="seconds"),
        )
        storage.set_setting(
            "debt_job_result",
            json.dumps(result, ensure_ascii=False),
        )
    return result


def run_debt_analysis(
    send_messages: bool = False,
    force: bool = False,
    calendar_only: bool = False,
) -> dict[str, Any]:
    if not send_messages:
        return _run_debt_analysis_core(
            send_messages=False,
            force=force,
            calendar_only=calendar_only,
        )
    if not debt_job_lock.acquire(blocking=False):
        return {
            "checked": 0,
            "debtors": 0,
            "sent": 0,
            "skipped": 0,
            "failed": 0,
            "calendar_wait": 0,
            "busy": 1,
            "statement_date": excel_statement_date(),
        }
    try:
        return _run_debt_analysis_core(
            send_messages=True,
            force=force,
            calendar_only=calendar_only,
        )
    finally:
        debt_job_lock.release()


def debt_analysis_summary(
    result: dict[str, Any],
    send_messages: bool = False,
) -> str:
    if result.get("busy"):
        return (
            "⏳ Qarzdorlik ogohlantirishlarini yuborish jarayoni "
            "allaqachon davom etmoqda. Ikkinchi jarayon ishga tushirilmadi."
        )
    if result.get("calendar_wait"):
        next_run = next_debt_calendar_date().strftime("%d.%m.%Y")
        return (
            "🔴 Qarzdorlik tahlili\n\n"
            "Bugun avtomatik ogohlantirish kuni emas.\n"
            "Avtomatik kunlar: har oyning 10 va 25-sanasi.\n"
            f"Keyingi avtomatik sana: {next_run}."
        )

    note = "📨 Ogohlantirish yuborish rejimi" if send_messages else "ℹ️ Bu faqat tahlil. Ogohlantirish yuborilmadi."
    return (
        "🔴 Qarzdorlik tahlili yakunlandi\n\n"
        f"📅 Excel sanasi: {result.get('statement_date', '-')}\n"
        f"🔢 Hisob raqamlari tekshirildi: {result['checked']}\n"
        f"🔴 Qarzdor hisoblar: {result['debtors']}\n"
        f"✅ Ogohlantirish yuborildi: {result['sent']}\n"
        f"⏭ O'tkazib yuborildi: {result['skipped']}\n"
        f"❌ Xatolik: {result['failed']}\n\n"
        f"{note}"
    )


def debt_job_worker(
    chat_id: int | None,
    actor_id: int | None,
    reason: str,
    force: bool,
    calendar_only: bool,
    run_id: str,
) -> None:
    try:
        result = _run_debt_analysis_core(
            send_messages=True,
            force=force,
            calendar_only=calendar_only,
            run_id=run_id,
        )
    except Exception as error:
        storage.set_setting("debt_job_status", "failed")
        storage.set_setting(
            "debt_job_completed_at",
            datetime.now().isoformat(timespec="seconds"),
        )
        notify_error("Qarzdorlik ogohlantirishlarini yuborishda xatolik", error)
        return
    finally:
        debt_job_lock.release()

    audit(actor_id, "debt_warning_job_completed", f"{reason}; {result}")
    summary = debt_analysis_summary(result, send_messages=True)
    if chat_id is not None:
        try:
            bot.send_message(chat_id, summary)
        except Exception as error:
            notify_error("Qarzdorlik natijasi adminga yuborilmadi", error)
    elif result.get("sent") or result.get("failed"):
        for admin_id in config.admin_ids:
            try:
                bot.send_message(
                    admin_id,
                    "🤖 Avtomatik qarzdorlik ogohlantirishlari\n\n"
                    + summary,
                )
            except Exception:
                continue


def start_debt_warning_job(
    *,
    chat_id: int | None,
    actor_id: int | None,
    reason: str,
    force: bool = False,
    calendar_only: bool = False,
    run_id: str = "",
) -> bool:
    if not debt_job_lock.acquire(blocking=False):
        return False
    resolved_run_id = (
        run_id or datetime.now().strftime("%Y%m%dT%H%M%S%f")
    )
    threading.Thread(
        target=debt_job_worker,
        args=(
            chat_id,
            actor_id,
            reason,
            force,
            calendar_only,
            resolved_run_id,
        ),
        name=f"debt-warning-{reason}",
        daemon=True,
    ).start()
    return True


def debt_analysis_worker() -> None:
    first_check = True
    while not shutdown_event.is_set():
        try:
            interrupted = (
                first_check
                and storage.get_setting("debt_job_status", "") == "running"
            )
            interrupted_run_id = (
                storage.get_setting("debt_job_run_id", "")
                if interrupted
                else ""
            )
            start_debt_warning_job(
                chat_id=None,
                actor_id=None,
                reason=(
                    "recovered_after_restart"
                    if interrupted
                    else "automatic_calendar"
                ),
                force=False,
                calendar_only=not interrupted,
                run_id=interrupted_run_id,
            )
        except Exception as error:
            notify_error("Avtomatik qarzdorlik tahlilida xatolik", error)
        first_check = False
        now = datetime.now()
        next_check = (now + timedelta(days=1)).replace(
            hour=0,
            minute=5,
            second=0,
            microsecond=0,
        )
        shutdown_event.wait(max((next_check - now).total_seconds(), 60))


def start_debt_analysis_worker() -> None:
    thread = threading.Thread(
        target=debt_analysis_worker,
        name="debt-analysis",
        daemon=True,
    )
    thread.start()


def contact_center_text(language: str) -> str:
    if language == "ru":
        title = "☎️ Контактный центр"
        phones_label = "Телефоны"
        address_label = "Адрес"
    elif language == "uz_cyrl":
        title = "☎️ Алоқа маркази"
        phones_label = "Телефонлар"
        address_label = "Манзил"
    else:
        title = "☎️ Aloqa markazi"
        phones_label = "Telefonlar"
        address_label = "Manzil"

    return (
        f"{title}\n\n"
        f"📞 {phones_label}: {config.contact_phones}\n"
        f"📍 {address_label}: {config.district_address}"
    )


def work_hours_text(language: str) -> str:
    return TEXTS.get(language, TEXTS["uz"])["work_hours_content"]


def faq_text(language: str) -> str:
    return TEXTS.get(language, TEXTS["uz"])["faq_content"]


def handle_public_service(message: types.Message, language: str = "uz") -> bool:
    text = TEXTS.get(language, TEXTS["uz"])
    message_text = message.text or ""
    if message_text == text["service"] or (language != "uz" and message_text == TEXTS["uz"]["service"]):
        sess = session(message)
        sess["state"] = states.PUBLIC_SERVICE
        bot.send_message(message.chat.id, text["service"], reply_markup=public_service_keyboard(language))
        return True
    if message_text == text["contact_center"]:
        bot.send_message(message.chat.id, contact_center_text(language), reply_markup=public_service_keyboard(language))
        if config.district_lat is not None and config.district_lon is not None:
            bot.send_location(message.chat.id, config.district_lat, config.district_lon)
        return True
    if message_text == text["work_hours"]:
        bot.send_message(message.chat.id, work_hours_text(language), reply_markup=public_service_keyboard(language))
        return True
    if message_text == text["faq"]:
        bot.send_message(message.chat.id, faq_text(language), reply_markup=public_service_keyboard(language))
        return True
    if message_text == text["back"]:
        clear_session(message)
        bot.send_message(message.chat.id, WELCOME, reply_markup=start_keyboard())
        return True
    return False


def recent_submissions_text() -> str:
    submissions = storage.recent_submissions(10)
    if not submissions:
        return "🧾 Oxirgi arizalar\n\nHozircha arizalar yo'q."

    lines = ["🧾 Oxirgi arizalar"]
    for index, item in enumerate(submissions, 1):
        abonent = book.find(str(item["account"]))
        if abonent:
            fio = abonent.fio or "-"
            phone = abonent.phone or "-"
            area = abonent.area or "-"
            inspector = abonent.inspector_fio or "Vakant"
        else:
            fio = "-"
            phone = "-"
            area = "-"
            inspector = "Excel faylda topilmadi"

        lines.append(
            "\n"
            f"{index}. Ariza ID: {item['id']}\n"
            f"🔢 Hisob raqami: {item['account']}\n"
            f"👨‍💼 F.I.O: {fio}\n"
            f"📞 Telefon: {phone}\n"
            f"🧭 Hudud: {area}\n"
            f"👮 Inspektor: {inspector}\n"
            f"📸 Rasmlar soni: {item['photo_count']}\n"
            f"🌐 Til: {item.get('language') or '-'}\n"
            f"📌 Status: {item.get('status') or '-'}\n"
            f"🔍 O'xshashlik: {item.get('similarity_score') if item.get('similarity_score') is not None else '-'}\n"
            f"👤 Qaror bergan: {item.get('decided_by') or '-'}\n"
            f"📅 Sana: {item['created_at']}"
        )
    return "\n".join(lines)


def audit_logs_text() -> str:
    logs = storage.recent_audit_logs(15)
    if not logs:
        return "📜 Audit loglar\n\nHozircha log yo'q."

    lines = ["📜 Oxirgi audit loglar"]
    for log in logs:
        details = sanitize_audit_details(log.get("details"))
        lines.append(
            "\n"
            f"#{log['id']} | {log['created_at']}\n"
            f"👤 Actor: {log.get('actor_id') or '-'}\n"
            f"🔧 Amal: {log['action']}\n"
            f"📝 Izoh: {details}"
        )
    return "\n".join(lines)


def backup_result_text(paths: list[Path]) -> str:
    if not paths:
        return "💾 Backup yaratilmadi: zaxira qilinadigan fayl topilmadi."
    lines = ["💾 Backup yaratildi", ""]
    for path in paths:
        lines.append(f"✅ {path}")
    return "\n".join(lines)


def backup_list_text(limit: int = 20) -> str:
    backup_dir = Path(config.backup_dir)
    if not backup_dir.exists():
        return "📂 Backup ro'yxati\n\nHozircha backup papkasi yo'q."

    backups = list_backup_sets(
        backup_dir,
        backup_encryption_password(),
    )
    if not backups:
        return "📂 Backup ro'yxati\n\nHozircha yangi formatdagi backup yo'q."

    lines = ["📂 Oxirgi tekshirilgan backup to'plamlari"]
    for backup in backups[:limit]:
        size_kb = sum(path.stat().st_size for path in backup.files) / 1024
        created = backup.created_at.strftime("%d.%m.%Y %H:%M")
        lines.append(
            f"\n✅ {created}\n"
            f"📝 Sabab: {backup.reason or '-'}\n"
            f"📦 Fayllar: {len(backup.files)}\n"
            f"💽 Hajm: {size_kb:.1f} KB"
        )
    return "\n".join(lines)


def system_status_text() -> str:
    def exists_icon(path: Path) -> str:
        return "✅" if path.exists() else "❌"

    db_path = Path(config.database_file)
    abonents_path = Path(config.excel_file)
    inspectors_path = Path(config.inspectors_file)
    backup_dir = Path(config.backup_dir)
    last_abonents = storage.get_setting("last_abonents.xlsx_upload", "-")
    last_inspectors = storage.get_setting("last_inspectors.xlsx_upload", "-")
    latest_backup = latest_valid_backup(
        backup_dir,
        backup_encryption_password(),
    )
    latest_backup_text = (
        latest_backup.created_at.strftime("%d.%m.%Y %H:%M")
        if latest_backup
        else "sog'lom backup topilmadi"
    )
    debt_job_status = storage.get_setting("debt_job_status", "-")
    debt_job_progress = storage.get_setting("debt_job_progress", "0")
    debt_job_completed = storage.get_setting("debt_job_completed_at", "-")
    health_path = Path("logs/health.json")
    health_state = "topilmadi"
    health_heartbeat = "-"
    if health_path.exists():
        try:
            health = json.loads(health_path.read_text(encoding="utf-8"))
            health_state = str(health.get("state") or "noma'lum")
            health_heartbeat = str(health.get("heartbeat_at") or "-")
        except Exception:
            health_state = "o'qilmadi"
    preflight_path = Path("logs/preflight.json")
    preflight_state = "topilmadi"
    preflight_checked = "-"
    if preflight_path.exists():
        try:
            preflight = json.loads(
                preflight_path.read_text(encoding="utf-8")
            )
            preflight_state = str(
                preflight.get("status") or "noma'lum"
            )
            preflight_checked = str(
                preflight.get("checked_at") or "-"
            )
        except Exception:
            preflight_state = "o'qilmadi"
    return (
        "🩺 Tizim holati\n\n"
        f"{exists_icon(abonents_path)} Abonentlar Excel: {abonents_path}\n"
        f"{exists_icon(inspectors_path)} Inspektorlar Excel: {inspectors_path}\n"
        f"{exists_icon(db_path)} Baza: {db_path}\n"
        f"{exists_icon(backup_dir)} Backup papka: {backup_dir}\n"
        f"💾 Oxirgi sog'lom backup: {latest_backup_text}\n"
        f"⏱ Backup oralig'i: {config.backup_interval_hours:g} soat\n"
        f"🗓 Saqlash: {config.backup_retention_days} kun / "
        f"{config.backup_keep_sets} to'plam\n"
        f"🔐 Backup shifrlash: {'yoqilgan' if config.encryption_enabled else 'o‘chirilgan'}\n"
        f"🫀 Bot holati: {health_state}\n"
        f"🕒 Oxirgi heartbeat: {health_heartbeat}\n"
        f"🛡 Preflight: {preflight_state}\n"
        f"🔎 Oxirgi integrity tekshiruv: {preflight_checked}\n"
        f"🛡 Faol flood bloklar: {request_limiter.active_blocks()}\n"
        f"🚫 Jami flood bloklar: {request_limiter.total_blocks}\n"
        f"📎 Excel limiti: {config.max_excel_upload_mb:g} MB\n"
        f"📷 Rasm limiti: {config.max_meter_photos} ta / "
        f"{config.max_meter_photo_mb:g} MB\n"
        f"📚 Xotiradagi abonentlar: {book.count()}\n"
        f"📅 Excel o'qilgan sana: {book.loaded_at or '-'}\n"
        f"📁 Oxirgi abonents.xlsx upload: {last_abonents}\n"
        f"👮 Oxirgi inspectors.xlsx upload: {last_inspectors}\n"
        f"📨 Bulk delay: {config.bulk_send_delay} soniya\n"
        f"🔴 Qarzdorlik jarayoni: {debt_job_status}\n"
        f"🔢 Qarzdorlik progressi: {debt_job_progress} hisob\n"
        f"🕒 Qarzdorlik yakunlangan vaqt: {debt_job_completed}\n"
        f"🚨 Error channel: {config.error_channel_id or '-'}"
    )


def user_stats_text() -> str:
    status_counts = storage.user_status_counts()
    language_counts = storage.user_language_counts()
    approved = storage.all_approved_accounts()
    connected_debtors = 0
    total_debit = 0.0
    for user in approved:
        abonent = book.find(user["account"])
        if not abonent:
            continue
        debit = parse_money(abonent.debit)
        if debit > 0:
            connected_debtors += 1
            total_debit += debit

    status_lines = "\n".join(f"{status}: {count}" for status, count in sorted(status_counts.items())) or "-"
    language_lines = "\n".join(f"{language}: {count}" for language, count in sorted(language_counts.items())) or "-"
    return (
        "👥 Foydalanuvchi statistikasi\n\n"
        "📌 Statuslar:\n"
        f"{status_lines}\n\n"
        "🌐 Tillar:\n"
        f"{language_lines}\n\n"
        f"🔴 Ulangan qarzdorlar: {connected_debtors}\n"
        f"💰 Ulangan qarzdorlik jami: {total_debit:,.2f}"
    )


def export_audit_csv(actor_id: int) -> Path:
    export_dir = Path("outputs")
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"audit_logs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    logs = [masked_row(row) for row in storage.recent_audit_logs(1000)]
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=["id", "actor_id", "action", "details", "created_at"])
        writer.writeheader()
        writer.writerows(csv_safe_rows(logs))
    audit(actor_id, "audit_exported", str(path))
    return path


def csv_safe_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def csv_safe_rows(
    rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    return [
        {
            key: csv_safe_value(value)
            for key, value in row.items()
        }
        for row in rows
    ]


def write_csv(name: str, rows: list[dict[str, Any]], fieldnames: list[str]) -> Path:
    export_dir = Path("outputs")
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"{name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(csv_safe_rows(rows))
    return path


def secure_export_after_send(path: Path, label: str = "export") -> None:
    if not config.encryption_enabled or not config.encryption_password:
        return
    encrypted_path = path.with_name(f"{path.name}.uwenc")
    encrypt_file(
        path,
        encrypted_path,
        config.encryption_password,
        label=label,
    )
    path.unlink(missing_ok=True)


def cleanup_pending_excel(current: dict[str, Any]) -> None:
    pending = current.get("pending_excel") or {}
    path_text = str(pending.get("path") or "")
    if path_text:
        Path(path_text).unlink(missing_ok=True)


def export_users_csv(actor_id: int) -> Path:
    rows = [masked_row(row) for row in storage.all_user_records()]
    audit(actor_id, "users_exported", f"rows={len(rows)}")
    return write_csv("users", rows, ["telegram_id", "account", "language", "phone", "status", "created_at", "updated_at"])


def export_debtors_csv(actor_id: int) -> Path:
    rows: list[dict[str, Any]] = []
    for user in storage.all_approved_accounts():
        abonent = book.find(user["account"])
        if not abonent:
            continue
        debit = parse_money(abonent.debit)
        if debit <= 0:
            continue
        rows.append({
            "telegram_id": mask_telegram_id(user["telegram_id"]),
            "account": mask_account(abonent.account),
            "fio": masked_row({"fio": abonent.fio})["fio"],
            "phone": masked_row({"phone": user.get("phone") or abonent.phone})["phone"],
            "language": user.get("language") or "uz",
            "debit": abonent.debit,
            "area": abonent.area,
            "inspector_fio": abonent.inspector_fio,
        })
    audit(actor_id, "debtors_exported", f"rows={len(rows)}")
    return write_csv("connected_debtors", rows, ["telegram_id", "account", "fio", "phone", "language", "debit", "area", "inspector_fio"])


def export_submissions_csv(actor_id: int) -> Path:
    rows = [masked_row(row) for row in storage.recent_submissions(10000)]
    audit(actor_id, "submissions_exported", f"rows={len(rows)}")
    return write_csv(
        "submissions",
        rows,
        [
            "id", "telegram_id", "account", "photo_count", "language",
            "inspector_fio", "status", "similarity_score",
            "decision_reason", "decided_by", "decided_at", "created_at",
        ],
    )


def export_debt_history_csv(actor_id: int) -> Path:
    rows = [masked_row(row) for row in storage.all_debt_warning_logs(10000)]
    audit(actor_id, "debt_history_exported", f"rows={len(rows)}")
    return write_csv("debt_warning_history", rows, ["id", "telegram_id", "account", "debit", "language", "statement_date", "status", "error", "created_at"])


def debt_history_text() -> str:
    logs = storage.recent_debt_warning_logs(15)
    if not logs:
        return "📨 Ogohlantirish tarixi\n\nHozircha tarix yo'q."
    lines = ["📨 Oxirgi qarzdorlik ogohlantirishlari"]
    for log in logs:
        lines.append(
            "\n"
            f"#{log['id']} | {log['created_at']}\n"
            f"🔢 Hisob raqami: {mask_account(log['account'])}\n"
            f"🆔 Telegram ID: {mask_telegram_id(log['telegram_id'])}\n"
            f"💰 Summa: {log['debit']}\n"
            f"🌐 Til: {log.get('language') or '-'}\n"
            f"📌 Status: {log['status']}"
        )
    return "\n".join(lines)


def settings_text() -> str:
    return (
        "⚙️ Sozlamalar\n\n"
        f"REPORT_CHANNEL_ID: {mask_telegram_id(config.report_channel_id)}\n"
        f"ERROR_CHANNEL_ID: {mask_telegram_id(config.error_channel_id)}\n"
        f"EXCEL_FILE: {config.excel_file}\n"
        f"INSPECTORS_FILE: {config.inspectors_file}\n"
        f"DATABASE_FILE: {config.database_file}\n"
        f"BACKUP_DIR: {config.backup_dir}\n"
        f"BACKUP_INTERVAL_HOURS: {config.backup_interval_hours:g}\n"
        f"BACKUP_RETENTION_DAYS: {config.backup_retention_days}\n"
        f"BACKUP_KEEP_SETS: {config.backup_keep_sets}\n"
        f"ENCRYPTION_ENABLED: {config.encryption_enabled}\n"
        f"BULK_SEND_DELAY: {config.bulk_send_delay}\n"
        f"CONTACT_PHONES: {config.contact_phones}\n"
        f"DISTRICT_ADDRESS: {config.district_address}"
    )


def inspector_users_text() -> str:
    inspectors = storage.inspector_users()
    if not inspectors:
        return "👷 Ulangan inspektorlar\n\nHozircha inspektor ulanmagan."
    lines = ["👷 Ulangan inspektorlar"]
    for index, item in enumerate(inspectors, 1):
        lines.append(
            "\n"
            f"{index}. Telegram ID: {item['telegram_id']}\n"
            f"🧭 Hudud: {item.get('area') or '-'}\n"
            f"👤 F.I.O: {item.get('fio') or '-'}\n"
            f"☎️ Telefon: {item.get('phone') or '-'}\n"
            f"📌 Status: {item.get('status') or '-'}"
        )
    return "\n".join(lines)


def inspector_profile_text(inspector: dict[str, Any]) -> str:
    inspector_id = int(inspector["telegram_id"])
    fio = html.escape(str(inspector.get("fio") or "F.I.O kiritilmagan"))
    area = html.escape(str(inspector.get("area") or "-"))
    phone = html.escape(str(inspector.get("phone") or "-"))
    return (
        f'<a href="tg://user?id={inspector_id}">{fio}</a>\n'
        f"🆔 Telegram ID: <code>{inspector_id}</code>\n"
        f"🧭 Hudud: {area}\n"
        f"☎️ Telefon: {phone}"
    )


def inspector_panel_text(inspector: dict[str, Any]) -> str:
    return (
        "👷 Inspektor paneli\n\n"
        f"👤 F.I.O: {inspector.get('fio') or '-'}\n"
        f"🧭 Hudud: {inspector.get('area') or '-'}\n"
        f"☎️ Telefon: {inspector.get('phone') or '-'}\n\n"
        "Rahbariyatdan live lokatsiya buyrug'i kelganda Telegramdagi "
        "📎 biriktirish menyusi orqali Location → Share My Live Location "
        "ni tanlang. Oddiy lokatsiya qabul qilinmaydi."
    )


def inspector_location_status_text(inspector_id: int) -> str:
    active = storage.active_inspector_location_session()
    if not active:
        return "📍 Hozir faol live lokatsiya nazorati mavjud emas."
    responses = storage.inspector_location_responses(int(active["id"]))
    response = next(
        (
            item
            for item in responses
            if int(item["inspector_id"]) == inspector_id
        ),
        None,
    )
    if not response:
        return "📍 Siz ushbu nazorat sessiyasi ro'yxatiga kiritilmagansiz."
    status_labels = {
        "pending": "⏳ Live lokatsiya kutilmoqda",
        "reminded": "⚠️ Ogohlantirish yuborilgan, live lokatsiya kutilmoqda",
        "live": "✅ Live lokatsiya yoqilgan",
        "delivery_failed": "❌ Buyruqni yetkazishda xatolik qayd etilgan",
    }
    return (
        "📍 Live lokatsiya holati\n\n"
        f"{status_labels.get(str(response['status']), response['status'])}\n"
        f"🕒 Buyruq vaqti: {active['requested_at']}\n"
        f"⏰ Yakuniy muddat: {active['final_at']}"
    )


def inspector_location_session_summary(session_id: int) -> str:
    responses = storage.inspector_location_responses(session_id)
    live = sum(item["status"] == "live" for item in responses)
    pending = len(responses) - live
    return (
        "📡 Live lokatsiya nazorati\n\n"
        f"👷 Jami inspektorlar: {len(responses)}\n"
        f"✅ Live lokatsiya yoqilgan: {live}\n"
        f"⏳ Bajarilmagan: {pending}"
    )


def inspector_noncompliance_report_chunks(
    session: dict[str, Any],
    responses: list[dict[str, Any]],
) -> list[str]:
    missing = [item for item in responses if item["status"] != "live"]
    header_lines = [
        "🚨 LIVE LOKATSIYA BUYRUG'I BAJARILMAGAN",
        "",
        f"🕒 Buyruq vaqti: {html.escape(str(session['requested_at']))}",
        f"⏰ Yakuniy muddat: {html.escape(str(session['final_at']))}",
        f"👷 Bajarmaganlar: {len(missing)} ta",
    ]
    if not missing:
        return [
            "\n".join(
                header_lines
                + ["", "✅ Barcha inspektorlar buyruqni bajardi."]
            )
        ]
    chunks: list[str] = []
    current = "\n".join(header_lines)
    for index, item in enumerate(missing, 1):
        block = "\n".join(
            (
                f"{index}. {inspector_profile_text(item)}",
                (
                    "📌 Holat: buyruq yetkazilmadi"
                    if item["status"] == "delivery_failed"
                    else "📌 Holat: live lokatsiya yoqilmadi"
                ),
            )
        )
        candidate = f"{current}\n\n{block}"
        if len(candidate) > 3800:
            chunks.append(current)
            current = (
                "🚨 BUYRUQNI BAJARMAGANLAR DAVOMI\n\n"
                + block
            )
        else:
            current = candidate
    chunks.append(current)
    return chunks


def send_inspector_location_command(
    requester_id: int,
    chat_id: int,
) -> None:
    active = storage.active_inspector_location_session()
    if active:
        bot.send_message(
            chat_id,
            "⏳ Oldingi live lokatsiya nazorati hali yakunlanmagan.\n\n"
            + inspector_location_session_summary(int(active["id"])),
            reply_markup=inspectors_admin_keyboard(requester_id),
        )
        return
    inspectors = storage.active_inspector_users()
    if not inspectors:
        bot.send_message(
            chat_id,
            "❌ Faol inspektorlar ro'yxati bo'sh. Avval inspektorlarni "
            "Telegram ID orqali ulang.",
            reply_markup=inspectors_admin_keyboard(requester_id),
        )
        return

    now = datetime.now()
    reminder_at = now + timedelta(
        minutes=INSPECTOR_LOCATION_RESPONSE_MINUTES
    )
    final_at = reminder_at + timedelta(
        minutes=INSPECTOR_LOCATION_FINAL_GRACE_MINUTES
    )
    session_id, session_inspectors = (
        storage.create_inspector_location_session(
            requester_id,
            now.isoformat(timespec="seconds"),
            reminder_at.isoformat(timespec="seconds"),
            final_at.isoformat(timespec="seconds"),
        )
    )
    delivered = 0
    failed = 0
    command_text = (
        "📡 RAHBARIYAT BUYRUG'I\n\n"
        "Live lokatsiyangizni hozir yoqing va kamida 1 soat davomida "
        "faol qoldiring.\n\n"
        "Telegramda: 📎 → Location → Share My Live Location → 1 hour.\n"
        "Oddiy statik lokatsiya bajarilgan deb hisoblanmaydi.\n\n"
        f"⏰ Birinchi muddat: {reminder_at.strftime('%H:%M')}\n"
        f"🚨 Yakuniy muddat: {final_at.strftime('%H:%M')}"
    )
    for inspector in session_inspectors:
        inspector_id = int(inspector["telegram_id"])
        try:
            bot.send_message(
                inspector_id,
                command_text,
                reply_markup=inspector_keyboard(),
            )
            time.sleep(max(config.bulk_send_delay, 0))
            storage.mark_inspector_command_delivery(
                session_id,
                inspector_id,
                delivered=True,
                sent_at=datetime.now().isoformat(timespec="seconds"),
            )
            delivered += 1
        except Exception as error:
            storage.mark_inspector_command_delivery(
                session_id,
                inspector_id,
                delivered=False,
                sent_at=datetime.now().isoformat(timespec="seconds"),
                error=type(error).__name__,
            )
            failed += 1
    audit(
        requester_id,
        "inspector_location_command_started",
        f"session={session_id}; delivered={delivered}; failed={failed}",
    )
    bot.send_message(
        chat_id,
        "✅ Live lokatsiya buyrug'i yuborildi.\n\n"
        f"👷 Inspektorlar: {len(session_inspectors)}\n"
        f"✅ Yetkazildi: {delivered}\n"
        f"❌ Yetkazilmadi: {failed}\n"
        f"⚠️ Ogohlantirish: {reminder_at.strftime('%d.%m.%Y %H:%M')}\n"
        f"🚨 Yakuniy hisobot: {final_at.strftime('%d.%m.%Y %H:%M')}",
        reply_markup=inspectors_admin_keyboard(requester_id),
    )


def relay_inspector_live_location(
    session_id: int,
    inspector: dict[str, Any],
    latitude: float,
    longitude: float,
    live_period: int,
) -> None:
    inspector_id = int(inspector["telegram_id"])
    relays = storage.inspector_location_relays(
        session_id,
        inspector_id,
    )
    if relays:
        for relay in relays:
            try:
                bot.edit_message_live_location(
                    latitude=latitude,
                    longitude=longitude,
                    chat_id=int(relay["admin_id"]),
                    message_id=int(relay["message_id"]),
                )
            except Exception:
                continue
        return

    relay_period = min(max(live_period, 60), 86400)
    intro_text = "✅ Inspektor live lokatsiyani yoqdi\n\n" + inspector_profile_text(inspector)
    destinations: list[int | str] = list(config.admin_ids)
    if config.location_channel_id:
        destinations.append(config.location_channel_id)
    for dest in destinations:
        try:
            sent = bot.send_location(
                dest,
                latitude,
                longitude,
                live_period=relay_period,
            )
            storage.save_inspector_location_relay(
                session_id,
                inspector_id,
                int(sent.chat.id),
                int(sent.message_id),
            )
            bot.send_message(
                dest,
                intro_text,
                parse_mode="HTML",
                disable_web_page_preview=True,
            )
        except Exception as error:
            notify_error(
                "Inspektor live lokatsiyasi uzatilmadi",
                f"inspector_id={inspector_id}; dest={dest}; {type(error).__name__}",
            )


def handle_inspector_location(
    message: types.Message,
    inspector: dict[str, Any],
    *,
    edited: bool = False,
) -> None:
    active = storage.active_inspector_location_session()
    if not active and edited:
        active = storage.inspector_location_session_for_update(
            int(inspector["telegram_id"]),
            int(message.message_id),
        )
    if not active:
        if not edited:
            bot.send_message(
                message.chat.id,
                "📍 Hozir faol live lokatsiya buyrug'i mavjud emas.",
                reply_markup=inspector_keyboard(),
            )
        return
    location = message.location
    live_period = int(getattr(location, "live_period", 0) or 0)
    if live_period < INSPECTOR_MINIMUM_LIVE_PERIOD_SECONDS:
        if not edited:
            bot.send_message(
                message.chat.id,
                "❌ Oddiy yoki 1 soatdan qisqa lokatsiya qabul qilinmadi.\n\n"
                "📎 → Location → Share My Live Location → 1 hour "
                "tartibida yuboring.",
                reply_markup=inspector_keyboard(),
            )
            if config.location_channel_id:
                try:
                    bot.send_message(
                        config.location_channel_id,
                        "⚠️ Inspektor live lokatsiya yoqmadi\n\n"
                        + inspector_profile_text(inspector)
                        + f"\n\nYuborilgan: oddiy lokatsiya (live_period={live_period}s)",
                        parse_mode="HTML",
                        disable_web_page_preview=True,
                    )
                except Exception:
                    logging.debug("Inspektor lokatsiya kanal xabari yuborishda xato", exc_info=True)
        return

    received_at = datetime.now().isoformat(timespec="seconds")
    saved = storage.save_inspector_live_location(
        int(active["id"]),
        int(inspector["telegram_id"]),
        latitude=float(location.latitude),
        longitude=float(location.longitude),
        live_period=live_period,
        source_message_id=int(message.message_id),
        received_at=received_at,
    )
    if not saved:
        if not edited:
            bot.send_message(
                message.chat.id,
                "❌ Siz ushbu nazorat sessiyasi ro'yxatida emassiz.",
                reply_markup=inspector_keyboard(),
            )
        return
    relay_inspector_live_location(
        int(active["id"]),
        inspector,
        float(location.latitude),
        float(location.longitude),
        live_period,
    )
    if not edited:
        audit(
            int(inspector["telegram_id"]),
            "inspector_live_location_started",
            f"session={active['id']}; area={inspector.get('area') or '-'}",
        )
        bot.send_message(
            message.chat.id,
            "✅ Live lokatsiyangiz qabul qilindi va rahbariyatga uzatildi.\n"
            "Translyatsiyani belgilangan vaqt davomida o'chirmang.",
            reply_markup=inspector_keyboard(),
        )


def handle_inspector_message(
    message: types.Message,
    inspector: dict[str, Any],
) -> None:
    if getattr(message, "location", None):
        handle_inspector_location(message, inspector)
        return
    message_text = message.text or ""
    if message_text == INSPECTOR_LOCATION_STATUS:
        bot.send_message(
            message.chat.id,
            inspector_location_status_text(message.from_user.id),
            reply_markup=inspector_keyboard(),
        )
        return
    if message_text == INSPECTOR_LOCATION_HELP:
        bot.send_message(
            message.chat.id,
            "ℹ️ Live lokatsiyani yuborish\n\n"
            "1. Telegramdagi 📎 belgisini bosing.\n"
            "2. Location bo'limini tanlang.\n"
            "3. Share My Live Location ni bosing.\n"
            "4. 1 hour yoki undan uzoq muddatni tanlang.\n\n"
            "Send My Current Location oddiy lokatsiya bo'lib, nazorat "
            "talabini bajarmaydi.",
            reply_markup=inspector_keyboard(),
        )
        return
    if message_text == INSPECTOR_EXIT:
        clear_session(message)
        bot.send_message(
            message.chat.id,
            "Inspektor paneli yopildi. Qayta kirish uchun Start bosing.",
            reply_markup=start_keyboard(),
        )
        return
    bot.send_message(
        message.chat.id,
        inspector_panel_text(inspector),
        reply_markup=inspector_keyboard(),
    )


def process_inspector_location_deadlines(
    now: datetime | None = None,
) -> None:
    current = now or datetime.now()
    active = storage.active_inspector_location_session()
    if not active:
        return
    session_id = int(active["id"])
    reminder_at = parse_db_datetime(str(active["reminder_at"]))
    final_at = parse_db_datetime(str(active["final_at"]))
    responses = storage.inspector_location_responses(session_id)

    if final_at and current >= final_at:
        if not storage.complete_inspector_location_session(
            session_id,
            current.isoformat(timespec="seconds"),
        ):
            return
        reports = inspector_noncompliance_report_chunks(active, responses)
        recipients = set(config.admin_ids)
        recipients.add(int(active["requested_by"]))
        for admin_id in recipients:
            for report in reports:
                try:
                    bot.send_message(
                        admin_id,
                        report,
                        parse_mode="HTML",
                        disable_web_page_preview=True,
                    )
                except Exception as error:
                    notify_error(
                        "Inspektorlar yakuniy hisoboti yuborilmadi",
                        f"admin_id={admin_id}; {type(error).__name__}",
                    )
        audit(
            None,
            "inspector_location_session_completed",
            f"session={session_id}; missing="
            f"{sum(item['status'] != 'live' for item in responses)}",
        )
        return

    if (
        reminder_at
        and current >= reminder_at
        and active["status"] == "active"
    ):
        for response in responses:
            if response["status"] == "live":
                continue
            inspector_id = int(response["inspector_id"])
            try:
                bot.send_message(
                    inspector_id,
                    "⚠️ OGOHLANTIRISH\n\n"
                    "Rahbariyatning live lokatsiya yoqish buyrug'i hali "
                    "bajarilmadi. Darhol kamida 1 soatlik live lokatsiyani "
                    "yoqing.\n\n"
                    f"🚨 Yakuniy muddat: "
                    f"{final_at.strftime('%d.%m.%Y %H:%M') if final_at else '-'}",
                    reply_markup=inspector_keyboard(),
                )
                time.sleep(max(config.bulk_send_delay, 0))
            except Exception:
                logging.exception("inspector reminder failed inspector_id=%s session=%s", inspector_id, session_id)
            storage.mark_inspector_location_reminded(
                session_id,
                inspector_id,
                current.isoformat(timespec="seconds"),
            )
        storage.set_inspector_location_session_status(
            session_id,
            "reminded",
        )
        audit(
            None,
            "inspector_location_reminders_sent",
            f"session={session_id}",
        )


def inspector_location_monitor_worker() -> None:
    while not shutdown_event.is_set():
        try:
            process_inspector_location_deadlines()
        except Exception as error:
            notify_error("Inspektor live lokatsiya nazorati xatosi", error)
        shutdown_event.wait(15)


def start_inspector_location_monitor_worker() -> None:
    threading.Thread(
        target=inspector_location_monitor_worker,
        name="inspector-location-monitor",
        daemon=True,
    ).start()


def money_invalid(value: str) -> bool:
    if not str(value or "").strip():
        return False
    cleaned = str(value).replace(" ", "").replace(",", ".").replace("-", "")
    try:
        float(cleaned)
        return False
    except ValueError:
        return True


def phone_invalid(value: str) -> bool:
    if not str(value or "").strip():
        return False
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    return len(digits) < 7 or len(digits) > 15


def excel_quality_checks(path: Path, temp_book: AbonentBook) -> dict[str, Any]:
    validate_safe_xlsx(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    mapping = temp_book._detect_mapping(rows)
    start = temp_book._data_start_index(rows, mapping)
    seen: dict[str, int] = {}
    duplicates: set[str] = set()
    blank_fio = 0
    bad_phone = 0
    bad_money = 0
    no_area = 0

    for row in rows[start:]:
        account = temp_book.clean_account(temp_book._cell(row, mapping, "account"))
        fio = temp_book._cell(row, mapping, "fio")
        if not account and not fio:
            continue
        if account:
            seen[account] = seen.get(account, 0) + 1
            if seen[account] > 1:
                duplicates.add(account)
        if account and not fio:
            blank_fio += 1
        phone = temp_book._cell(row, mapping, "new_phone") or temp_book._cell(row, mapping, "phone")
        if phone_invalid(phone):
            bad_phone += 1
        if money_invalid(temp_book._cell(row, mapping, "debit")) or money_invalid(temp_book._cell(row, mapping, "credit")):
            bad_money += 1
        if account and not temp_book._cell(row, mapping, "area"):
            no_area += 1

    inspector_missing = sum(1 for abonent in temp_book.all_items() if not abonent.inspector_fio)
    return {
        "duplicates": len(duplicates),
        "duplicate_samples": ", ".join(list(sorted(duplicates))[:5]),
        "blank_fio": blank_fio,
        "bad_phone": bad_phone,
        "bad_money": bad_money,
        "no_area": no_area,
        "inspector_missing": inspector_missing,
    }


def validate_abonents_excel(path: Path) -> tuple[dict[str, int], str]:
    temp_book = AbonentBook(str(path), config.inspectors_file)
    items = temp_book.all_items()
    if not items:
        raise RuntimeError(
            "Abonentlar Excel faylida iste'molchi topilmadi"
        )
    quality = excel_quality_checks(path, temp_book)
    debtors = sum(1 for abonent in items if parse_money(abonent.debit) > 0)
    no_phone = sum(1 for abonent in items if not abonent.phone)
    no_meter = sum(1 for abonent in items if abonent.has_meter.strip() == "-")
    vacant = sum(1 for abonent in items if not abonent.inspector_fio)
    stats = {
        "abonents": len(items),
        "debtors": debtors,
        "no_phone": no_phone,
        "no_meter": no_meter,
        "vacant": vacant,
        **quality,
    }
    text = (
        "📁 Abonentlar Excel tekshirildi\n\n"
        f"📚 Abonentlar: {stats['abonents']}\n"
        f"🔴 Qarzdorlar: {stats['debtors']}\n"
        f"☎️ Telefonsiz: {stats['no_phone']}\n"
        f"🚱 Hisoblagichsiz (-): {stats['no_meter']}\n"
        f"👷 Inspektor vakant/topilmadi: {stats['vacant']}\n\n"
        "🔎 Sifat nazorati:\n"
        f"🔁 Dublikat hisob raqamlar: {stats['duplicates']}\n"
        f"👤 Bo'sh F.I.O qatorlar: {stats['blank_fio']}\n"
        f"☎️ Noto'g'ri telefon formatlari: {stats['bad_phone']}\n"
        f"💰 Noto'g'ri summa formatlari: {stats['bad_money']}\n"
        f"🧭 Hududsiz qatorlar: {stats['no_area']}\n"
        f"👷 Inspektor mos kelmaganlar: {stats['inspector_missing']}\n"
        f"Namuna dublikatlar: {stats['duplicate_samples'] or '-'}\n\n"
        "Excelni almashtirishni tasdiqlaysizmi?"
    )
    return stats, text


def validate_inspectors_excel(path: Path) -> tuple[dict[str, int], str]:
    validate_safe_xlsx(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    total = 0
    with_phone = 0
    for row in rows[1:]:
        values = [str(value or "").strip() for value in row]
        if not any(values):
            continue
        total += 1
        if any(ch.isdigit() for ch in " ".join(values)):
            with_phone += 1
    if total == 0:
        raise RuntimeError("Inspektorlar Excel faylida ma'lumot topilmadi")
    AbonentBook(config.excel_file, str(path))
    stats = {"inspectors": total, "with_phone": with_phone}
    text = (
        "👮 Inspektorlar Excel tekshirildi\n\n"
        f"👷 Qatorlar: {stats['inspectors']}\n"
        f"☎️ Telefon ko'rsatilgan qatorlar: {stats['with_phone']}\n\n"
        "Excelni almashtirishni tasdiqlaysizmi?"
    )
    return stats, text


def send_pending_users(chat_id: int, requester_id: int) -> None:
    total = storage.pending_account_count()
    users = storage.pending_users(20)
    if not users:
        bot.send_message(chat_id, "⏳ Tasdiqlash so'rovlari\n\nHozircha kutayotgan so'rovlar yo'q.", reply_markup=requests_admin_keyboard(requester_id))
        return

    shown_text = (
        f"⏳ Tasdiqlash so'rovlari: jami {total} ta\n"
        f"Eng eski {len(users)} ta so'rov ko'rsatilmoqda."
    )
    bot.send_message(
        chat_id,
        shown_text,
        reply_markup=requests_admin_keyboard(requester_id),
    )
    for user in users:
        abonent = book.find(user["account"])
        if abonent:
            text = approval_request_text(abonent, int(user["telegram_id"]))
        else:
            text = (
                "🆕 Yangi ulanish so'rovi\n\n"
                f"🔢 Hisob raqami: {user['account']}\n"
                "❌ Excel faylda topilmadi\n"
                f"🆔 Telegram ID: {user['telegram_id']}\n\n"
                "Tasdiqlaysizmi?"
            )
        bot.send_message(
            chat_id,
            text,
            reply_markup=approval_keyboard(
                int(user["telegram_id"]),
                str(user["account"]),
            ),
        )


@serialized_message_handler
def handle_start(message: types.Message) -> None:
    if not request_allowed(message.from_user.id, message.chat.id):
        return
    if not is_private_chat(message):
        bot.send_message(
            message.chat.id,
            private_chat_required_text(message.from_user.id),
        )
        return
    clear_session(message)
    if is_super_admin_id(message.from_user.id):
        if not is_private_chat(message):
            bot.send_message(message.chat.id, "⛔ Admin panel faqat bot bilan shaxsiy chatda ishlaydi.")
            return
        activate_admin_session(message)
        bot.send_message(message.chat.id, "🛠 Super admin panel\n\nKerakli amalni tanlang:", reply_markup=admin_keyboard(message.from_user.id))
        return
    inspector = storage.get_inspector_user(message.from_user.id)
    if inspector and inspector.get("status") == "active":
        bot.send_message(
            message.chat.id,
            inspector_panel_text(inspector),
            reply_markup=inspector_keyboard(),
        )
        return
    operator = storage.get_operator(message.from_user.id)
    if operator:
        if not is_private_chat(message):
            bot.send_message(message.chat.id, "⛔ Operator panel faqat bot bilan shaxsiy chatda ishlaydi.")
            return
        if operator.get("language"):
            activate_admin_session(message)
            bot.send_message(message.chat.id, "🛠 Operator panel\n\nKerakli amalni tanlang:", reply_markup=admin_keyboard(message.from_user.id))
            return
        current = session(message)
        current["state"] = states.CHOOSING_OPERATOR_LANGUAGE
        bot.send_message(message.chat.id, TEXTS["uz"]["choose_language"], reply_markup=language_keyboard())
        return
    if handle_existing_consumer_start(message):
        return
    bot.send_message(message.chat.id, WELCOME, reply_markup=start_keyboard())


@serialized_message_handler
def handle_myid(message: types.Message) -> None:
    if not request_allowed(message.from_user.id, message.chat.id):
        return
    if not is_private_chat(message):
        bot.send_message(
            message.chat.id,
            private_chat_required_text(message.from_user.id),
        )
        return
    bot.send_message(message.chat.id, f"Sizning Telegram ID: `{message.from_user.id}`", parse_mode="Markdown")


@serialized_message_handler
def handle_admin(message: types.Message) -> None:
    if not request_allowed(message.from_user.id, message.chat.id):
        return
    if not is_private_chat(message):
        audit(message.from_user.id, "admin_group_access_denied", "non_private_chat")
        bot.send_message(message.chat.id, "⛔ Admin panel faqat bot bilan shaxsiy chatda ishlaydi.")
        return
    if not is_staff(message):
        attempt = register_unauthorized_admin_attempt(message.from_user.id)
        if attempt.get("blocked"):
            bot.send_message(
                message.chat.id,
                f"⛔ Juda ko'p urinish. Admin panelga kirish {ADMIN_BLOCK_MINUTES} daqiqaga bloklandi.",
            )
            return
        bot.send_message(message.chat.id, admin_denied_text(message), parse_mode="Markdown")
        return
    clear_session(message)
    activate_admin_session(message)
    operator = storage.get_operator(message.from_user.id)
    if operator and not operator.get("language") and not is_super_admin_id(message.from_user.id):
        current = session(message)
        current["state"] = states.CHOOSING_OPERATOR_LANGUAGE
        bot.send_message(message.chat.id, TEXTS["uz"]["choose_language"], reply_markup=language_keyboard())
        return
    title = "🛠 Super admin panel" if is_super_admin_id(message.from_user.id) else "🛠 Operator panel"
    bot.send_message(message.chat.id, f"{title}\n\nKerakli amalni tanlang:", reply_markup=admin_keyboard(message.from_user.id))


@serialized_message_handler
def handle_approval_callback(call: types.CallbackQuery) -> None:
    if not request_allowed(call.from_user.id):
        bot.answer_callback_query(
            call.id,
            "Juda ko'p so'rov. Birozdan keyin qayta urining.",
            show_alert=True,
        )
        return
    if not is_staff_id(call.from_user.id):
        register_unauthorized_admin_attempt(call.from_user.id)
        bot.answer_callback_query(call.id, "Sizda ruxsat yo'q.")
        return
    if getattr(call.message.chat, "type", "") != "private":
        bot.answer_callback_query(call.id, "Tasdiqlash faqat shaxsiy chatda ishlaydi.")
        return

    if not call.data or ":" not in call.data:
        bot.answer_callback_query(call.id, "Noto'g'ri so'rov.")
        return
    callback_parts = call.data.split(":")
    if len(callback_parts) not in {2, 3}:
        bot.answer_callback_query(call.id, "Noto'g'ri so'rov.")
        return
    action = callback_parts[0]
    try:
        telegram_id = int(callback_parts[1])
    except (TypeError, ValueError):
        bot.answer_callback_query(call.id, "Noto'g'ri so'rov.")
        return
    user = storage.get_user(telegram_id)
    if not user:
        bot.answer_callback_query(call.id, "Foydalanuvchi topilmadi.")
        return

    account = callback_parts[2] if len(callback_parts) == 3 else user["account"]
    if not account:
        bot.answer_callback_query(call.id, "Noto'g'ri so'rov.")
        return
    linked_account = storage.get_linked_account(telegram_id, account)
    if not linked_account:
        bot.answer_callback_query(call.id, "Hisob raqam bog'lanishi topilmadi.")
        return

    abonent = book.find(account)
    language = user.get("language") or "uz"

    if action not in {"approve", "reject"}:
        bot.answer_callback_query(call.id, "Noto'g'ri amal.")
        return
    if action == "approve" and not abonent:
        bot.answer_callback_query(
            call.id,
            "Hisob raqami joriy Excel faylda topilmadi. Excelni tekshiring.",
            show_alert=True,
        )
        audit(
            call.from_user.id,
            "stale_approval_blocked",
            f"telegram_id={telegram_id}; account={account}",
        )
        return

    target_status = "approved" if action == "approve" else "rejected"
    if not storage.transition_account_status(
        telegram_id,
        account,
        "pending",
        target_status,
        f"operator:{call.from_user.id}",
    ):
        latest_account = storage.get_linked_account(telegram_id, account)
        latest_status = latest_account.get("status") if latest_account else "unknown"
        bot.answer_callback_query(
            call.id,
            f"Bu ariza avval ko'rib chiqilgan: {latest_status}.",
            show_alert=True,
        )
        return

    if action == "approve":
        if user.get("status") != "approved":
            storage.set_active_account(telegram_id, account)
        audit(call.from_user.id, "user_approved", f"telegram_id={telegram_id}; account={account}")
        if abonent:
            notify_registration(abonent, user.get("phone", ""))
            bot.send_message(
                telegram_id,
                TEXTS[language]["operator_approved"].format(
                    account=mask_account(account),
                ),
                reply_markup=main_menu(language),
            )
        bot.answer_callback_query(call.id, "Tasdiqlandi.")
        result = "✅ Tasdiqlandi"
    else:
        if user.get("status") != "approved":
            storage.update_status(telegram_id, "rejected")
        audit(call.from_user.id, "user_rejected", f"telegram_id={telegram_id}; account={account}")
        bot.send_message(
            telegram_id,
            TEXTS[language]["operator_rejected"].format(
                account=mask_account(account),
            ),
            reply_markup=(
                main_menu(language)
                if user.get("status") == "approved"
                else start_keyboard()
            ),
        )
        bot.answer_callback_query(call.id, "Rad etildi.")
        result = "❌ Rad etildi"

    try:
        bot.edit_message_text(
            f"{call.message.text}\n\nNatija: {result}\nOperator ID: {call.from_user.id}",
            chat_id=call.message.chat.id,
            message_id=call.message.message_id,
        )
    except Exception:
        logging.debug("Tasdiqlash xabari tahrirlashda xato (message_id=%s)", call.message.message_id, exc_info=True)


@serialized_message_handler
def handle_meter_review_callback(call: types.CallbackQuery) -> None:
    if not request_allowed(call.from_user.id):
        bot.answer_callback_query(call.id, "Birozdan keyin qayta urining.", show_alert=True)
        return
    if not is_staff_id(call.from_user.id):
        register_unauthorized_admin_attempt(call.from_user.id)
        bot.answer_callback_query(call.id, "Sizda ruxsat yo'q.", show_alert=True)
        return
    try:
        action, raw_submission_id = (call.data or "").split(":", 1)
        submission_id = int(raw_submission_id)
    except (TypeError, ValueError):
        bot.answer_callback_query(call.id, "Noto'g'ri so'rov.", show_alert=True)
        return
    if action not in {"meterapprove", "meterreject"}:
        bot.answer_callback_query(call.id, "Noto'g'ri amal.", show_alert=True)
        return

    submission = storage.meter_submission(submission_id)
    if not submission:
        bot.answer_callback_query(call.id, "Ariza topilmadi.", show_alert=True)
        return
    approved = action == "meterapprove"
    target_status = "approved" if approved else "rejected"
    if not storage.transition_meter_submission(
        submission_id,
        "pending_review",
        target_status,
        call.from_user.id,
        "operator_approved" if approved else "operator_rejected",
        create_references=approved,
    ):
        latest = storage.meter_submission(submission_id)
        bot.answer_callback_query(
            call.id,
            f"Ariza avval ko'rib chiqilgan: {(latest or {}).get('status', '-')}.",
            show_alert=True,
        )
        return

    telegram_id = int(submission["telegram_id"])
    language = str(submission.get("language") or "uz")
    text = TEXTS.get(language, TEXTS["uz"])
    if approved:
        next_date = (datetime.now() + timedelta(days=30)).strftime("%d.%m.%Y")
        user_text = text["meter_operator_approved"].format(date=next_date)
        answer = "Tasdiqlandi va etalon sifatida saqlandi."
    else:
        user_text = text["meter_operator_rejected"]
        answer = "Rad etildi. Iste'molchidan qayta rasm so'raldi."
    try:
        bot.send_message(
            telegram_id,
            user_text,
            reply_markup=main_menu(language),
        )
    except Exception as error:
        notify_error("Hisoblagich qarori iste'molchiga yuborilmadi", error)
    from_channel = (
        call.message.chat.id == config.report_channel_id
        if config.report_channel_id
        else False
    )
    source_label = "kanal" if from_channel else "bot"
    audit(
        call.from_user.id,
        f"meter_submission_{target_status}",
        (
            f"submission_id={submission_id}; account={submission['account']}; "
            f"source={source_label}; chat_id={call.message.chat.id}"
        ),
    )
    result_text = "✅ Tasdiqlandi" if approved else "❌ Rad etildi"
    operator_id = call.from_user.id
    bot.answer_callback_query(call.id, answer)

    review_messages = storage.get_review_messages(submission_id)
    for msg in review_messages:
        try:
            bot.edit_message_text(
                f"🧾 Hisoblagich arizasi #{submission_id}\n\n"
                f"Natija: {result_text}\n"
                f"Operator: {operator_id}\n"
                f"Manba: {source_label}",
                chat_id=int(msg["chat_id"]),
                message_id=int(msg["message_id"]),
                reply_markup=None,
            )
        except Exception:
            logging.debug("Ko'rib chiqish xabari tahrirlashda xato (msg_id=%s)", msg["message_id"], exc_info=True)
    storage.delete_review_messages(submission_id)

    if not review_messages:
        try:
            original = call.message.text or f"Hisoblagich arizasi #{submission_id}"
            bot.edit_message_text(
                f"{original}\n\nNatija: {result_text}\nOperator: {operator_id}",
                chat_id=call.message.chat.id,
                message_id=call.message.message_id,
                reply_markup=None,
            )
        except Exception:
            logging.debug("Callback xabari tahrirlashda xato (message_id=%s)", call.message.message_id, exc_info=True)


@serialized_message_handler
def handle_message(message: types.Message) -> None:
    if not request_allowed(message.from_user.id, message.chat.id):
        return
    if not is_private_chat(message):
        bot.send_message(
            message.chat.id,
            private_chat_required_text(message.from_user.id),
        )
        return
    current = session(message)
    state = current.get("state")
    message_text = message.text or ""

    if message_text.strip().casefold() in {"/start", "start", "▶️ start"}:
        cleanup_meter_session(current)
        current.clear()
        if is_super_admin_id(message.from_user.id):
            if not is_private_chat(message):
                bot.send_message(message.chat.id, "⛔ Admin panel faqat bot bilan shaxsiy chatda ishlaydi.")
                return
            activate_admin_session(message)
            bot.send_message(message.chat.id, "🛠 Super admin panel\n\nKerakli amalni tanlang:", reply_markup=admin_keyboard(message.from_user.id))
            return
        inspector = storage.get_inspector_user(message.from_user.id)
        if inspector and inspector.get("status") == "active":
            bot.send_message(
                message.chat.id,
                inspector_panel_text(inspector),
                reply_markup=inspector_keyboard(),
            )
            return
        operator = storage.get_operator(message.from_user.id)
        if operator:
            if not is_private_chat(message):
                bot.send_message(message.chat.id, "⛔ Operator panel faqat bot bilan shaxsiy chatda ishlaydi.")
                return
            if operator.get("language"):
                activate_admin_session(message)
                bot.send_message(message.chat.id, "🛠 Operator panel\n\nKerakli amalni tanlang:", reply_markup=admin_keyboard(message.from_user.id))
                return
            current["state"] = states.CHOOSING_OPERATOR_LANGUAGE
            bot.send_message(message.chat.id, TEXTS["uz"]["choose_language"], reply_markup=language_keyboard())
            return
        if handle_existing_consumer_start(message):
            return
        current["state"] = states.CHOOSING_LANGUAGE
        bot.send_message(message.chat.id, TEXTS["uz"]["choose_language"], reply_markup=language_keyboard())
        return

    inspector = storage.get_inspector_user(message.from_user.id)
    if inspector and inspector.get("status") == "active":
        handle_inspector_message(message, inspector)
        return

    if state == states.PUBLIC_SERVICE:
        _lang = current.get("language") or "uz"
        if handle_public_service(message, _lang):
            return
        bot.send_message(
            message.chat.id,
            "⚠️ Quyidagi xizmat tugmalaridan birini tanlang.",
            reply_markup=public_service_keyboard(_lang),
        )
        return

    if state == states.CHOOSING_OPERATOR_LANGUAGE:
        choose_operator_language(message, current)
        return

    if state == states.CHOOSING_LANGUAGE:
        choose_language(message, current)
        return

    if state == states.ENTERING_ACCOUNT:
        enter_account(message, current)
        return

    if state == states.VERIFYING_ACCOUNT_PHONE:
        verify_account_phone(message, current)
        return

    if state == states.CHOOSING_ACCOUNT:
        choose_linked_account(message, current)
        return

    if state in states.ADMIN_STATES:
        if not require_admin_session(message):
            return
        handle_admin_state(message, current)
        return

    if is_staff(message):
        if not require_admin_session(message):
            return
        if handle_admin_menu(message, current):
            return
        bot.send_message(
            message.chat.id,
            "⚠️ Admin paneldagi tugmalardan birini tanlang.",
            reply_markup=admin_keyboard(message.from_user.id),
        )
        return

    saved_user = storage.get_user(message.from_user.id)
    if saved_user and saved_user.get("status") != "approved":
        saved_language = saved_user.get("language") or "uz"
        bot.send_message(
            message.chat.id,
            access_waiting_text(
                saved_user.get("status", ""),
                saved_language,
            ),
            reply_markup=start_keyboard(),
        )
        return

    context = user_context(message)
    if not context:
        _lang = current.get("language") or "uz"
        if handle_public_service(message, _lang):
            return
        bot.send_message(message.chat.id, WELCOME, reply_markup=start_keyboard())
        return

    user, abonent, language = context
    text = TEXTS[language]

    if message_text == text["back"]:
        clear_session(message)
        bot.send_message(message.chat.id, text["main_menu"], reply_markup=main_menu(language))
        return

    if state == states.EDITING_PHONE:
        save_phone(message, abonent, language)
        return

    if state == states.WAITING_METER_PHOTO:
        handle_meter_photo(message, current, user, abonent, language)
        return

    try:
        if message_text == text["cabinet"]:
            bot.send_message(
                message.chat.id,
                cabinet_text(abonent, language, user.get("phone", "")),
                reply_markup=cabinet_keyboard(language),
            )
        elif message_text == text["edit_phone"]:
            current["state"] = states.EDITING_PHONE
            bot.send_message(message.chat.id, text["enter_phone"], reply_markup=phone_keyboard(language))
        elif message_text == text["send_meter"]:
            current["state"] = states.WAITING_METER_PHOTO
            current["photo_count"] = 0
            current["meter_photos"] = []
            bot.send_message(message.chat.id, meter_intro_text(abonent, language), reply_markup=meter_keyboard(language))
        elif message_text == text["balance"]:
            bot.send_message(message.chat.id, balance_text(abonent, language, book.loaded_at), reply_markup=main_menu(language))
        elif message_text == text["my_accounts"]:
            show_linked_accounts(message, language)
        elif message_text == text["service"]:
            bot.send_message(message.chat.id, text["service"], reply_markup=service_keyboard(language))
        elif message_text == text["change_language"]:
            current["state"] = states.CHOOSING_LANGUAGE
            current["change_only"] = True
            bot.send_message(message.chat.id, text["choose_language"], reply_markup=language_keyboard())
        elif message_text == text["guide"]:
            bot.send_message(message.chat.id, GUIDE[language], reply_markup=service_keyboard(language))
        elif message_text == text["contact_center"]:
            bot.send_message(message.chat.id, contact_center_text(language), reply_markup=service_keyboard(language))
            if config.district_lat is not None and config.district_lon is not None:
                bot.send_location(message.chat.id, config.district_lat, config.district_lon)
        elif message_text == text["work_hours"]:
            bot.send_message(message.chat.id, work_hours_text(language), reply_markup=service_keyboard(language))
        elif message_text == text["faq"]:
            bot.send_message(message.chat.id, faq_text(language), reply_markup=service_keyboard(language))
        else:
            bot.send_message(message.chat.id, text["main_menu"], reply_markup=main_menu(language))
    except Exception:
        logging.exception("handle_message menu send failed user_id=%s", message.from_user.id)


def handle_edited_inspector_location(message: types.Message) -> None:
    if getattr(message.chat, "type", "") != "private":
        return
    now = time.monotonic()
    uid = message.from_user.id
    if now - _location_last_update.get(uid, 0.0) < _LOCATION_THROTTLE_SECONDS:
        return
    _location_last_update[uid] = now
    inspector = storage.get_inspector_user(uid)
    if (
        not inspector
        or inspector.get("status") != "active"
        or not message.location
    ):
        return
    handle_inspector_location(message, inspector, edited=True)


def choose_language(message: types.Message, current: dict[str, Any]) -> None:
    language = LANGUAGE_BY_BUTTON.get(message.text or "")
    if not language:
        bot.send_message(message.chat.id, TEXTS["uz"]["choose_language"], reply_markup=language_keyboard())
        return

    user = storage.get_user(message.from_user.id)
    if current.get("change_only") and user:
        storage.save_user(message.from_user.id, user["account"], language, user.get("phone", ""))
        current.clear()
        bot.send_message(message.chat.id, TEXTS[language]["main_menu"], reply_markup=main_menu(language))
        return

    current["language"] = language
    current["state"] = states.ENTERING_ACCOUNT
    bot.send_message(message.chat.id, TEXTS[language]["enter_account"])


def choose_operator_language(message: types.Message, current: dict[str, Any]) -> None:
    language = LANGUAGE_BY_BUTTON.get(message.text or "")
    if not language:
        bot.send_message(message.chat.id, TEXTS["uz"]["choose_language"], reply_markup=language_keyboard())
        return

    storage.save_operator_language(message.from_user.id, language)
    current.clear()
    activate_admin_session(message)
    bot.send_message(message.chat.id, "🛠 Operator panel\n\nKerakli amalni tanlang:", reply_markup=admin_keyboard(message.from_user.id))


def normalized_phone(value: str) -> str:
    digits = "".join(character for character in str(value or "") if character.isdigit())
    return digits[-9:] if len(digits) >= 9 else digits


def normalized_uzbek_phone(value: str) -> str | None:
    digits = "".join(
        character
        for character in str(value or "")
        if character.isdigit()
    )
    if len(digits) == 9:
        return f"+998{digits}"
    if len(digits) == 12 and digits.startswith("998"):
        return f"+{digits}"
    return None


def account_phone_matches(abonent: Abonent, phone: str) -> bool:
    candidate = normalized_phone(phone)
    if len(candidate) != 9:
        return False
    excel_phones = {
        normalized_phone(abonent.phone),
        normalized_phone(abonent.original_phone),
        normalized_phone(abonent.edited_phone),
    }
    excel_phones.discard("")
    return candidate in excel_phones


def begin_account_verification(
    message: types.Message,
    current: dict[str, Any],
    abonent: Abonent,
    language: str,
) -> None:
    current["state"] = states.VERIFYING_ACCOUNT_PHONE
    current["pending_account"] = abonent.account
    current["language"] = language
    bot.send_message(
        message.chat.id,
        TEXTS[language]["verify_account_phone"],
        reply_markup=phone_keyboard(language),
    )


def verify_account_phone(message: types.Message, current: dict[str, Any]) -> None:
    language = current.get("language") or "uz"
    text = TEXTS[language]
    existing_user = storage.get_user(message.from_user.id)

    if message.text == text["back"]:
        current.clear()
        bot.send_message(
            message.chat.id,
            text["main_menu"] if existing_user and existing_user.get("status") == "approved" else WELCOME,
            reply_markup=(
                main_menu(language)
                if existing_user and existing_user.get("status") == "approved"
                else start_keyboard()
            ),
        )
        return

    contact = message.contact
    if (
        not contact
        or contact.user_id is None
        or int(contact.user_id) != int(message.from_user.id)
    ):
        bot.send_message(
            message.chat.id,
            text["own_contact_required"],
            reply_markup=phone_keyboard(language),
        )
        return

    account = str(current.get("pending_account") or "")
    abonent = book.find(account)
    if not abonent:
        current.clear()
        bot.send_message(
            message.chat.id,
            text["account_not_found"],
            reply_markup=main_menu(language) if existing_user else start_keyboard(),
        )
        return

    phone = contact.phone_number or ""
    verified = account_phone_matches(abonent, phone)
    status = "approved" if verified else "pending"
    method = "telegram_phone_match" if verified else "operator_required"

    if existing_user:
        storage.link_account(message.from_user.id, account, status, method)
        storage.update_phone(message.from_user.id, phone)
    else:
        storage.create_user_if_absent(
            message.from_user.id,
            account,
            language,
            phone,
            status=status,
        )
        storage.link_account(message.from_user.id, account, status, method)

    current.clear()
    if verified:
        storage.set_active_account(message.from_user.id, account)
        notify_registration(abonent, phone)
        audit(
            message.from_user.id,
            "account_auto_verified",
            f"telegram_id={message.from_user.id}; account={account}",
        )
        bot.send_message(
            message.chat.id,
            text["account_auto_approved"],
            reply_markup=main_menu(language),
        )
        return

    notify_approval_request(abonent, message.from_user.id)
    audit(
        message.from_user.id,
        "account_manual_verification_requested",
        f"telegram_id={message.from_user.id}; account={account}",
    )
    bot.send_message(
        message.chat.id,
        text["account_pending"],
        reply_markup=(
            main_menu(language)
            if existing_user and existing_user.get("status") == "approved"
            else start_keyboard()
        ),
    )


def show_linked_accounts(message: types.Message, language: str) -> None:
    accounts = storage.linked_accounts(message.from_user.id)
    markup, mapping = account_list_keyboard(language, accounts)
    current = session(message)
    current["state"] = states.CHOOSING_ACCOUNT
    current["account_choices"] = mapping
    bot.send_message(
        message.chat.id,
        TEXTS[language]["accounts_title"],
        reply_markup=markup,
    )


def choose_linked_account(message: types.Message, current: dict[str, Any]) -> None:
    user = storage.get_user(message.from_user.id)
    language = (user or {}).get("language") or "uz"
    text = TEXTS[language]
    message_text = message.text or ""

    if message_text == text["back"]:
        current.clear()
        bot.send_message(message.chat.id, text["main_menu"], reply_markup=main_menu(language))
        return
    if message_text == text["add_account"]:
        current.clear()
        current["state"] = states.ENTERING_ACCOUNT
        current["language"] = language
        current["adding_account"] = True
        bot.send_message(message.chat.id, text["enter_account"])
        return

    account = (current.get("account_choices") or {}).get(message_text)
    if not account or not storage.set_active_account(message.from_user.id, account):
        show_linked_accounts(message, language)
        return

    current.clear()
    bot.send_message(
        message.chat.id,
        text["account_selected"],
        reply_markup=main_menu(language),
    )


def enter_account(message: types.Message, current: dict[str, Any]) -> None:
    language = current.get("language", "uz")
    account = book.clean_account(message.text or "")
    abonent = book.find(account)
    if not abonent:
        bot.send_message(message.chat.id, TEXTS[language]["account_not_found"])
        return

    linked = storage.get_linked_account(message.from_user.id, abonent.account)
    if linked and linked.get("status") == "approved":
        storage.set_active_account(message.from_user.id, abonent.account)
        current.clear()
        bot.send_message(
            message.chat.id,
            TEXTS[language]["main_menu"],
            reply_markup=main_menu(language),
        )
        return
    if linked and linked.get("status") == "pending":
        current.clear()
        existing_user = storage.get_user(message.from_user.id)
        bot.send_message(
            message.chat.id,
            access_waiting_text("pending", language),
            reply_markup=(
                main_menu(language)
                if existing_user and existing_user.get("status") == "approved"
                else start_keyboard()
            ),
        )
        return

    begin_account_verification(message, current, abonent, language)


def handle_admin_menu(message: types.Message, current: dict[str, Any]) -> bool:
    message_text = message.text or ""
    if message_text == ADMIN_BACK:
        bot.send_message(message.chat.id, "🛠 Admin panel\n\nBo'limni tanlang:", reply_markup=admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_EXCEL:
        bot.send_message(message.chat.id, "📁 Excel bo'limi", reply_markup=excel_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_ABONENTS:
        bot.send_message(message.chat.id, "🔎 Abonentlar bo'limi", reply_markup=abonents_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_REQUESTS:
        bot.send_message(message.chat.id, "📋 Arizalar bo'limi", reply_markup=requests_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_OPERATORS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu bo'lim faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, "👥 Operatorlar bo'limi", reply_markup=operators_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_BROADCAST:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu bo'lim faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        current["state"] = states.ADMIN_BROADCAST_UZ
        current.pop("broadcast_texts", None)
        bot.send_message(
            message.chat.id,
            "📢 Xabarlar bo'limi\n\n"
            "1/3 — O'zbek lotin tilidagi xabar matnini kiriting:",
            reply_markup=admin_input_keyboard(),
        )
        return True
    if message_text == ADMIN_SECTION_SERVICE:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu bo'lim faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, "🧹 Xizmat bo'limi", reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_REPORTS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu bo'lim faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, "📊 Hisobotlar bo'limi", reply_markup=reports_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SECTION_TABLE_REPORTS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        bot.send_message(
            message.chat.id,
            "📑 Jadval hisobotlari\n\nKerakli bo'limni tanlang:",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        return True
    if message_text == ADMIN_TABLE_BACK:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        bot.send_message(
            message.chat.id,
            "📑 Jadval hisobotlari\n\nKerakli bo'limni tanlang:",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        return True
    if message_text == ADMIN_TABLE_TEMPLATES:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        bot.send_message(
            message.chat.id,
            "📋 Shablonlar\n\n"
            "Yuklab olinadigan hisobotni tanlang:",
            reply_markup=table_templates_admin_keyboard(message.from_user.id),
        )
        return True
    if message_text == ADMIN_TEMPLATE_THREE_MONTHS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        report_path = latest_three_month_report("outputs/table_reports")
        if not report_path:
            bot.send_message(
                message.chat.id,
                "❌ 3 oylik hisobot hali yaratilmagan.",
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
            return True
        with report_path.open("rb") as file:
            bot.send_document(
                message.chat.id,
                file,
                caption=(
                    "📆 3 oylik ko'rsatkich olinmagan iste'molchilar\n\n"
                    "H ustun: faqat +\n"
                    "L ustun: joriy va oldingi ikki oy chiqarilgan\n"
                    "R ustun: olib tashlangan"
                ),
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
        audit(
            message.from_user.id,
            "table_report_downloaded",
            f"template=three_months; file={report_path.name}",
        )
        return True
    if message_text == ADMIN_TEMPLATE_DEBT_LIST:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        report_path = latest_debt_report("outputs/table_reports")
        if not report_path:
            bot.send_message(
                message.chat.id,
                "⏳ Qarzdorlar ro'yxati hali yaratilmagan.\nHisobot yaratilmoqda, iltimos kuting...",
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
            try:
                source = Path("data/report_templates/umumiy_shablon.xlsx")
                if not source.exists():
                    bot.send_message(
                        message.chat.id,
                        "❌ Umumiy shablon (F3 fayli) topilmadi.",
                        reply_markup=table_templates_admin_keyboard(message.from_user.id),
                    )
                    return True
                from datetime import date as _date
                today_str = _date.today().strftime("%Y-%m")
                out = Path(f"outputs/table_reports/qarzdorlar_royxati_{today_str}.xlsx")
                t = Path("data/report_templates/template_qarzdorlar.xlsx")
                result = generate_debt_report(source, out, template_path=t if t.exists() else None)
                report_path = result.output_path
                audit(
                    message.from_user.id,
                    "debt_report_generated",
                    f"source_rows={result.source_rows}; kept={result.kept_rows}",
                )
            except Exception as exc:
                bot.send_message(
                    message.chat.id,
                    f"❌ Hisobot yaratishda xatolik: {exc}",
                    reply_markup=table_templates_admin_keyboard(message.from_user.id),
                )
                return True
        with report_path.open("rb") as file:
            bot.send_document(
                message.chat.id,
                file,
                caption=(
                    "💸 Qarzdorlar ro'yxati\n\n"
                    "A–N ustunlar: abonent ma'lumotlari\n"
                    "O ustun: qarz miqdori (faqat > 0)\n"
                    "P–R ustunlar: olib tashlangan"
                ),
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
        audit(
            message.from_user.id,
            "debt_report_downloaded",
            f"file={report_path.name}",
        )
        return True
    if message_text == ADMIN_TEMPLATE_CREDITOR_LIST:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        report_path = latest_creditor_report("outputs/table_reports")
        if not report_path:
            bot.send_message(
                message.chat.id,
                "⏳ Kreditor ro'yxat hali yaratilmagan.\nHisobot yaratilmoqda, iltimos kuting...",
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
            try:
                source = Path("data/report_templates/umumiy_shablon.xlsx")
                if not source.exists():
                    bot.send_message(
                        message.chat.id,
                        "❌ Umumiy shablon (F3 fayli) topilmadi.",
                        reply_markup=table_templates_admin_keyboard(message.from_user.id),
                    )
                    return True
                from datetime import date as _date
                today_str = _date.today().strftime("%Y-%m")
                out = Path(f"outputs/table_reports/kreditor_royxati_{today_str}.xlsx")
                t = Path("data/report_templates/template_kreditor.xlsx")
                result = generate_creditor_report(source, out, template_path=t if t.exists() else None)
                report_path = result.output_path
                audit(
                    message.from_user.id,
                    "creditor_report_generated",
                    f"source_rows={result.source_rows}; kept={result.kept_rows}",
                )
            except Exception as exc:
                bot.send_message(
                    message.chat.id,
                    f"❌ Hisobot yaratishda xatolik: {exc}",
                    reply_markup=table_templates_admin_keyboard(message.from_user.id),
                )
                return True
        with report_path.open("rb") as file:
            bot.send_document(
                message.chat.id,
                file,
                caption=(
                    "💰 Kreditor iste'molchilar ro'yxati\n\n"
                    "A–N ustunlar: abonent ma'lumotlari\n"
                    "O ustun (P asl): balans (faqat < 0)\n"
                    "O va R asl ustunlar olib tashlangan"
                ),
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
        audit(
            message.from_user.id,
            "creditor_report_downloaded",
            f"file={report_path.name}",
        )
        return True
    if message_text == ADMIN_TEMPLATE_NO_ACCOUNTING:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        report_path = latest_no_accounting_report("outputs/table_reports")
        if not report_path:
            bot.send_message(
                message.chat.id,
                "⏳ Hisob yuritilmayotganlar ro'yxati hali yaratilmagan.\nHisobot yaratilmoqda, iltimos kuting...",
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
            try:
                source = Path("data/report_templates/umumiy_shablon.xlsx")
                if not source.exists():
                    bot.send_message(
                        message.chat.id,
                        "❌ Umumiy shablon (F3 fayli) topilmadi.",
                        reply_markup=table_templates_admin_keyboard(message.from_user.id),
                    )
                    return True
                from datetime import date as _date
                today_str = _date.today().strftime("%Y-%m")
                out = Path(f"outputs/table_reports/hisob_yuritilmayotgan_{today_str}.xlsx")
                t = Path("data/report_templates/template_hisob.xlsx")
                result = generate_no_accounting_report(source, out, template_path=t if t.exists() else None)
                report_path = result.output_path
                audit(
                    message.from_user.id,
                    "no_accounting_report_generated",
                    f"source_rows={result.source_rows}; kept={result.kept_rows}",
                )
            except Exception as exc:
                bot.send_message(
                    message.chat.id,
                    f"❌ Hisobot yaratishda xatolik: {exc}",
                    reply_markup=table_templates_admin_keyboard(message.from_user.id),
                )
                return True
        with report_path.open("rb") as file:
            bot.send_document(
                message.chat.id,
                file,
                caption=(
                    "📭 Hisob kitobi yuritilmayotgan iste'molchilar\n\n"
                    "Barcha 20 ustun saqlanadi (A–T)\n"
                    "N ustun: 0,00 yoki hisobsiz (< 1)"
                ),
                reply_markup=table_templates_admin_keyboard(message.from_user.id),
            )
        audit(
            message.from_user.id,
            "no_accounting_report_downloaded",
            f"file={report_path.name}",
        )
        return True
    if message_text == ADMIN_TABLE_UPLOAD_F3:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        current["state"] = states.ADMIN_UPLOAD_F3
        f3_path = Path("data/report_templates/f3_full.xlsx")
        exists_note = (
            f"\n\n✅ Joriy F3 fayli mavjud: <code>{f3_path.name}</code>"
            if f3_path.exists()
            else "\n\n⚠️ Hozircha F3 fayli yuklanmagan."
        )
        bot.send_message(
            message.chat.id,
            "📥 <b>To'liq F3 fayli yuklash</b>\n\n"
            "104 ustunli to'liq F3 Excel faylini yuboring.\n"
            "Bu fayl tuman pasportini avtomatik hisoblash uchun ishlatiladi.\n"
            "(Eslatma: bu 20 ustunli umumiy shablon emas, balki to'liq F3!)\n\n"
            "⚠️ <b>Fayl hajmi 20 MB dan katta bo'lsa</b>, Telegram orqali yuklash "
            "imkonsiz. Buning o'rniga fayl joylashgan to'liq yo'lni matn sifatida "
            "yuboring:\n"
            "<code>C:\\Users\\user\\Documents\\f3.xlsx</code>"
            + exists_note,
            parse_mode="HTML",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        return True
    if message_text == ADMIN_TABLE_DISTRICT_PASSPORT:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu bo'lim faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        f3_path = Path("data/report_templates/f3_full.xlsx")
        if not f3_path.exists():
            bot.send_message(
                message.chat.id,
                "⚠️ To'liq F3 fayli topilmadi.\n\n"
                "Tuman pasportini yaratish uchun avval "
                "<b>«📥 Yangi F3 jadvalni yuborish»</b> tugmasini bosib "
                "104 ustunli F3 faylini yuboring.",
                parse_mode="HTML",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            return True
        from datetime import date as _date
        today_str = _date.today().strftime("%Y-%m")
        report_path = latest_district_passport("outputs/table_reports")
        if report_path:
            with open(report_path, "rb") as f:
                bot.send_document(
                    message.chat.id,
                    f,
                    caption=(
                        f"🏢 Tuman pasporti\n📅 {report_path.stem}"
                    ),
                    reply_markup=table_reports_admin_keyboard(message.from_user.id),
                )
            audit(
                message.from_user.id,
                "district_passport_downloaded",
                f"file={report_path.name}",
            )
            return True
        bot.send_message(
            message.chat.id,
            "⏳ Tuman pasporti hali yaratilmagan.\nHisobot yaratilmoqda, iltimos kuting...",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        try:
            out = Path(f"outputs/table_reports/tuman_pasporti_{today_str}.xlsx")
            result = generate_district_passport(f3_path, out)
            report_path = result.output_path
            audit(
                message.from_user.id,
                "district_passport_generated",
                f"ref_date={result.reference_date}; me={result.me_rows}; chd={result.chd_rows}",
            )
        except Exception as exc:
            bot.send_message(
                message.chat.id,
                f"❌ Pasport yaratishda xatolik: {exc}",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            return True
        with open(report_path, "rb") as f:
            bot.send_document(
                message.chat.id,
                f,
                caption=(
                    f"🏢 Tuman pasporti\n"
                    f"📅 Sana: {result.reference_date.strftime('%d.%m.%Y')}\n"
                    f"🏠 Ko'p qavatli: {result.me_rows:,}\n"
                    f"🏡 Xususiy: {result.chd_rows:,}"
                ),
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
        return True
    if message_text == ADMIN_INSPECTORS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu bo'lim faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, "👷 Inspektorlar bo'limi", reply_markup=inspectors_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_STATS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, admin_stats_text(), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_RELOAD:
        reload_excel(message)
        return True
    if message_text == ADMIN_SEARCH:
        current["state"] = states.ADMIN_SEARCH
        bot.send_message(message.chat.id, "🔍 Hisob raqamini kiriting:", reply_markup=abonents_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_BROADCAST:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        current["state"] = states.ADMIN_BROADCAST_UZ
        current.pop("broadcast_texts", None)
        bot.send_message(
            message.chat.id,
            "📢 1/3\n\nO'zbek lotin matnini kiriting:",
            reply_markup=admin_input_keyboard(),
        )
        return True
    if message_text == ADMIN_RECENT_SUBMISSIONS:
        bot.send_message(message.chat.id, recent_submissions_text(), reply_markup=requests_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_PENDING_USERS:
        send_pending_users(message.chat.id, message.from_user.id)
        return True
    if message_text == ADMIN_UPLOAD_ABONENTS:
        current["state"] = states.ADMIN_UPLOAD_ABONENTS
        bot.send_message(message.chat.id, "📁 Yangi abonents.xlsx faylini yuboring.", reply_markup=excel_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_UPLOAD_INSPECTORS:
        current["state"] = states.ADMIN_UPLOAD_INSPECTORS
        bot.send_message(message.chat.id, "👮 Yangi inspectors.xlsx faylini yuboring.", reply_markup=excel_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_AUDIT_LOGS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, audit_logs_text(), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_AUDIT_EXPORT:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        export_path = export_audit_csv(message.from_user.id)
        try:
            with export_path.open("rb") as file:
                bot.send_document(message.chat.id, file, caption="📤 Audit loglar CSV eksport qilindi.", reply_markup=service_admin_keyboard(message.from_user.id))
        finally:
            secure_export_after_send(export_path, "audit_export")
        return True
    if message_text == ADMIN_BACKUP:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        paths = create_backup(message.from_user.id, "admin_button")
        bot.send_message(message.chat.id, backup_result_text(paths), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_BACKUP_LIST:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, backup_list_text(), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_BACKUP_RESTORE:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        backup = latest_valid_backup(
            config.backup_dir,
            backup_encryption_password(),
        )
        if backup is None:
            bot.send_message(
                message.chat.id,
                "❌ Tiklash uchun sog'lom backup topilmadi.",
                reply_markup=service_admin_keyboard(message.from_user.id),
            )
            return True
        current["state"] = states.ADMIN_BACKUP_RESTORE_CONFIRM
        current["backup_restore_path"] = str(backup.directory)
        bot.send_message(
            message.chat.id,
            "⚠️ Oxirgi sog'lom backup tiklanadi.\n\n"
            f"📅 Sana: {backup.created_at.strftime('%d.%m.%Y %H:%M:%S')}\n"
            f"📝 Sabab: {backup.reason or '-'}\n"
            f"📦 Fayllar: {len(backup.files)}\n\n"
            "Joriy baza va Excel fayllari almashtiriladi. Tiklashdan oldin "
            "joriy holatning xavfsizlik nusxasi olinadi.",
            reply_markup=backup_restore_confirm_keyboard(),
        )
        return True
    if message_text == ADMIN_SYSTEM_STATUS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, system_status_text(), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_USER_STATS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, user_stats_text(), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_SETTINGS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, settings_text(), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_DEBT_HISTORY:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, debt_history_text(), reply_markup=reports_admin_keyboard(message.from_user.id))
        return True
    export_actions = {
        ADMIN_EXPORT_USERS: export_users_csv,
        ADMIN_EXPORT_DEBTORS: export_debtors_csv,
        ADMIN_EXPORT_SUBMISSIONS: export_submissions_csv,
        ADMIN_EXPORT_DEBT_HISTORY: export_debt_history_csv,
    }
    if message_text in export_actions:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        export_path = export_actions[message_text](message.from_user.id)
        try:
            with export_path.open("rb") as file:
                bot.send_document(message.chat.id, file, caption=f"{message_text} tayyor.", reply_markup=reports_admin_keyboard(message.from_user.id))
        finally:
            secure_export_after_send(export_path)
        return True
    if message_text == ADMIN_LIST_INSPECTORS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, inspector_users_text(), reply_markup=inspectors_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_INSPECTOR_LOCATION_COMMAND:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu amal faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        send_inspector_location_command(
            message.from_user.id,
            message.chat.id,
        )
        return True
    if message_text == ADMIN_INSPECTOR_LOCATION_STATUS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(
                message.chat.id,
                "⛔ Bu amal faqat super admin uchun.",
                reply_markup=admin_keyboard(message.from_user.id),
            )
            return True
        active = storage.active_inspector_location_session()
        text = (
            inspector_location_session_summary(int(active["id"]))
            if active
            else "📡 Hozir faol live lokatsiya nazorati mavjud emas."
        )
        bot.send_message(
            message.chat.id,
            text,
            reply_markup=inspectors_admin_keyboard(message.from_user.id),
        )
        return True
    if message_text == ADMIN_ADD_INSPECTOR:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        current["state"] = states.ADMIN_ADD_INSPECTOR_ID
        bot.send_message(message.chat.id, "👷 Inspektorning Telegram ID raqamini kiriting:", reply_markup=inspectors_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_ADD_OPERATOR:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        current["state"] = states.ADMIN_ADD_OPERATOR
        bot.send_message(message.chat.id, "👨‍💼 Operator Telegram ID raqamini kiriting:", reply_markup=operators_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_LIST_OPERATORS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, operators_text(), reply_markup=operators_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_REMOVE_OPERATOR:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        current["state"] = states.ADMIN_REMOVE_OPERATOR
        bot.send_message(message.chat.id, "🗑 O'chiriladigan operator Telegram ID raqamini kiriting:", reply_markup=operators_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_CLEAR_TEST_DATA:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        current["state"] = states.ADMIN_CLEAR_TEST_DATA_CONFIRM
        bot.send_message(
            message.chat.id,
            "⚠️ Diqqat!\n\nBu amal test iste'molchilar va arizalarni o'chiradi.\nOperatorlar saqlanib qoladi.\n\nDavom etasizmi?",
            reply_markup=clear_test_data_keyboard(),
        )
        return True
    if message_text == ADMIN_DEBT_ANALYSIS:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        bot.send_message(message.chat.id, "⏳ Qarzdorlik tahlili bajarilmoqda...", reply_markup=service_admin_keyboard(message.from_user.id))
        result = run_debt_analysis(send_messages=False)
        bot.send_message(message.chat.id, debt_analysis_summary(result, send_messages=False), reply_markup=service_admin_keyboard(message.from_user.id))
        return True
    if message_text == ADMIN_DEBT_SEND:
        if not is_super_admin_id(message.from_user.id):
            bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
            return True
        preview = run_debt_analysis(send_messages=False)
        current["state"] = states.ADMIN_DEBT_SEND_CONFIRM
        bot.send_message(
            message.chat.id,
            "⚠️ Qarzdor iste'molchilarga ogohlantirish yuborish\n\n"
            f"📅 Excel sanasi: {preview['statement_date']}\n"
            f"🔢 Tekshirilgan hisoblar: {preview['checked']}\n"
            f"🔴 Qarzdor hisoblar: {preview['debtors']}\n"
            f"⏭ Excelda topilmaganlar: {preview['skipped']}\n\n"
            "Bugun avval muvaffaqiyatli ogohlantirilgan hisob raqamlariga "
            "xabar takroran yuborilmaydi.\n\n"
            "Davom etasizmi?",
            reply_markup=debt_send_confirm_keyboard(),
        )
        return True
    if message_text == ADMIN_EXIT:
        clear_session(message)
        bot.send_message(message.chat.id, "Admin panel yopildi.", reply_markup=start_keyboard())
        return True
    return False


def handle_admin_state(message: types.Message, current: dict[str, Any]) -> None:
    if not is_staff(message):
        clear_session(message)
        bot.send_message(message.chat.id, admin_denied_text(message), parse_mode="Markdown")
        return

    if message.text == ADMIN_BACK:
        clear_admin_flow(message)
        bot.send_message(
            message.chat.id,
            "🛠 Admin panel\n\nBo'limni tanlang:",
            reply_markup=admin_keyboard(message.from_user.id),
        )
        return
    if message.text == ADMIN_EXIT:
        clear_session(message)
        bot.send_message(message.chat.id, "Admin panel yopildi.", reply_markup=start_keyboard())
        return
    if message.text == CANCEL:
        state = current.get("state")
        if state == states.ADMIN_CONFIRM_EXCEL_UPLOAD:
            cleanup_pending_excel(current)
        clear_admin_flow(message)
        if state == states.ADMIN_CLEAR_TEST_DATA_CONFIRM:
            bot.send_message(message.chat.id, "🧹 Tozalash bekor qilindi.", reply_markup=service_admin_keyboard(message.from_user.id))
        elif state == states.ADMIN_DEBT_SEND_CONFIRM:
            bot.send_message(message.chat.id, "🔴 Qarzdorlik ogohlantirishi bekor qilindi.", reply_markup=service_admin_keyboard(message.from_user.id))
        elif state == states.ADMIN_CONFIRM_EXCEL_UPLOAD:
            bot.send_message(message.chat.id, "📁 Excel almashtirish bekor qilindi.", reply_markup=excel_admin_keyboard(message.from_user.id))
        elif state == states.ADMIN_BACKUP_RESTORE_CONFIRM:
            bot.send_message(message.chat.id, "♻️ Backup tiklash bekor qilindi.", reply_markup=service_admin_keyboard(message.from_user.id))
        elif str(state).startswith("admin_broadcast"):
            bot.send_message(message.chat.id, "📢 Xabar yuborish bekor qilindi.", reply_markup=broadcast_admin_keyboard(message.from_user.id))
        else:
            bot.send_message(message.chat.id, "Amal bekor qilindi.", reply_markup=admin_keyboard(message.from_user.id))
        return

    state = current.get("state")
    if state == states.ADMIN_SEARCH:
        admin_search(message)
    elif state == states.ADMIN_UPLOAD_ABONENTS:
        admin_upload_excel(message, Path(config.excel_file), "abonents.xlsx")
    elif state == states.ADMIN_UPLOAD_INSPECTORS:
        admin_upload_excel(message, Path(config.inspectors_file), "inspectors.xlsx")
    elif state == states.ADMIN_CONFIRM_EXCEL_UPLOAD:
        admin_confirm_excel_upload(message, current)
    elif state == states.ADMIN_ADD_OPERATOR:
        admin_add_operator(message)
    elif state == states.ADMIN_REMOVE_OPERATOR:
        admin_remove_operator(message)
    elif state == states.ADMIN_REMOVE_OPERATOR_CONFIRM:
        admin_confirm_remove_operator(message, current)
    elif state == states.ADMIN_CLEAR_TEST_DATA_CONFIRM:
        admin_clear_test_data(message)
    elif state == states.ADMIN_DEBT_SEND_CONFIRM:
        admin_debt_send(message)
    elif state == states.ADMIN_BACKUP_RESTORE_CONFIRM:
        admin_restore_backup(message, current)
    elif state == states.ADMIN_BROADCAST_UZ:
        admin_broadcast_collect(message, current, "uz", "📢 2/3\n\nЎзбек кирилл матнини киритинг:", states.ADMIN_BROADCAST_UZ_CYRL)
    elif state == states.ADMIN_BROADCAST_UZ_CYRL:
        admin_broadcast_collect(message, current, "uz_cyrl", "📢 3/3\n\nРусский текст введите:", states.ADMIN_BROADCAST_RU)
    elif state == states.ADMIN_BROADCAST_RU:
        admin_broadcast_collect(message, current, "ru", "", states.ADMIN_BROADCAST_CONFIRM)
    elif state == states.ADMIN_BROADCAST_CONFIRM:
        admin_broadcast_send(message, current)
    elif state == states.ADMIN_UPLOAD_F3:
        admin_upload_f3_full(message, current)
    elif state in {
        states.ADMIN_ADD_INSPECTOR_ID,
        states.ADMIN_ADD_INSPECTOR_AREA,
        states.ADMIN_ADD_INSPECTOR_FIO,
        states.ADMIN_ADD_INSPECTOR_PHONE,
    }:
        admin_add_inspector_flow(message, current)
    else:
        clear_admin_flow(message)
        bot.send_message(
            message.chat.id,
            "⚠️ Amal holati eskirgan. Kerakli bo'limni qayta tanlang.",
            reply_markup=admin_keyboard(message.from_user.id),
        )


def admin_upload_f3_full(message: types.Message, current: dict) -> None:
    """104 ustunli to'liq F3 faylini qabul qiladi va saqlaydi.

    Ikki usul:
    1. Telegram document (≤20 MB) — yuklab saqlanadi.
    2. Lokal fayl yo'li (matn) — fayl serverda mavjud bo'lsa to'g'ridan o'qiladi.
    """
    import shutil

    dest = Path("data/report_templates/f3_full.xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)

    # --- Usul 2: lokal yo'l matn sifatida ---
    if message.document is None:
        raw = (message.text or "").strip().strip('"').strip("'")
        if not raw:
            bot.send_message(
                message.chat.id,
                "⚠️ Excel fayl yoki fayl yo'lini yuboring.",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            return
        local_path = Path(raw)
        if not local_path.exists():
            bot.send_message(
                message.chat.id,
                f"❌ Fayl topilmadi:\n<code>{raw}</code>",
                parse_mode="HTML",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            return
        if not local_path.suffix.lower() == ".xlsx":
            bot.send_message(
                message.chat.id,
                "❌ Faqat .xlsx formatidagi fayl qabul qilinadi.",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            return
        bot.send_message(
            message.chat.id,
            "⏳ Fayl tekshirilmoqda...",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        try:
            err = validate_f3_full(local_path)
            if err:
                bot.send_message(
                    message.chat.id,
                    f"❌ Fayl yaroqsiz: {err}",
                    reply_markup=table_reports_admin_keyboard(message.from_user.id),
                )
                return
            size_mb = local_path.stat().st_size / 1_048_576
            if local_path.resolve() != dest.resolve():
                shutil.copy2(str(local_path), str(dest))
            clear_admin_flow(message)
            audit(
                message.from_user.id,
                "f3_full_uploaded_local",
                f"path={raw}; size={size_mb:.2f}MB",
            )
            _auto_regenerate_all_reports(message, dest)
        except Exception as exc:
            bot.send_message(
                message.chat.id,
                f"❌ Xatolik: {exc}",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
        return

    # --- Usul 1: Telegram document ---
    if not message.document.file_name.lower().endswith(".xlsx"):
        bot.send_message(
            message.chat.id,
            "❌ Faqat .xlsx formatidagi fayl qabul qilinadi.",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        return
    file_size_mb = (message.document.file_size or 0) / 1_048_576
    if file_size_mb > 19:
        bot.send_message(
            message.chat.id,
            f"❌ Fayl hajmi {file_size_mb:.1f} MB — Telegram orqali max 19 MB.\n\n"
            "Fayl yo'lini matn sifatida yuboring:\n"
            "<code>C:\\Users\\user\\Documents\\f3.xlsx</code>",
            parse_mode="HTML",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )
        return
    bot.send_message(
        message.chat.id,
        "⏳ Fayl yuklanmoqda va tekshirilmoqda...",
        reply_markup=table_reports_admin_keyboard(message.from_user.id),
    )
    try:
        work_path = Path("work/uploads/f3_full_incoming.xlsx")
        work_path.parent.mkdir(parents=True, exist_ok=True)
        file_info = bot.get_file(message.document.file_id)
        downloaded = bot.download_file(file_info.file_path)
        with open(work_path, "wb") as fh:
            fh.write(downloaded)
        err = validate_f3_full(work_path)
        if err:
            bot.send_message(
                message.chat.id,
                f"❌ Fayl yaroqsiz: {err}",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            work_path.unlink(missing_ok=True)
            return
        shutil.move(str(work_path), str(dest))
        clear_admin_flow(message)
        audit(
            message.from_user.id,
            "f3_full_uploaded",
            f"file={message.document.file_name}; size={file_size_mb:.2f}MB",
        )
        # Barcha shablonlar va pasportni avtomatik yangilash
        _auto_regenerate_all_reports(message, dest)
    except Exception as exc:
        bot.send_message(
            message.chat.id,
            f"❌ Xatolik: {exc}",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )


def _auto_regenerate_all_reports(
    message: types.Message,
    f3_path: Path,
) -> None:
    """F3 yuklangandan so'ng barcha shablonlar va pasportni qayta yaratadi."""
    from datetime import date as _date
    import threading

    def _run() -> None:
        today_str = _date.today().strftime("%Y-%m")
        out_dir = Path("outputs/table_reports")
        out_dir.mkdir(parents=True, exist_ok=True)
        shablon_path = Path("data/report_templates/umumiy_shablon.xlsx")
        lines: list[str] = []
        errors: list[str] = []

        # 1. umumiy_shablon.xlsx yaratish
        try:
            bot.send_message(
                message.chat.id,
                "⏳ <b>F3 qayta ishlanmoqda...</b>\n"
                "1/6 umumiy shablon yaratilmoqda (~40 soniya)...",
                parse_mode="HTML",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            row_count, ref_date = extract_umumiy_shablon(f3_path, shablon_path)
            ref_str = ref_date.strftime("%d.%m.%Y") if ref_date else "?"
            lines.append(f"✅ Umumiy shablon: {row_count:,} qator ({ref_str})")
        except Exception as exc:
            errors.append(f"❌ Umumiy shablon: {exc}")
            bot.send_message(
                message.chat.id,
                f"❌ Umumiy shablon yaratishda xatolik: {exc}",
                reply_markup=table_reports_admin_keyboard(message.from_user.id),
            )
            return

        # 2–5. To'rt shablon (original template formatini saqlash uchun template_path beriladi)
        tmpl_dir = Path("data/report_templates")
        template_jobs = [
            ("3 oylik ko'rsatkich olinmaganlar", generate_three_month_report,
             out_dir / f"3_oylik_korsatgich_olinmagan_{today_str}.xlsx", "three_month",
             tmpl_dir / "template_3oylik.xlsx"),
            ("Qarzdorlar ro'yxati", generate_debt_report,
             out_dir / f"qarzdorlar_royxati_{today_str}.xlsx", "debt",
             tmpl_dir / "template_qarzdorlar.xlsx"),
            ("Kreditor iste'molchilar", generate_creditor_report,
             out_dir / f"kreditor_royxati_{today_str}.xlsx", "creditor",
             tmpl_dir / "template_kreditor.xlsx"),
            ("Hisob yuritilmayotganlar", generate_no_accounting_report,
             out_dir / f"hisob_yuritilmayotgan_{today_str}.xlsx", "no_accounting",
             tmpl_dir / "template_hisob.xlsx"),
        ]
        for step, (label, gen_fn, out_path, _key, tmpl_path) in enumerate(template_jobs, 2):
            try:
                out_path.unlink(missing_ok=True)
                t = tmpl_path if tmpl_path.exists() else None
                result = gen_fn(shablon_path, out_path, template_path=t)
                lines.append(
                    f"✅ {label}: {result.kept_rows:,}/{result.source_rows:,} qator"
                )
                audit(message.from_user.id, f"{_key}_report_regenerated",
                      f"rows={result.kept_rows}")
            except Exception as exc:
                errors.append(f"❌ {label}: {exc}")

        # 6. Tuman pasporti
        try:
            passport_out = out_dir / f"tuman_pasporti_{today_str}.xlsx"
            passport_out.unlink(missing_ok=True)
            pr = generate_district_passport(f3_path, passport_out)
            lines.append(
                f"✅ Tuman pasporti: ME={pr.me_rows:,} | CHD={pr.chd_rows:,} "
                f"({pr.reference_date.strftime('%d.%m.%Y')})"
            )
            audit(message.from_user.id, "district_passport_regenerated",
                  f"ref={pr.reference_date}; me={pr.me_rows}; chd={pr.chd_rows}")
        except Exception as exc:
            errors.append(f"❌ Tuman pasporti: {exc}")

        summary = "\n".join(lines + errors)
        bot.send_message(
            message.chat.id,
            f"<b>F3 asosida yangilash yakunlandi:</b>\n\n{summary}",
            parse_mode="HTML",
            reply_markup=table_reports_admin_keyboard(message.from_user.id),
        )

    threading.Thread(target=_run, daemon=True).start()


def reload_excel(message: types.Message) -> None:
    try:
        book.reload()
    except Exception as error:
        audit(
            message.from_user.id,
            "excel_reload_failed",
            safe_exception_summary(error),
        )
        notify_error("Excel qayta yuklanmadi", error)
        bot.send_message(
            message.chat.id,
            "❌ Excel qayta yuklanmadi. Tafsilot xatolik jurnaliga yozildi.",
            reply_markup=excel_admin_keyboard(message.from_user.id),
        )
        return
    audit(message.from_user.id, "excel_reloaded", f"abonents={book.count()}")
    bot.send_message(message.chat.id, f"✅ Excel qayta yuklandi.\n📚 Abonentlar soni: {book.count()}", reply_markup=excel_admin_keyboard(message.from_user.id))


def operators_text() -> str:
    operators = storage.operators()
    if not operators:
        return "👥 Operatorlar ro'yxati\n\nHozircha operatorlar yo'q."

    lines = ["👥 Operatorlar ro'yxati"]
    for index, operator in enumerate(operators, 1):
        language = operator.get("language") or "til tanlanmagan"
        lines.append(
            "\n"
            f"{index}. Telegram ID: {operator['telegram_id']}\n"
            f"🌐 Til: {language}\n"
            f"📅 Qo'shilgan sana: {operator['created_at']}"
        )
    return "\n".join(lines)


def _parse_telegram_id(text: str) -> int | None:
    raw = text.strip()
    if raw.isdigit():
        return int(raw)
    return None


def admin_add_operator(message: types.Message) -> None:
    if not is_super_admin_id(message.from_user.id):
        bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
        return

    operator_id = _parse_telegram_id(message.text or "")
    if operator_id is None:
        bot.send_message(message.chat.id, "⚠️ Telegram ID faqat raqamlardan iborat bo'lishi kerak.", reply_markup=operators_admin_keyboard(message.from_user.id))
        return

    storage.add_operator(operator_id)
    audit(message.from_user.id, "operator_added", f"operator_id={operator_id}")
    clear_admin_flow(message)
    bot.send_message(
        message.chat.id,
        f"✅ Operator qo'shildi.\n\nTelegram ID: {operator_id}\n\nOperator /start bosganda til tanlaydi va operator panelga kiradi.",
        reply_markup=operators_admin_keyboard(message.from_user.id),
    )


def admin_remove_operator(message: types.Message) -> None:
    if not is_super_admin_id(message.from_user.id):
        bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
        return

    operator_id = _parse_telegram_id(message.text or "")
    if operator_id is None:
        bot.send_message(message.chat.id, "⚠️ Telegram ID faqat raqamlardan iborat bo'lishi kerak.", reply_markup=operators_admin_keyboard(message.from_user.id))
        return

    current = session(message)
    current["operator_id_to_remove"] = operator_id
    current["state"] = states.ADMIN_REMOVE_OPERATOR_CONFIRM
    bot.send_message(
        message.chat.id,
        "⚠️ Operatorni o'chirishni tasdiqlang.\n\n"
        f"Telegram ID: {operator_id}",
        reply_markup=remove_operator_confirm_keyboard(),
    )


def admin_confirm_remove_operator(
    message: types.Message,
    current: dict[str, Any],
) -> None:
    if not is_super_admin_id(message.from_user.id):
        clear_session(message)
        bot.send_message(
            message.chat.id,
            "⛔ Bu amal faqat super admin uchun.",
            reply_markup=admin_keyboard(message.from_user.id),
        )
        return
    if message.text != REMOVE_OPERATOR_CONFIRM:
        bot.send_message(
            message.chat.id,
            "Tasdiqlash uchun o'chirish tugmasini bosing.",
            reply_markup=remove_operator_confirm_keyboard(),
        )
        return

    operator_id = str(current.get("operator_id_to_remove") or "")
    if not operator_id:
        clear_admin_flow(message)
        bot.send_message(
            message.chat.id,
            "❌ Operator ID topilmadi. Amalni qayta boshlang.",
            reply_markup=operators_admin_keyboard(message.from_user.id),
        )
        return

    storage.remove_operator(int(operator_id))
    audit(message.from_user.id, "operator_removed", f"operator_id={operator_id}")
    clear_admin_flow(message)
    bot.send_message(
        message.chat.id,
        f"✅ Operator o'chirildi.\n\nTelegram ID: {operator_id}",
        reply_markup=operators_admin_keyboard(message.from_user.id),
    )


def admin_add_inspector_flow(message: types.Message, current: dict[str, Any]) -> None:
    if not is_super_admin_id(message.from_user.id):
        clear_session(message)
        bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
        return

    state = current.get("state")
    text = (message.text or "").strip()
    if state == states.ADMIN_ADD_INSPECTOR_ID:
        inspector_id = _parse_telegram_id(text)
        if inspector_id is None:
            bot.send_message(message.chat.id, "⚠️ Telegram ID faqat raqamlardan iborat bo'lishi kerak.", reply_markup=inspectors_admin_keyboard(message.from_user.id))
            return
        current["inspector_id"] = inspector_id
        current["state"] = states.ADMIN_ADD_INSPECTOR_AREA
        bot.send_message(message.chat.id, "🧭 Inspektor hududini kiriting, masalan: 39")
        return

    if state == states.ADMIN_ADD_INSPECTOR_AREA:
        if not text:
            bot.send_message(message.chat.id, "🧭 Hudud bo'sh bo'lmasin.")
            return
        current["inspector_area"] = text
        current["state"] = states.ADMIN_ADD_INSPECTOR_FIO
        bot.send_message(message.chat.id, "👤 Inspektor F.I.O ma'lumotini kiriting:")
        return

    if state == states.ADMIN_ADD_INSPECTOR_FIO:
        if not text:
            bot.send_message(message.chat.id, "👤 F.I.O bo'sh bo'lmasin.")
            return
        current["inspector_fio"] = text
        current["state"] = states.ADMIN_ADD_INSPECTOR_PHONE
        bot.send_message(message.chat.id, "☎️ Inspektor telefon raqamini kiriting yoki '-' yuboring:")
        return

    phone = "" if text == "-" else text
    inspector_id = int(current["inspector_id"])
    area = str(current["inspector_area"])
    fio = str(current["inspector_fio"])
    storage.add_inspector_user(inspector_id, area, fio, phone)
    audit(message.from_user.id, "inspector_user_added", f"telegram_id={inspector_id}; area={area}; fio={fio}")
    clear_admin_flow(message)
    bot.send_message(
        message.chat.id,
        "✅ Inspektor ulash uchun tayyorlandi.\n\n"
        f"Telegram ID: {inspector_id}\n"
        f"Hudud: {area}\n"
        f"F.I.O: {fio}\n"
        f"Telefon: {phone or '-'}",
        reply_markup=inspectors_admin_keyboard(message.from_user.id),
    )


def admin_clear_test_data(message: types.Message) -> None:
    if not is_super_admin_id(message.from_user.id):
        bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
        return
    if message.text != CLEAR_TEST_CONFIRM:
        bot.send_message(message.chat.id, "Tasdiqlash uchun 🧹 Ha, tozalash tugmasini bosing.", reply_markup=clear_test_data_keyboard())
        return

    storage.clear_test_data()
    audit(message.from_user.id, "test_data_cleared", "users/submissions/debt_warnings")
    clear_admin_flow(message)
    bot.send_message(
        message.chat.id,
        "✅ Test ma'lumotlar tozalandi.\n\nO'chirildi: iste'molchilar, pending/approved ro'yxatlar va arizalar.\nOperatorlar saqlanib qoldi.",
        reply_markup=service_admin_keyboard(message.from_user.id),
    )


def admin_debt_send(message: types.Message) -> None:
    if not is_super_admin_id(message.from_user.id):
        clear_session(message)
        bot.send_message(message.chat.id, "⛔ Bu amal faqat super admin uchun.", reply_markup=admin_keyboard(message.from_user.id))
        return

    if "ha" not in (message.text or "").lower() or "yuborish" not in (message.text or "").lower():
        bot.send_message(message.chat.id, "Tasdiqlash uchun 📨 Ha, yuborish tugmasini bosing.", reply_markup=debt_send_confirm_keyboard())
        return

    clear_admin_flow(message)
    started = start_debt_warning_job(
        chat_id=message.chat.id,
        actor_id=message.from_user.id,
        reason="manual_admin",
        force=False,
        calendar_only=False,
    )
    if not started:
        bot.send_message(
            message.chat.id,
            "⏳ Qarzdorlik ogohlantirishlarini yuborish jarayoni "
            "allaqachon davom etmoqda.",
            reply_markup=service_admin_keyboard(message.from_user.id),
        )
        return
    audit(message.from_user.id, "debt_warning_job_started", "manual_admin")
    bot.send_message(
        message.chat.id,
        "📨 Qarzdorlik ogohlantirishlarini yuborish fon rejimida "
        "boshlandi.\n\n"
        "Botdan foydalanishda davom etishingiz mumkin. Jarayon "
        "yakunlangach natija shu yerga yuboriladi. Bugun avval "
        "ogohlantirilgan hisob raqamlarga xabar takroran yuborilmaydi.",
        reply_markup=service_admin_keyboard(message.from_user.id),
    )


def admin_restore_backup(
    message: types.Message,
    current: dict[str, Any],
) -> None:
    global storage, book
    if not is_super_admin_id(message.from_user.id):
        clear_session(message)
        bot.send_message(
            message.chat.id,
            "⛔ Bu amal faqat super admin uchun.",
            reply_markup=admin_keyboard(message.from_user.id),
        )
        return
    if message.text != BACKUP_RESTORE_CONFIRM:
        bot.send_message(
            message.chat.id,
            "Tasdiqlash uchun backupni tiklash tugmasini bosing.",
            reply_markup=backup_restore_confirm_keyboard(),
        )
        return

    backup_path = str(current.get("backup_restore_path") or "")
    try:
        with maintenance_lock:
            password = backup_encryption_password()
            selected_backup = load_backup_set(
                backup_path,
                verify_hashes=True,
                encryption_password=password,
            )
            create_backup(message.from_user.id, "before_restore")
            restored = restore_backup_set(
                selected_backup,
                {
                    "database": Path(config.database_file),
                    "abonents": Path(config.excel_file),
                    "inspectors": Path(config.inspectors_file),
                    "meter_references": Path("work/restored_meter_reference_photos.zip"),
                },
                encryption_password=password,
            )
            restored_reference_archive = Path("work/restored_meter_reference_photos.zip")
            if restored_reference_archive in restored:
                restore_meter_reference_archive(restored_reference_archive)
                restored_reference_archive.unlink(missing_ok=True)
            storage = BotStorage(config.database_file)
            book = AbonentBook(
                config.excel_file,
                config.inspectors_file,
            )
        audit(
            message.from_user.id,
            "backup_restored",
            f"backup={selected_backup.directory.name}; files={len(restored)}",
        )
    except Exception as error:
        notify_error("Backup tiklanmadi", error)
        bot.send_message(
            message.chat.id,
            "❌ Backup tiklanmadi. Amaldagi fayllar saqlab qolindi; "
            "tafsilot xatolik jurnaliga yozildi.",
            reply_markup=service_admin_keyboard(message.from_user.id),
        )
        return

    clear_admin_flow(message)
    bot.send_message(
        message.chat.id,
        "✅ Backup muvaffaqiyatli tiklandi.\n\n"
        f"📅 Backup sanasi: {selected_backup.created_at.strftime('%d.%m.%Y %H:%M:%S')}\n"
        f"📦 Tiklangan fayllar: {len(restored)}",
        reply_markup=service_admin_keyboard(message.from_user.id),
    )


def admin_search(message: types.Message) -> None:
    account = book.clean_account(message.text or "")
    abonent = book.find(account)
    if not abonent:
        audit(message.from_user.id, "admin_search_failed", f"account={account}")
        bot.send_message(message.chat.id, "❌ Abonent topilmadi. Boshqa hisob raqam kiriting:", reply_markup=abonents_admin_keyboard(message.from_user.id))
        return

    audit(message.from_user.id, states.ADMIN_SEARCH, f"account={account}")
    bot.send_message(
        message.chat.id,
        "🔍 Abonent topildi\n\n" + report_caption(abonent, masked=False),
        reply_markup=abonents_admin_keyboard(message.from_user.id),
    )


def admin_upload_excel(message: types.Message, target: Path, expected_name: str) -> None:
    if not message.document:
        bot.send_message(message.chat.id, f"⚠️ Iltimos, {expected_name} Excel faylini document sifatida yuboring.", reply_markup=excel_admin_keyboard(message.from_user.id))
        return

    file_name = message.document.file_name or ""
    if not file_name.lower().endswith(".xlsx"):
        bot.send_message(message.chat.id, "⚠️ Faqat .xlsx fayl qabul qilinadi.", reply_markup=excel_admin_keyboard(message.from_user.id))
        return
    file_size = int(message.document.file_size or 0)
    maximum_bytes = int(config.max_excel_upload_mb * 1024 * 1024)
    if not file_size_allowed(file_size, config.max_excel_upload_mb):
        audit(
            message.from_user.id,
            "excel_upload_size_rejected",
            f"name={expected_name}; size={file_size}; max={maximum_bytes}",
        )
        bot.send_message(
            message.chat.id,
            f"⚠️ Excel fayl hajmi {config.max_excel_upload_mb:g} MB dan "
            "oshmasligi kerak.",
            reply_markup=excel_admin_keyboard(message.from_user.id),
        )
        return

    pending_path: Path | None = None
    try:
        file_info = bot.get_file(message.document.file_id)
        downloaded = bot.download_file(file_info.file_path)
        if not file_size_allowed(
            len(downloaded),
            config.max_excel_upload_mb,
        ):
            raise RuntimeError(
                f"Yuklangan faylning haqiqiy hajmi "
                f"{config.max_excel_upload_mb:g} MB limitdan oshdi."
            )
        upload_dir = Path("work") / "uploads"
        upload_dir.mkdir(parents=True, exist_ok=True)
        pending_path = upload_dir / f"{uuid4().hex}_{expected_name}"
        pending_path.write_bytes(downloaded)
        if expected_name == "abonents.xlsx":
            stats, validation_text = validate_abonents_excel(pending_path)
        else:
            stats, validation_text = validate_inspectors_excel(pending_path)
    except Exception as error:
        if pending_path:
            pending_path.unlink(missing_ok=True)
        safe_error = safe_exception_summary(error)
        audit(message.from_user.id, "excel_upload_validation_failed", f"{expected_name}: {safe_error}")
        notify_error("Excel yuklash tekshiruvida xatolik", error)
        bot.send_message(
            message.chat.id,
            "❌ Fayl xavfsizlik yoki tuzilma tekshiruvidan o'tmadi. "
            "Tafsilot administrator xatolik jurnaliga yozildi.",
            reply_markup=excel_admin_keyboard(message.from_user.id),
        )
        return

    current = session(message)
    current["state"] = states.ADMIN_CONFIRM_EXCEL_UPLOAD
    current["pending_excel"] = {
        "path": str(pending_path),
        "target": str(target),
        "expected_name": expected_name,
        "stats": stats,
    }
    audit(message.from_user.id, "excel_upload_validated", f"{expected_name}: {stats}")
    bot.send_message(message.chat.id, validation_text, reply_markup=excel_confirm_keyboard())


def admin_confirm_excel_upload(message: types.Message, current: dict[str, Any]) -> None:
    if message.text != EXCEL_REPLACE_CONFIRM:
        bot.send_message(message.chat.id, "Tasdiqlash uchun ✅ Excelni almashtirish tugmasini bosing.", reply_markup=excel_confirm_keyboard())
        return

    pending = current.get("pending_excel") or {}
    pending_path_text = str(pending.get("path") or "")
    target_text = str(pending.get("target") or "")
    pending_path = Path(pending_path_text) if pending_path_text else Path()
    target = Path(target_text) if target_text else Path()
    expected_name = str(pending.get("expected_name") or "Excel")
    if (
        not pending_path_text
        or not target_text
        or not pending_path.is_file()
    ):
        clear_admin_flow(message)
        bot.send_message(message.chat.id, "❌ Vaqtinchalik Excel fayl topilmadi. Qayta yuklang.", reply_markup=excel_admin_keyboard(message.from_user.id))
        return

    staged_path = target.with_name(f".{target.name}.replace.tmp")
    rollback_path = target.with_name(
        f".{target.name}.replace.rollback"
    )
    try:
        with maintenance_lock:
            create_backup(
                message.from_user.id,
                f"before_{expected_name}_replace",
            )
            target.parent.mkdir(parents=True, exist_ok=True)
            staged_path.unlink(missing_ok=True)
            rollback_path.unlink(missing_ok=True)
            shutil.copy2(pending_path, staged_path)
            if target.exists():
                shutil.copy2(target, rollback_path)
            staged_path.replace(target)
            try:
                book.reload()
            except Exception:
                if rollback_path.exists():
                    rollback_path.replace(target)
                    book.reload()
                else:
                    target.unlink(missing_ok=True)
                raise
        storage.set_setting(f"last_{expected_name}_upload", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    except Exception as error:
        audit(
            message.from_user.id,
            "excel_replace_failed",
            f"{expected_name}: {safe_exception_summary(error)}",
        )
        notify_error("Excel almashtirishda xatolik", error)
        bot.send_message(
            message.chat.id,
            "❌ Excel almashtirilmadi. Xavfsiz rollback bajarildi; "
            "tafsilot xatolik jurnaliga yozildi.",
            reply_markup=excel_admin_keyboard(message.from_user.id),
        )
        return
    finally:
        staged_path.unlink(missing_ok=True)
        rollback_path.unlink(missing_ok=True)

    pending_path.unlink(missing_ok=True)
    clear_admin_flow(message)
    audit(message.from_user.id, "excel_replaced", f"{expected_name}; abonents={book.count()}")
    bot.send_message(
        message.chat.id,
        f"✅ {expected_name} yangilandi va Excel qayta yuklandi.\n📚 Abonentlar soni: {book.count()}",
        reply_markup=excel_admin_keyboard(message.from_user.id),
    )
    if expected_name == "abonents.xlsx":
        due_today = debt_calendar_due(datetime.now())
        started = False
        if due_today:
            started = start_debt_warning_job(
                chat_id=message.chat.id,
                actor_id=message.from_user.id,
                reason="excel_upload_calendar",
                force=False,
                calendar_only=True,
            )
        audit(
            message.from_user.id,
            "debt_warning_job_after_excel",
            f"due_today={due_today}; started={started}",
        )
        next_run = next_debt_calendar_date().strftime("%d.%m.%Y")
        bot.send_message(
            message.chat.id,
            (
                "📨 Yangi Excel bo'yicha qarzdorlik ogohlantirishlari "
                "fon rejimida boshlandi. Yakuniy natija alohida yuboriladi."
                if due_today and started
                else "⏳ Qarzdorlik ogohlantirishlarini yuborish jarayoni "
                "allaqachon davom etmoqda. Ikkinchi jarayon ochilmadi."
                if due_today
                else "ℹ️ Excel yangilandi. Bugun avtomatik ogohlantirish "
                "kuni emas, shu sabab xabar yuborilmadi.\n"
                f"Keyingi avtomatik sana: {next_run}."
            ),
            reply_markup=excel_admin_keyboard(message.from_user.id),
        )


def admin_broadcast_collect(
    message: types.Message,
    current: dict[str, Any],
    language: str,
    next_prompt: str,
    next_state: str,
) -> None:
    text = (message.text or "").strip()
    if not text:
        bot.send_message(
            message.chat.id,
            "⚠️ Xabar matni bo'sh bo'lmasin.",
            reply_markup=admin_input_keyboard(),
        )
        return

    broadcast_texts = current.setdefault("broadcast_texts", {})
    broadcast_texts[language] = text
    current["state"] = next_state

    if next_state != states.ADMIN_BROADCAST_CONFIRM:
        bot.send_message(
            message.chat.id,
            next_prompt,
            reply_markup=admin_input_keyboard(),
        )
        return

    admin_broadcast_preview(message, current)


def admin_broadcast_preview(message: types.Message, current: dict[str, Any]) -> None:
    translations = current.get("broadcast_texts", {})
    users_count = len(storage.all_user_ids())
    bot.send_message(
        message.chat.id,
        "📢 Xabar quyidagi foydalanuvchilarga yuboriladi:\n"
        f"👥 Qabul qiluvchilar soni: {users_count}\n\n"
        "🇺🇿 O'zbek (lotin):\n"
        f"{translations['uz']}\n\n"
        "🇺🇿 Ўзбек (кирилл):\n"
        f"{translations['uz_cyrl']}\n\n"
        "🇷🇺 Русский:\n"
        f"{translations['ru']}\n\n"
        "Tasdiqlaysizmi?",
        reply_markup=broadcast_confirm_keyboard(),
    )


def _broadcast_worker(
    chat_id: int,
    actor_id: int,
    texts: dict[str, str],
) -> None:
    success = 0
    failed = 0
    for user in storage.all_users():
        if shutdown_event.is_set():
            break
        telegram_id = int(user["telegram_id"])
        language = user.get("language") or "uz"
        text = texts.get(language) or texts.get("uz") or ""
        try:
            if not send_message_slow(telegram_id, text):
                failed += 1
                continue
            success += 1
        except Exception:
            failed += 1
    audit(actor_id, "broadcast_sent", f"success={success}; failed={failed}")
    try:
        bot.send_message(
            chat_id,
            "📢 Xabar yuborish yakunlandi.\n\n"
            f"✅ Yuborildi: {success}\n"
            f"❌ Xatolik: {failed}",
            reply_markup=broadcast_admin_keyboard(actor_id),
        )
    except Exception:
        logging.exception("broadcast result message failed user_id=%s", actor_id)


def admin_broadcast_send(message: types.Message, current: dict[str, Any]) -> None:
    if message.text != CONFIRM_SEND:
        bot.send_message(
            message.chat.id,
            "Tasdiqlash uchun ✅ Yuborish tugmasini bosing.",
            reply_markup=broadcast_confirm_keyboard(),
        )
        return

    texts = current.get("broadcast_texts", {})
    clear_admin_flow(message)
    threading.Thread(
        target=_broadcast_worker,
        args=(message.chat.id, message.from_user.id, texts),
        name="broadcast-send",
        daemon=True,
    ).start()
    bot.send_message(
        message.chat.id,
        "📢 Xabar yuborish boshlandi. Jarayon fonda davom etadi.",
        reply_markup=broadcast_admin_keyboard(message.from_user.id),
    )


def save_phone(message: types.Message, abonent: Abonent, language: str) -> None:
    raw_phone = ""
    if message.contact and message.contact.phone_number:
        if (
            message.contact.user_id is not None
            and int(message.contact.user_id) != int(message.from_user.id)
        ):
            bot.send_message(
                message.chat.id,
                TEXTS[language]["enter_phone"],
                reply_markup=phone_keyboard(language),
            )
            return
        raw_phone = message.contact.phone_number
    elif message.text:
        raw_phone = message.text.strip()

    phone = normalized_uzbek_phone(raw_phone)
    if not phone:
        bot.send_message(
            message.chat.id,
            TEXTS[language]["phone_invalid"],
            reply_markup=phone_keyboard(language),
        )
        return

    previous_excel_phone = abonent.edited_phone
    try:
        with maintenance_lock:
            book.update_edited_phone(abonent.account, phone)
            try:
                storage.update_phone(message.from_user.id, phone)
            except Exception:
                book.update_edited_phone(
                    abonent.account,
                    previous_excel_phone,
                )
                raise
    except Exception as error:
        audit(
            message.from_user.id,
            "phone_update_failed",
            f"account={abonent.account}; error={error}",
        )
        notify_error("Telefon raqami saqlanmadi", error)
        bot.send_message(
            message.chat.id,
            TEXTS[language]["phone_save_failed"],
            reply_markup=phone_keyboard(language),
        )
        return
    clear_session(message)
    notify_channel(
        "Telefon raqam tahrirlandi:\n"
        + report_caption(abonent, phone, language, masked=False)
    )
    bot.send_message(message.chat.id, TEXTS[language]["phone_saved"], reply_markup=main_menu(language))


def handle_meter_photo(
    message: types.Message,
    current: dict[str, Any],
    user: dict[str, str],
    abonent: Abonent,
    language: str,
) -> None:
    text = TEXTS[language]
    photo_count = int(current.get("photo_count", 0))
    pending_photos = current.setdefault("meter_photos", [])

    if message.text == text["finish"]:
        if photo_count == 0:
            bot.send_message(message.chat.id, text["photo_needed"], reply_markup=meter_keyboard(language))
            return
        references = storage.meter_reference_fingerprints(abonent.account)
        if references:
            match = compare_photo_sets(
                [str(photo["fingerprint"]) for photo in pending_photos],
                references,
            )
            decision = match.decision
            similarity_score = match.score
        else:
            decision = "review"
            similarity_score = None

        status = {
            "approve": "auto_approved",
            "reject": "auto_rejected",
            "review": "pending_review",
        }[decision]
        try:
            submission_id = storage.save_meter_submission(
                message.from_user.id,
                abonent.account,
                language,
                abonent.inspector_fio or "Vakant",
                status,
                pending_photos,
                similarity_score,
                f"image_match_{decision}",
            )
        except Exception as error:
            audit(
                message.from_user.id,
                "meter_submission_save_failed",
                f"account={abonent.account}; error={error}",
            )
            notify_error("Hisoblagich arizasi saqlanmadi", error)
            bot.send_message(
                message.chat.id,
                text["photo_storage_failed"],
                reply_markup=meter_keyboard(language),
            )
            return

        current.pop("meter_photos", None)
        clear_session(message)
        bot.send_message(message.chat.id, text["meters_count"].format(count=photo_count))
        delivered = notify_meter_submission(
            submission_id,
            abonent,
            user.get("phone", ""),
            language,
            decision=decision,
        )
        if not delivered:
            audit(
                message.from_user.id,
                "meter_review_delivery_failed",
                f"submission_id={submission_id}; account={abonent.account}",
            )

        if decision == "approve":
            next_date = (datetime.now() + timedelta(days=30)).strftime("%d.%m.%Y")
            response = text["meter_auto_approved"].format(date=next_date)
        elif decision == "reject":
            response = text["meter_auto_rejected"]
        else:
            response = text["meter_review_pending"]
        bot.send_message(
            message.chat.id,
            response,
            reply_markup=main_menu(language),
        )
        audit(
            message.from_user.id,
            "meter_submission_created",
            f"submission_id={submission_id}; account={abonent.account}; "
            f"status={status}; photos={photo_count}",
        )
        return

    if not message.photo:
        bot.send_message(message.chat.id, text["photo_needed"], reply_markup=meter_keyboard(language))
        return
    if not item_count_allowed(photo_count, config.max_meter_photos):
        audit(
            message.from_user.id,
            "meter_photo_count_rejected",
            f"account={abonent.account}; max={config.max_meter_photos}",
        )
        bot.send_message(
            message.chat.id,
            text["photo_count_limit"].format(
                count=config.max_meter_photos,
            ),
            reply_markup=meter_keyboard(language),
        )
        return
    photo_size = int(message.photo[-1].file_size or 0)
    maximum_photo_bytes = int(config.max_meter_photo_mb * 1024 * 1024)
    if not file_size_allowed(photo_size, config.max_meter_photo_mb):
        audit(
            message.from_user.id,
            "meter_photo_size_rejected",
            f"account={abonent.account}; size={photo_size}; "
            f"max={maximum_photo_bytes}",
        )
        bot.send_message(
            message.chat.id,
            text["photo_size_limit"].format(
                size=f"{config.max_meter_photo_mb:g}",
            ),
            reply_markup=meter_keyboard(language),
        )
        return

    if not config.encryption_enabled or not config.encryption_password:
        bot.send_message(
            message.chat.id,
            text["photo_storage_failed"],
            reply_markup=meter_keyboard(language),
        )
        notify_error(
            "Hisoblagich rasmi saqlanmadi",
            "ENCRYPTION_ENABLED yoki shifrlash kaliti mavjud emas.",
        )
        return

    encrypted_path = ""
    try:
        telegram_file = bot.get_file(message.photo[-1].file_id)
        payload = bot.download_file(telegram_file.file_path)
        if not file_size_allowed(len(payload), config.max_meter_photo_mb):
            bot.send_message(
                message.chat.id,
                text["photo_size_limit"].format(size=f"{config.max_meter_photo_mb:g}"),
                reply_markup=meter_keyboard(language),
            )
            return
        fingerprint = fingerprint_image(payload)
        encrypted_path, digest = store_meter_photo(
            payload,
            config.encryption_password,
        )
        if any(photo.get("sha256") == digest for photo in pending_photos):
            remove_meter_photo(encrypted_path)
            bot.send_message(
                message.chat.id,
                text["photo_duplicate"],
                reply_markup=meter_keyboard(language),
            )
            return
        pending_photos.append(
            {
                "telegram_file_id": message.photo[-1].file_id,
                "encrypted_path": encrypted_path,
                "sha256": digest,
                "fingerprint": fingerprint,
            }
        )
    except InvalidMeterImage:
        if encrypted_path:
            remove_meter_photo(encrypted_path)
        bot.send_message(
            message.chat.id,
            text["photo_invalid"],
            reply_markup=meter_keyboard(language),
        )
        return
    except Exception as error:
        if encrypted_path:
            remove_meter_photo(encrypted_path)
        audit(
            message.from_user.id,
            "meter_photo_store_failed",
            f"account={abonent.account}; error={error}",
        )
        notify_error("Hisoblagich rasmi shifrlanib saqlanmadi", error)
        bot.send_message(
            message.chat.id,
            text["photo_storage_failed"],
            reply_markup=meter_keyboard(language),
        )
        return

    photo_count += 1
    current["photo_count"] = photo_count
    bot.send_message(
        message.chat.id,
        text["send_more_photo"],
        reply_markup=meter_keyboard(language),
    )


def _http_health_server_worker(port: int, stop_event: threading.Event) -> None:
    import http.server
    import socketserver

    health_path = Path("logs/health.json")

    class _Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            if self.path in ("/health", "/", "/status"):
                try:
                    data = health_path.read_bytes()
                    self.send_response(200)
                    self.send_header("Content-Type", "application/json; charset=utf-8")
                    self.end_headers()
                    self.wfile.write(data)
                except Exception:
                    self.send_response(503)
                    self.end_headers()
            else:
                self.send_response(404)
                self.end_headers()

        def log_message(self, *_: object) -> None:
            pass

    try:
        with socketserver.TCPServer(("127.0.0.1", port), _Handler) as server:
            server.timeout = 1
            while not stop_event.is_set():
                try:
                    server.handle_request()
                except Exception:
                    logging.debug("HTTP health server so'rovi xatosi", exc_info=True)
    except Exception as err:
        logging.exception("HTTP health server port=%s xatosi: %s", port, err)


def request_bot_shutdown() -> None:
    try:
        bot.stop_polling()
    except Exception:
        logging.debug("Bot polling to'xtatishda xato", exc_info=True)


def main() -> int:
    global config, book, storage, bot, process_started_at
    global last_preflight_status
    global request_limiter
    process_lock = ProcessLock("storage/uzwatermu.lock")
    try:
        process_lock.acquire()
    except AlreadyRunningError as error:
        print(str(error), file=sys.stderr)
        return 3

    shutdown_event.clear()
    process_started_at = datetime.now().isoformat(timespec="seconds")
    initialized = False
    try:
        verify_code_integrity()
        config = load_config()
        ensure_runtime_dirs()
        startup_report = run_system_preflight()
        if startup_report.critical:
            raise RuntimeError(preflight_summary(startup_report))
        book = AbonentBook(config.excel_file, config.inspectors_file)
        storage = BotStorage(config.database_file)
        cleanup_expired_meter_photos()
        bot = telebot.TeleBot(config.bot_token)
        request_limiter = RequestRateLimiter(
            config.request_window_seconds,
            config.flood_block_seconds,
        )
        install_shutdown_handlers(shutdown_event, request_bot_shutdown)
        bot.register_message_handler(handle_start, commands=["start"])
        bot.register_message_handler(handle_myid, commands=["myid"])
        bot.register_message_handler(handle_admin, commands=["admin"])
        bot.register_callback_query_handler(
            handle_approval_callback,
            func=lambda call: bool(call.data)
            and call.data.startswith(("approve:", "reject:")),
        )
        bot.register_callback_query_handler(
            handle_meter_review_callback,
            func=lambda call: bool(call.data)
            and call.data.startswith(("meterapprove:", "meterreject:")),
        )
        bot.register_message_handler(
            handle_message,
            content_types=[
                "text",
                "photo",
                "contact",
                "document",
                "location",
            ],
        )
        bot.register_edited_message_handler(
            handle_edited_inspector_location,
            content_types=["location"],
        )
        start_debt_analysis_worker()
        start_inspector_location_monitor_worker()
        start_automatic_backup_worker()
        last_preflight_status = "|".join(
            f"{item.name}:{item.level}:{item.message}"
            for item in startup_report.results
        )
        if startup_report.status != "ok":
            notify_error(
                "Startup preflight ogohlantirishi",
                preflight_summary(startup_report),
            )
        start_integrity_monitor_worker()
        threading.Thread(
            target=health_worker,
            args=(
                shutdown_event,
                "logs/health.json",
                process_started_at,
            ),
            name="health-monitor",
            daemon=True,
        ).start()
        if config.health_port:
            threading.Thread(
                target=_http_health_server_worker,
                args=(config.health_port, shutdown_event),
                name="health-http",
                daemon=True,
            ).start()
        initialized = True
        audit(None, "bot_started", "infinity_polling")
        while not shutdown_event.is_set():
            try:
                bot.infinity_polling(
                    skip_pending=True,
                    timeout=30,
                    long_polling_timeout=30,
                )
            except Exception as error:
                if shutdown_event.is_set():
                    break
                notify_error(
                    "Bot polling to'xtadi, qayta uriniladi",
                    redact_sensitive_text(
                        f"{safe_exception_summary(error)}\n\n"
                        f"{traceback.format_exc()[-2500:]}",
                        3000,
                    ),
                )
                shutdown_event.wait(10)
        return 0
    finally:
        shutdown_event.set()
        request_bot_shutdown()
        if initialized:
            try:
                audit(None, "bot_stopped", "graceful_shutdown")
            except Exception:
                logging.debug("Shutdown audit yozishda xato", exc_info=True)
            write_health_status(
                "logs/health.json",
                state="stopped",
                started_at=process_started_at,
                details="graceful_shutdown",
            )
        process_lock.release()


if __name__ == "__main__":
    raise SystemExit(main())

