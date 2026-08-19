import shutil
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from bot.excel_loader import AbonentBook


class ExcelLoaderCompletenessTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path("work/tests/excel_loader_completeness")
        shutil.rmtree(self.root, ignore_errors=True)
        self.root.mkdir(parents=True)
        self.path = self.root / "abonents.xlsx"

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def test_account_with_blank_fio_is_not_dropped(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "mahalla",
                "ko'cha",
                "uy",
                "xonadon",
                "fio",
                "hisob raqam",
            ]
        )
        sheet.append(
            [
                "Zakovat",
                "Olimlar",
                "51",
                "1",
                "",
                "1251050256",
            ]
        )
        workbook.save(self.path)
        workbook.close()

        book = AbonentBook(str(self.path))
        abonent = book.find("1251050256")

        self.assertIsNotNone(abonent)
        self.assertEqual(abonent.fio, "-")
        self.assertEqual(len(book._items), 1)

    def test_edited_phone_is_written_to_excel_and_reloaded(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "mahalla",
                "ko'cha",
                "uy",
                "xonadon",
                "fio",
                "hisob raqam",
                "telefon",
                "yangi telefon",
            ]
        )
        sheet.append(
            [
                "Zakovat",
                "Olimlar",
                "51",
                "1",
                "Test User",
                "1251050256",
                "901112233",
                "",
            ]
        )
        workbook.save(self.path)
        workbook.close()

        book = AbonentBook(str(self.path))
        book.update_edited_phone("1251050256", "+998909998877")

        updated = load_workbook(self.path, read_only=True, data_only=True)
        try:
            self.assertEqual(
                "+998909998877",
                updated.active.cell(row=2, column=8).value,
            )
        finally:
            updated.close()
        self.assertEqual(
            "+998909998877",
            book.find("1251050256").edited_phone,
        )
        self.assertEqual(
            "+998909998877",
            book.find("1251050256").phone,
        )

    def test_phone_update_requires_new_phone_column(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["hisob raqam", "telefon"])
        sheet.append(["1251050256", "901112233"])
        workbook.save(self.path)
        workbook.close()

        book = AbonentBook(str(self.path))
        with self.assertRaisesRegex(RuntimeError, "yangi telefon"):
            book.update_edited_phone("1251050256", "+998909998877")


if __name__ == "__main__":
    unittest.main()
