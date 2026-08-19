from __future__ import annotations

import json
import shutil
import sqlite3
import unittest
from contextlib import closing
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from bot.backup_manager import (
    BackupFile,
    apply_retention,
    create_backup_set,
    latest_valid_backup,
    list_backup_sets,
    load_backup_set,
    restore_backup_set,
)
from bot.encryption import InvalidPasswordOrData


ENCRYPTION_PASSWORD = "Backup-Encryption-Password-2026!"


class BackupManagerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path("work/tests/backup_manager") / self._testMethodName
        shutil.rmtree(self.root, ignore_errors=True)
        self.source_dir = self.root / "source"
        self.backup_dir = self.root / "backups"
        self.source_dir.mkdir(parents=True, exist_ok=True)

        self.database = self.source_dir / "bot.sqlite3"
        with closing(sqlite3.connect(self.database)) as connection:
            connection.execute("create table sample (value text not null)")
            connection.execute("insert into sample values ('original')")
            connection.commit()

        self.abonents = self.source_dir / "abonents.xlsx"
        self.inspectors = self.source_dir / "inspectors.xlsx"
        self._write_excel(self.abonents, "original-abonent")
        self._write_excel(self.inspectors, "original-inspector")

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    @staticmethod
    def _write_excel(path: Path, value: str) -> None:
        workbook = Workbook()
        workbook.active["A1"] = value
        workbook.save(path)
        workbook.close()

    def _sources(self) -> list[BackupFile]:
        return [
            BackupFile("database", self.database, "bot.sqlite3"),
            BackupFile("abonents", self.abonents, "abonents.xlsx"),
            BackupFile("inspectors", self.inspectors, "inspectors.xlsx"),
        ]

    def test_backup_set_contains_valid_manifest_and_all_files(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "unit_test",
            datetime(2026, 6, 13, 10, 0, 0),
        )

        loaded = load_backup_set(backup.directory, verify_hashes=True)
        manifest = json.loads(
            (backup.directory / "manifest.json").read_text(encoding="utf-8")
        )

        self.assertTrue(loaded.valid)
        self.assertEqual(3, len(loaded.files))
        self.assertEqual("unit_test", manifest["reason"])
        self.assertEqual(
            {"database", "abonents", "inspectors"},
            {item["label"] for item in manifest["files"]},
        )

    def test_changed_backup_file_is_rejected_by_hash_check(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "hash_test",
        )
        (backup.directory / "abonents.xlsx").write_bytes(b"corrupted")

        with self.assertRaises(RuntimeError):
            load_backup_set(backup.directory, verify_hashes=True)
        self.assertIsNone(latest_valid_backup(self.backup_dir))

    def test_retention_keeps_only_newest_configured_sets(self) -> None:
        start = datetime(2026, 6, 1, 8, 0, 0)
        for offset in range(4):
            create_backup_set(
                self.backup_dir,
                self._sources(),
                f"set_{offset}",
                start + timedelta(days=offset),
            )

        removed = apply_retention(
            self.backup_dir,
            keep_sets=2,
            retention_days=30,
            now=start + timedelta(days=4),
        )

        backups = list_backup_sets(self.backup_dir)
        self.assertEqual(2, len(removed))
        self.assertEqual(["set_3", "set_2"], [item.reason for item in backups])

    def test_restore_returns_database_and_excel_files_to_backup_state(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "restore_test",
        )

        with closing(sqlite3.connect(self.database)) as connection:
            connection.execute("update sample set value = 'changed'")
            connection.commit()
        self._write_excel(self.abonents, "changed-abonent")
        self._write_excel(self.inspectors, "changed-inspector")

        restored = restore_backup_set(
            backup,
            {
                "database": self.database,
                "abonents": self.abonents,
                "inspectors": self.inspectors,
            },
        )

        with closing(sqlite3.connect(self.database)) as connection:
            value = connection.execute("select value from sample").fetchone()[0]
        abonents_book = load_workbook(self.abonents, read_only=True)
        inspectors_book = load_workbook(self.inspectors, read_only=True)
        try:
            self.assertEqual("original", value)
            self.assertEqual("original-abonent", abonents_book.active["A1"].value)
            self.assertEqual(
                "original-inspector",
                inspectors_book.active["A1"].value,
            )
        finally:
            abonents_book.close()
            inspectors_book.close()

        self.assertEqual(
            {self.database, self.abonents, self.inspectors},
            set(restored),
        )

    def test_encrypted_backup_contains_no_plain_database_or_excel(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "encrypted_test",
            encryption_password=ENCRYPTION_PASSWORD,
        )

        names = {path.name for path in backup.files}
        self.assertEqual(
            {
                "bot.sqlite3.uwenc",
                "abonents.xlsx.uwenc",
                "inspectors.xlsx.uwenc",
            },
            names,
        )
        self.assertFalse((backup.directory / "bot.sqlite3").exists())
        self.assertFalse((backup.directory / "abonents.xlsx").exists())
        self.assertFalse((backup.directory / "inspectors.xlsx").exists())
        loaded = load_backup_set(
            backup.directory,
            verify_hashes=True,
            encryption_password=ENCRYPTION_PASSWORD,
        )
        self.assertEqual(3, len(loaded.files))

    def test_encrypted_backup_rejects_wrong_password(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "wrong_password_test",
            encryption_password=ENCRYPTION_PASSWORD,
        )
        with self.assertRaises(InvalidPasswordOrData):
            load_backup_set(
                backup.directory,
                verify_hashes=True,
                encryption_password="Wrong-backup-password-2026!",
            )

    def test_encrypted_backup_can_be_restored(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "encrypted_restore_test",
            encryption_password=ENCRYPTION_PASSWORD,
        )
        with closing(sqlite3.connect(self.database)) as connection:
            connection.execute("update sample set value = 'changed'")
            connection.commit()
        self._write_excel(self.abonents, "changed-abonent")

        restore_backup_set(
            backup,
            {
                "database": self.database,
                "abonents": self.abonents,
                "inspectors": self.inspectors,
            },
            encryption_password=ENCRYPTION_PASSWORD,
        )

        with closing(sqlite3.connect(self.database)) as connection:
            value = connection.execute("select value from sample").fetchone()[0]
        workbook = load_workbook(self.abonents, read_only=True)
        try:
            self.assertEqual("original", value)
            self.assertEqual("original-abonent", workbook.active["A1"].value)
        finally:
            workbook.close()

    def test_manifest_rejects_path_traversal(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "unsafe_manifest",
        )
        manifest_path = backup.directory / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["files"][0]["name"] = "../outside.sqlite3"
        manifest_path.write_text(
            json.dumps(manifest),
            encoding="utf-8",
        )

        with self.assertRaises(RuntimeError):
            load_backup_set(backup.directory)

    def test_empty_manifest_is_not_a_valid_backup(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "empty_manifest",
        )
        manifest_path = backup.directory / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["files"] = []
        manifest_path.write_text(
            json.dumps(manifest),
            encoding="utf-8",
        )

        with self.assertRaises(RuntimeError):
            load_backup_set(backup.directory)

    def test_partial_restore_failure_rolls_back_replaced_files(self) -> None:
        backup = create_backup_set(
            self.backup_dir,
            self._sources(),
            "rollback_test",
        )
        self._write_excel(self.abonents, "current-abonent")
        self._write_excel(self.inspectors, "current-inspector")
        original_replace = Path.replace

        def fail_on_inspector(source: Path, target: Path) -> Path:
            if (
                source.name == ".inspectors.restore.tmp.xlsx"
                and Path(target) == self.inspectors
            ):
                raise OSError("simulated restore failure")
            return original_replace(source, target)

        with patch.object(Path, "replace", fail_on_inspector):
            with self.assertRaises(OSError):
                restore_backup_set(
                    backup,
                    {
                        "database": self.database,
                        "abonents": self.abonents,
                        "inspectors": self.inspectors,
                    },
                )

        abonents_book = load_workbook(self.abonents, read_only=True)
        inspectors_book = load_workbook(self.inspectors, read_only=True)
        try:
            self.assertEqual(
                "current-abonent",
                abonents_book.active["A1"].value,
            )
            self.assertEqual(
                "current-inspector",
                inspectors_book.active["A1"].value,
            )
        finally:
            abonents_book.close()
            inspectors_book.close()


class BackupAdminRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.source = Path("bot/main.py").read_text(encoding="utf-8")

    def test_restore_is_super_admin_only_and_requires_confirmation(self) -> None:
        self.assertIn("if message_text == ADMIN_BACKUP_RESTORE:", self.source)
        self.assertIn("if not is_super_admin_id(message.from_user.id):", self.source)
        self.assertIn("if message.text != BACKUP_RESTORE_CONFIRM:", self.source)

    def test_restore_creates_safety_backup_before_replacement(self) -> None:
        restore_start = self.source.index("def admin_restore_backup(")
        restore_end = self.source.index("\ndef admin_search(", restore_start)
        restore_source = self.source[restore_start:restore_end]
        safety_index = restore_source.index(
            'create_backup(message.from_user.id, "before_restore")'
        )
        replacement_index = restore_source.index("restore_backup_set(")
        self.assertLess(safety_index, replacement_index)

    def test_automatic_worker_is_started(self) -> None:
        self.assertIn("def automatic_backup_worker()", self.source)
        self.assertIn("start_automatic_backup_worker()", self.source)


if __name__ == "__main__":
    unittest.main()
