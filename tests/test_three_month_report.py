import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from bot.table_reports import latest_three_month_report, three_month_cutoff


ROOT = Path(__file__).resolve().parents[1]
REPORT = (
    ROOT
    / "outputs"
    / "table_reports"
    / "3_oylik_korsatgich_olinmagan_2026-06.xlsx"
)


class ThreeMonthCutoffTests(unittest.TestCase):
    def test_june_cutoff_is_march_end(self) -> None:
        self.assertEqual(
            three_month_cutoff(date(2026, 6, 11)),
            date(2026, 3, 31),
        )

    def test_july_cutoff_is_april_end(self) -> None:
        self.assertEqual(
            three_month_cutoff(date(2026, 7, 1)),
            date(2026, 4, 30),
        )

    def test_year_boundary_is_supported(self) -> None:
        self.assertEqual(
            three_month_cutoff(date(2026, 1, 15)),
            date(2025, 10, 31),
        )


class ThreeMonthWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not REPORT.exists():
            raise AssertionError(f"Hisobot topilmadi: {REPORT}")
        cls.workbook = load_workbook(REPORT, read_only=True, data_only=False)
        cls.sheet = cls.workbook.active

    @classmethod
    def tearDownClass(cls) -> None:
        cls.workbook.close()

    def test_removed_column_shifts_area_and_inspector(self) -> None:
        self.assertEqual(self.sheet.max_column, 19)
        self.assertEqual(self.sheet["R3"].value, "Участок")
        self.assertEqual(self.sheet["S3"].value, "Инспектор")
        self.assertIsNone(self.sheet["T3"].value)

    def test_footer_formulas_follow_filtered_rows(self) -> None:
        last_data_row = self.sheet.max_row - 1
        self.assertEqual(
            self.sheet.cell(self.sheet.max_row, 9).value,
            f"=SUBTOTAL(9,I6:I{last_data_row})",
        )
        self.assertEqual(
            self.sheet.cell(self.sheet.max_row, 14).value,
            f"=SUBTOTAL(9,N6:N{last_data_row})",
        )

    def test_all_rows_match_meter_and_date_filters(self) -> None:
        cutoff = date(2026, 3, 31)
        for row in self.sheet.iter_rows(
            min_row=6,
            max_row=self.sheet.max_row - 1,
            min_col=8,
            max_col=12,
            values_only=True,
        ):
            self.assertEqual(str(row[0]).strip(), "+")
            last_reading_date = row[4]
            if hasattr(last_reading_date, "date"):
                last_reading_date = last_reading_date.date()
            if isinstance(last_reading_date, date):
                self.assertLessEqual(last_reading_date, cutoff)

    def test_latest_report_can_be_found_for_bot_download(self) -> None:
        latest = latest_three_month_report(REPORT.parent)
        self.assertEqual(latest, REPORT)


if __name__ == "__main__":
    unittest.main()
